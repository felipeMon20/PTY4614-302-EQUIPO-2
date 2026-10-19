import os
import shutil
import tempfile
import time
import unittest
from unittest.mock import patch

import openpyxl
import pandas as pd

import config.settings
from src.core.data_manager import DataManager
from src.ui.historial import HistorialModule


def _read_excel_con_reintento(path, intentos=5, espera_seg=0.3, **kwargs):
    """Igual que src.ui.provisiones._copiar_con_reintento, pero para releer
    (no copiar): en Windows, el SO/antivirus puede retener por unos
    milisegundos el lock sobre un .xlsx recien escrito (WinError 32), aunque
    el propio codigo de produccion ya haya hecho wb.save()+wb.close(). No es
    un handle que la app deje abierto (ver _clear_papelera en historial.py);
    es la misma condicion de carrera de SO que _copiar_con_reintento ya
    tolera en otro modulo."""
    ultimo_error = None
    for _ in range(intentos):
        try:
            return pd.read_excel(path, **kwargs)
        except PermissionError as e:
            ultimo_error = e
            time.sleep(espera_seg)
    raise ultimo_error


class TestClearLogsBloqueadoTemporalmente(unittest.TestCase):
    """_clear_logs() (boton 'Vaciar Pestaña Actual'): ambas pestañas quedan
    habilitadas (28-09-2026, a pedido explícito -- antes Logs de
    Modificaciones permanecía bloqueado porque esa operación no tenía
    protección de backup propia). Papelera (RegistroEliminado.xlsx) ya tenía
    su protección (ver TestVaciarPapeleraConBackupYMultiplesHojas); Logs de
    Modificaciones (LogApp.csv) ahora también respalda antes de vaciar (ver
    TestVaciarLogsModificacionesConBackup)."""

    def setUp(self):
        self.dm = DataManager()
        self.modulo = HistorialModule()

    def tearDown(self):
        for signal in (self.dm.logs_updated, self.dm.data_ready):
            try:
                signal.disconnect(self.modulo.load_data)
            except Exception:
                pass

    def test_boton_habilitado_en_pestana_modificaciones(self):
        self.modulo.tabs.setCurrentIndex(0)  # "Logs de Modificaciones"
        self.assertTrue(self.modulo.btn_clear.isEnabled(),
                         "Fallo: en Logs de Modificaciones el botón debe estar habilitado (ya tiene protección de backup).")
        self.assertIn("backup", self.modulo.btn_clear.toolTip().lower(),
                       "Fallo: el tooltip debe mencionar la protección de backup.")

    def test_boton_habilitado_en_pestana_papelera(self):
        self.modulo.tabs.setCurrentIndex(1)  # "Papelera"
        self.assertTrue(self.modulo.btn_clear.isEnabled(),
                         "Fallo: en Papelera el botón debe estar habilitado (ya tiene protección de backup).")


class TestVaciarLogsModificacionesConBackup(unittest.TestCase):
    """_clear_logs_modificaciones() (invocada por _clear_logs() en la
    pestaña Logs de Modificaciones, 28-09-2026): debe crear un backup previo
    de LogApp.csv (reutilizando ExcelWriter._crear_backup), abortar sin
    borrar nada si el backup falla, y dejar solo la línea de encabezado."""

    def setUp(self):
        self.dm = DataManager()
        self.temp_dir = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.log_csv_path = os.path.join(self.temp_dir.name, "LogApp_clear_test.csv")
        with open(self.log_csv_path, 'w', encoding='utf-8-sig', newline='') as f:
            f.write("FECHA Y HORA,MÓDULO,VIN,CAMPO,VALOR ANTERIOR,VALOR NUEVO\r\n")
            f.write("2026-08-01 09:00:00,Comercial,YV2AAA,CDB,100,200\r\n")
            f.write("2026-08-02 09:00:00,Comercial,YV2BBB,CDB,300,400\r\n")

        self._patcher = patch.object(config.settings, 'LOG_CSV_PATH', self.log_csv_path)
        self._patcher.start()

        self.modulo = HistorialModule()
        self.modulo.writer.log_csv_filepath = self.log_csv_path
        self.modulo.writer.backup_dir = os.path.join(self.temp_dir.name, "backups")

    def tearDown(self):
        self._patcher.stop()
        for signal in (self.dm.logs_updated, self.dm.data_ready):
            try:
                signal.disconnect(self.modulo.load_data)
            except Exception:
                pass
        self.temp_dir.cleanup()

    def _confirmar_dialogo_y_vaciar(self):
        from PyQt6.QtWidgets import QMessageBox
        with patch.object(QMessageBox, 'exec', return_value=QMessageBox.StandardButton.Yes), \
             patch('src.ui.historial.QMessageBox.information'), \
             patch('src.ui.historial.QMessageBox.warning'):
            self.modulo.tabs.setCurrentIndex(0)  # pestaña "Logs de Modificaciones"
            self.modulo._clear_logs()

    def test_backup_se_crea_antes_de_vaciar(self):
        self.assertFalse(os.path.exists(self.modulo.writer.backup_dir))
        self._confirmar_dialogo_y_vaciar()
        self.assertTrue(os.path.exists(self.modulo.writer.backup_dir),
                         "Fallo: debe crearse el directorio/archivo de backup antes de vaciar LogApp.csv.")
        backups = os.listdir(self.modulo.writer.backup_dir)
        self.assertTrue(any("LogApp_clear_test" in b for b in backups),
                         "Fallo: debe existir un backup del archivo real antes de la operación destructiva.")

    def test_backup_contiene_los_datos_originales(self):
        self._confirmar_dialogo_y_vaciar()
        backups = os.listdir(self.modulo.writer.backup_dir)
        backup_path = os.path.join(self.modulo.writer.backup_dir, backups[0])
        with open(backup_path, encoding='utf-8-sig') as f:
            lineas = f.readlines()
        self.assertEqual(len(lineas), 3, "Fallo: el backup debe conservar las filas originales.")

    def test_solo_queda_el_encabezado(self):
        self._confirmar_dialogo_y_vaciar()
        with open(self.log_csv_path, encoding='utf-8-sig') as f:
            lineas = f.readlines()
        self.assertEqual(len(lineas), 1, "Fallo: solo debe quedar la línea de encabezado.")
        self.assertTrue(lineas[0].startswith("FECHA Y HORA"), "Fallo: el encabezado debe conservarse intacto.")

    def test_si_backup_falla_no_se_elimina_ningun_registro(self):
        with open(self.log_csv_path, 'rb') as f:
            contenido_antes = f.read()

        from PyQt6.QtWidgets import QMessageBox
        with patch.object(QMessageBox, 'exec', return_value=QMessageBox.StandardButton.Yes), \
             patch('src.ui.historial.QMessageBox.information'), \
             patch('src.ui.historial.QMessageBox.warning') as mock_warning, \
             patch.object(self.modulo.writer, '_crear_backup', return_value=False):
            self.modulo.tabs.setCurrentIndex(0)
            self.modulo._clear_logs()

        mock_warning.assert_called_once()
        with open(self.log_csv_path, 'rb') as f:
            contenido_despues = f.read()
        self.assertEqual(contenido_antes, contenido_despues,
                          "Fallo: si el backup falla, LogApp.csv no debe modificarse.")

    def test_vaciar_ya_vacio_no_falla_y_avisa(self):
        self._confirmar_dialogo_y_vaciar()
        from PyQt6.QtWidgets import QMessageBox
        with patch.object(QMessageBox, 'exec', return_value=QMessageBox.StandardButton.Yes), \
             patch('src.ui.historial.QMessageBox.information') as mock_info, \
             patch('src.ui.historial.QMessageBox.warning') as mock_warning:
            self.modulo.tabs.setCurrentIndex(0)
            self.modulo._clear_logs()
        mock_warning.assert_not_called()
        mock_info.assert_called_once()
        with open(self.log_csv_path, encoding='utf-8-sig') as f:
            lineas = f.readlines()
        self.assertEqual(len(lineas), 1)


class TestVaciarPapeleraConBackupYMultiplesHojas(unittest.TestCase):
    """_clear_papelera() (invocada por _clear_logs() en la pestaña Papelera)
    debe: crear un backup previo reutilizando ExcelWriter._crear_backup,
    abortar sin borrar nada si el backup falla, y vaciar TODAS las hojas de
    RegistroEliminado.xlsx (una por módulo) sin hacerlas desaparecer."""

    def setUp(self):
        self.dm = DataManager()
        # ignore_cleanup_errors=True: en Windows, el SO/antivirus puede retener
        # por unos milisegundos el lock sobre el .xlsx recien escrito/leido de
        # esta prueba (WinError 32) incluso despues de que el codigo de
        # produccion ya cerro sus propios handles. No es un handle que la app
        # deje abierto (ver _clear_papelera en historial.py); es la misma
        # condicion de carrera de SO que _copiar_con_reintento ya tolera en
        # otro modulo. Un archivo temporal que sobrevive a un cleanup fallido
        # no afecta la correccion de la prueba, solo la prolijidad del temp dir.
        self.temp_dir = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.deleted_path = os.path.join(self.temp_dir.name, "RegistroEliminado_clear_test.xlsx")

        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "Comercial"
        ws1.append(["FECHA Y HORA DE ELIMINACIÓN", "MÓDULO", "VIN", "MODELO"])
        ws1.append(["2026-08-01 09:00:00", "Comercial", "YV2AAA", "FM"])
        ws1.append(["2026-08-02 09:00:00", "Comercial", "YV2BBB", "FMH"])
        ws2 = wb.create_sheet("Contactos")
        ws2.append(["FECHA Y HORA DE ELIMINACIÓN", "MÓDULO", "ID", "Nombre"])
        ws2.append(["2026-08-03 10:00:00", "Contactos", "458", "Ana"])
        wb.save(self.deleted_path)
        wb.close()

        self._patcher = patch.object(config.settings, 'DELETED_LOG_EXCEL_PATH', self.deleted_path)
        self._patcher.start()

        self.modulo = HistorialModule()
        self.modulo.writer.deleted_filepath = self.deleted_path
        self.modulo.writer.backup_dir = os.path.join(self.temp_dir.name, "backups")

    def tearDown(self):
        self._patcher.stop()
        for signal in (self.dm.logs_updated, self.dm.data_ready):
            try:
                signal.disconnect(self.modulo.load_data)
            except Exception:
                pass
        self.temp_dir.cleanup()

    def _confirmar_dialogo_y_vaciar(self):
        # Simula al usuario confirmando el diálogo de advertencia: se
        # parchea QMessageBox.exec para que devuelva "Yes" sin requerir una
        # interacción real.
        from PyQt6.QtWidgets import QMessageBox
        with patch.object(QMessageBox, 'exec', return_value=QMessageBox.StandardButton.Yes), \
             patch('src.ui.historial.QMessageBox.information'), \
             patch('src.ui.historial.QMessageBox.warning'):
            self.modulo.tabs.setCurrentIndex(1)  # pestaña "Papelera"
            self.modulo._clear_logs()

    def test_backup_se_crea_antes_de_vaciar(self):
        self.assertFalse(os.path.exists(self.modulo.writer.backup_dir))
        self._confirmar_dialogo_y_vaciar()
        self.assertTrue(os.path.exists(self.modulo.writer.backup_dir),
                         "Fallo: debe crearse el directorio/archivo de backup antes de vaciar la Papelera.")
        backups = os.listdir(self.modulo.writer.backup_dir)
        self.assertTrue(any("RegistroEliminado_clear_test" in b for b in backups),
                         "Fallo: debe existir un backup del archivo real antes de la operación destructiva.")

    def test_backup_contiene_los_datos_originales(self):
        self._confirmar_dialogo_y_vaciar()
        backups = os.listdir(self.modulo.writer.backup_dir)
        backup_path = os.path.join(self.modulo.writer.backup_dir, backups[0])
        wb_backup = openpyxl.load_workbook(backup_path)
        self.assertEqual(wb_backup["Comercial"].max_row, 3, "Fallo: el backup debe conservar los datos originales.")
        wb_backup.close()

    def test_todas_las_hojas_quedan_vacias_pero_existen(self):
        self._confirmar_dialogo_y_vaciar()
        wb = openpyxl.load_workbook(self.deleted_path)
        self.assertIn("Comercial", wb.sheetnames, "Fallo: la hoja Comercial no debe desaparecer.")
        self.assertIn("Contactos", wb.sheetnames, "Fallo: la hoja Contactos no debe desaparecer.")
        self.assertEqual(wb["Comercial"].max_row, 1, "Fallo: solo debe quedar el encabezado en Comercial.")
        self.assertEqual(wb["Contactos"].max_row, 1, "Fallo: solo debe quedar el encabezado en Contactos.")
        self.assertEqual(
            [c.value for c in wb["Comercial"][1]],
            ["FECHA Y HORA DE ELIMINACIÓN", "MÓDULO", "VIN", "MODELO"],
            "Fallo: el encabezado de Comercial debe conservarse intacto.",
        )
        wb.close()

    def test_si_backup_falla_no_se_elimina_ningun_registro(self):
        with open(self.deleted_path, 'rb') as f:
            contenido_antes = f.read()

        from PyQt6.QtWidgets import QMessageBox
        with patch.object(QMessageBox, 'exec', return_value=QMessageBox.StandardButton.Yes), \
             patch('src.ui.historial.QMessageBox.information'), \
             patch('src.ui.historial.QMessageBox.warning') as mock_warning, \
             patch.object(self.modulo.writer, '_crear_backup', return_value=False):
            self.modulo.tabs.setCurrentIndex(1)
            self.modulo._clear_logs()

        mock_warning.assert_called_once()
        with open(self.deleted_path, 'rb') as f:
            contenido_despues = f.read()
        self.assertEqual(contenido_antes, contenido_despues,
                          "Fallo: si el backup falla, RegistroEliminado.xlsx no debe modificarse.")

    def test_archivo_queda_valido_y_se_puede_recargar(self):
        self._confirmar_dialogo_y_vaciar()
        # No debe lanzar excepcion al releer el archivo (formato valido) y
        # DataManager/UI deben poder recargar todas las hojas normalmente.
        # Reintento acotado: el SO/antivirus puede retener el lock unos
        # milisegundos tras el wb.save()+wb.close() de _clear_papelera (ver
        # _read_excel_con_reintento arriba).
        hojas = _read_excel_con_reintento(self.deleted_path, sheet_name=None)
        self.assertIn("Comercial", hojas)
        self.assertIn("Contactos", hojas)
        self.assertEqual(len(hojas["Comercial"]), 0)
        self.assertEqual(len(hojas["Contactos"]), 0)

    def test_vaciar_papelera_ya_vacia_no_falla_y_avisa(self):
        # Primera pasada: vacía la papelera.
        self._confirmar_dialogo_y_vaciar()
        # Segunda pasada: ya no hay filas de datos en ninguna hoja.
        from PyQt6.QtWidgets import QMessageBox
        with patch.object(QMessageBox, 'exec', return_value=QMessageBox.StandardButton.Yes), \
             patch('src.ui.historial.QMessageBox.information') as mock_info, \
             patch('src.ui.historial.QMessageBox.warning') as mock_warning:
            self.modulo.tabs.setCurrentIndex(1)
            self.modulo._clear_logs()
        mock_warning.assert_not_called()
        mock_info.assert_called_once()
        wb = openpyxl.load_workbook(self.deleted_path)
        self.assertIn("Comercial", wb.sheetnames)
        self.assertIn("Contactos", wb.sheetnames)
        wb.close()


class TestPapeleraPorModulo(unittest.TestCase):
    """La pestaña Papelera debe mostrar cada hoja de RegistroEliminado.xlsx
    (un módulo por hoja, ver DataManager.deleted_sheets) por separado, sin
    mezclar columnas entre módulos, y 'Restaurar Fila' debe estar bloqueada
    en vez de intentar llamar a un método que no existe (ExcelWriter no
    implementa restore_row)."""

    def setUp(self):
        self.dm = DataManager()
        self.dm.deleted_sheets = {
            "Comercial": pd.DataFrame({
                "FECHA Y HORA DE ELIMINACIÓN": ["2026-08-01 09:00:00"],
                "MÓDULO": ["Comercial"],
                "VIN": ["YV2AAA"],
                "MODELO": ["FM"],
            }),
            "Contactos": pd.DataFrame({
                "FECHA Y HORA DE ELIMINACIÓN": ["2026-08-02 10:00:00"],
                "MÓDULO": ["Contactos"],
                "ID": ["458"],
                "Nombre": ["Ana"],
            }),
        }
        self.modulo = HistorialModule()

    def tearDown(self):
        for signal in (self.dm.logs_updated, self.dm.data_ready):
            try:
                signal.disconnect(self.modulo.load_data)
            except Exception:
                pass
        self.dm.deleted_sheets = {}

    def test_combo_lista_todos_los_modulos_con_registros_eliminados(self):
        items = [self.modulo.combo_papelera_modulo.itemText(i) for i in range(self.modulo.combo_papelera_modulo.count())]
        self.assertIn("Comercial", items)
        self.assertIn("Contactos", items)

    def test_seleccionar_un_modulo_no_mezcla_columnas_de_otro(self):
        self.modulo.combo_papelera_modulo.setCurrentText("Comercial")
        self.assertIn("VIN", self.modulo.filtered_del.columns)
        self.assertNotIn("ID", self.modulo.filtered_del.columns, "Fallo: no debe traer columnas de Contactos.")

        self.modulo.combo_papelera_modulo.setCurrentText("Contactos")
        self.assertIn("ID", self.modulo.filtered_del.columns)
        self.assertNotIn("VIN", self.modulo.filtered_del.columns, "Fallo: no debe traer columnas de Comercial.")

    def test_handle_restaurar_no_crashea_y_avisa_que_no_esta_disponible(self):
        with patch('src.ui.historial.QMessageBox.information') as mock_info:
            self.modulo._handle_restaurar(0, "YV2AAA")
        mock_info.assert_called_once()


if __name__ == '__main__':
    unittest.main(verbosity=2)

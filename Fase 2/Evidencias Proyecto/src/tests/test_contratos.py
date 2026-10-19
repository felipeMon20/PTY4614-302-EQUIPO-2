import os
import tempfile
import unittest
from unittest.mock import patch

import openpyxl
import pandas as pd
from PyQt6.QtWidgets import QApplication, QMessageBox, QFileDialog, QDialog

import config.settings
from src.ui.contratos import ContratosModule, SaveContratosThread
from src.core.data_manager import DataManager
from src.services.cm_bajas_service import (
    leer_casos_baja_cm, ESTADO_CASO_PENDIENTE, ESTADO_CASO_REAPARECIDO, ORIGEN_CASO_ESTADO_EXPIRED,
)


class TestContratosNoLimpiaAutomaticamenteVinAusente(unittest.TestCase):
    """Regresion critica (regla confirmada con la contraparte comercial): que un VIN
    desaparezca del Excel mensual de CM NO debe limpiar automaticamente sus
    datos CM en PruebasFM. Antes de esta correccion,
    ContratosModule.importar_excel_reporte() ponia '-'/0.0 en los campos CM
    de cualquier VIN de PruebasFM que no apareciera en el archivo importado,
    sin ninguna revision. Este test cubre exactamente ese camino."""

    @classmethod
    def setUpClass(cls):
        if QApplication.instance() is None:
            cls._app = QApplication([])

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.import_path = os.path.join(self.temp_dir.name, "CM_Mensual.xlsx")

        # Excel mensual de CM: solo trae VIN_PRESENTE. VIN_AUSENTE
        # (que ya tenia datos CM en PruebasFM) no aparece este mes.
        df_import = pd.DataFrame({
            "VIN No": ["VIN_PRESENTE"],
            "Contract No": ["3PC-999"],
            "Start Date": ["2026-01-01"],
            "Expiring Date": ["2027-01-01"],
            "Contract Status": ["Running"],
            "Contract Type": ["Gold"],
            "Net Amount": [20000],
        })
        df_import.to_excel(self.import_path, index=False)

        self.df_master_original = pd.DataFrame({
            "VIN": ["VIN_PRESENTE", "VIN_AUSENTE"],
            "CHASIS": ["C1", "C2"],
            "N CONTRATO 3PC": ["3PC-111", "3PC-222"],
            "FECHA INICIO CM": ["01-01-2025", "01-01-2025"],
            "FECHA TERMINO CM": ["01-01-2026", "01-01-2026"],
            "TIPO DE CM": ["Silver", "Silver"],
            "ESTADO DEL CM": ["Running", "Running"],
            "PAGO CM": [15000, 18000],
        })

        self.dm = DataManager()
        self.dm.df_master = self.df_master_original.copy()

        self.modulo = ContratosModule()
        self.modulo.data_manager.df_master = self.df_master_original.copy()
        self.modulo.load_data()

        self.casos_baja_cm_path = os.path.join(self.temp_dir.name, "CasosBajaCM_Test.csv")

        self._patches = [
            patch('src.ui.contratos.QFileDialog.getOpenFileName', return_value=(self.import_path, '')),
            # ResumenImportacionDialog (30-09-2026) reemplazó al QMessageBox.question
            # de confirmación de import masivo: se mockea su exec() en vez del
            # QMessageBox, para no abrir un diálogo modal real en los tests.
            patch('src.ui.contratos.ResumenImportacionDialog.exec', return_value=QDialog.DialogCode.Accepted),
            # _on_import_finished(True) muestra un QMessageBox.information real
            # al terminar: sin mockearlo, el dialogo modal queda esperando un
            # clic que nunca llega y el test se congela (bug real detectado
            # con faulthandler bajo QT_QPA_PLATFORM=offscreen).
            patch('src.ui.contratos.QMessageBox.information'),
            patch.object(SaveContratosThread, 'start', lambda self_thread: None),
            patch.object(config.settings, 'CASOS_BAJA_CM_PATH', self.casos_baja_cm_path),
        ]
        for p in self._patches:
            p.start()

    def tearDown(self):
        for p in self._patches:
            p.stop()
        self.temp_dir.cleanup()

    def test_vin_ausente_genera_caso_pendiente_de_baja_cm(self):
        self.modulo.importar_excel_reporte()

        df_casos = leer_casos_baja_cm(self.casos_baja_cm_path)
        self.assertEqual(len(df_casos), 1, "Fallo: el VIN ausente con datos CM debía generar un caso PENDIENTE.")
        self.assertEqual(df_casos.iloc[0]['VIN'], 'VIN_AUSENTE')
        self.assertEqual(df_casos.iloc[0]['ESTADO_CASO'], ESTADO_CASO_PENDIENTE)

    def test_vin_presente_no_genera_caso_de_baja_cm(self):
        self.modulo.importar_excel_reporte()

        df_casos = leer_casos_baja_cm(self.casos_baja_cm_path)
        self.assertNotIn('VIN_PRESENTE', df_casos['VIN'].tolist())

    def test_vin_ausente_no_entra_en_los_cambios_a_aplicar(self):
        self.modulo.importar_excel_reporte()

        self.assertNotIn(
            "VIN_AUSENTE", self.modulo.import_changes,
            "Fallo: un VIN ausente del Excel mensual de CM no debe generar ningún cambio a aplicar sobre PruebasFM.",
        )
        self.assertIn("VIN_PRESENTE", self.modulo.import_changes)

    def test_vin_presente_se_actualiza_normalmente(self):
        self.modulo.importar_excel_reporte()

        cambios = self.modulo.import_changes["VIN_PRESENTE"]["cambios"]
        self.assertEqual(cambios["N CONTRATO 3PC"], "3PC-999")
        self.assertEqual(cambios["ESTADO DEL CM"], "Running")

    def test_datos_cm_del_vin_ausente_quedan_intactos_tras_inyectar_en_ram(self):
        self.modulo.importar_excel_reporte()
        self.modulo._on_import_finished(True)

        fila = self.modulo.data_manager.df_master.set_index("VIN").loc["VIN_AUSENTE"]
        self.assertEqual(fila["N CONTRATO 3PC"], "3PC-222")
        self.assertEqual(fila["FECHA INICIO CM"], "01-01-2025")
        self.assertEqual(fila["FECHA TERMINO CM"], "01-01-2026")
        self.assertEqual(fila["TIPO DE CM"], "Silver")
        self.assertEqual(fila["ESTADO DEL CM"], "Running")
        self.assertEqual(fila["PAGO CM"], 18000)

    def test_vin_ausente_nunca_se_elimina_de_pruebasfm(self):
        self.modulo.importar_excel_reporte()
        self.modulo._on_import_finished(True)

        vins_finales = self.modulo.data_manager.df_master["VIN"].tolist()
        self.assertIn("VIN_AUSENTE", vins_finales)


class TestContratosReglaTerminatedCm(unittest.TestCase):
    """Misma regla que TestReglaTerminatedCm en test_comercial.py, pero para
    la edición manual de ESTADO DEL CM en Contratos CM
    (ContratosModule.on_item_edited): al elegir 'Terminated' se limpian a
    '-' N CONTRATO 3PC, FECHA INICIO CM, FECHA TERMINO CM, TIPO DE CM y
    PAGO CM. No debe confundirse con la detección automática de
    cm_bajas_service.py (ausencia del VIN en el 3PC mensual): esta regla es
    manual y limpia los campos de inmediato al guardar."""

    @classmethod
    def setUpClass(cls):
        if QApplication.instance() is None:
            cls._app = QApplication([])

    def setUp(self):
        self.df_master = pd.DataFrame({
            "VIN": ["VIN_TERMINATED_TEST"],
            "CHASIS": ["C1"],
            "RAZÓN SOCIAL": ["Empresa Test"],
            "N CONTRATO 3PC": ["3PC-1"],
            "FECHA INICIO CM": ["01-01-2025"],
            "FECHA TERMINO CM": ["01-01-2026"],
            "TIPO DE CM": ["Gold"],
            "ESTADO DEL CM": ["Running"],
            "PAGO CM": [15000],
        })
        self.dm = DataManager()
        self.dm.df_master = self.df_master.copy()

        self.modulo = ContratosModule()
        self.modulo.data_manager.df_master = self.df_master.copy()
        self.modulo.load_data()

    def test_running_a_terminated_limpia_cinco_campos(self):
        self.modulo.on_item_edited(row=0, col_name='ESTADO DEL CM', new_value='Terminated')

        cambios = self.modulo.pending_changes['VIN_TERMINATED_TEST']['cambios']
        self.assertEqual(cambios.get('N CONTRATO 3PC'), '-')
        self.assertEqual(cambios.get('FECHA INICIO CM'), '-')
        self.assertEqual(cambios.get('FECHA TERMINO CM'), '-')
        self.assertEqual(cambios.get('TIPO DE CM'), '-')
        self.assertEqual(cambios.get('PAGO CM'), '-')
        self.assertEqual(cambios.get('ESTADO DEL CM'), 'Terminated')

        df_actual = self.modulo.model.get_df()
        fila = df_actual[df_actual['VIN'].astype(str).str.strip().str.upper() == 'VIN_TERMINATED_TEST'].iloc[0]
        self.assertEqual(fila['N CONTRATO 3PC'], '-')
        self.assertEqual(fila['FECHA INICIO CM'], '-')
        self.assertEqual(fila['FECHA TERMINO CM'], '-')
        self.assertEqual(fila['TIPO DE CM'], '-')
        self.assertEqual(fila['PAGO CM'], '-')

    def test_terminated_a_otro_valor_no_restaura_datos(self):
        self.modulo.on_item_edited(row=0, col_name='ESTADO DEL CM', new_value='Terminated')

        # El usuario vuelve a cambiar ESTADO DEL CM a otro valor: NO debe
        # existir regla inversa que reconstruya los 5 campos.
        self.modulo.on_item_edited(row=0, col_name='ESTADO DEL CM', new_value='Running')

        cambios = self.modulo.pending_changes['VIN_TERMINATED_TEST']['cambios']
        self.assertEqual(cambios.get('ESTADO DEL CM'), 'Running')
        self.assertEqual(cambios.get('N CONTRATO 3PC'), '-')
        self.assertEqual(cambios.get('FECHA INICIO CM'), '-')
        self.assertEqual(cambios.get('FECHA TERMINO CM'), '-')
        self.assertEqual(cambios.get('TIPO DE CM'), '-')
        self.assertEqual(cambios.get('PAGO CM'), '-')

    def test_columna_n_contrato_3pc_no_genera_columna_duplicada(self):
        """Regresión: contratos.py ya no renombra 'N CONTRATO 3PC' a
        'N° CONTRATO 3PC' solo para mostrarla -- usa el mismo nombre real en
        todo el sistema (ver constants.py, comercial.py, cm_bajas_service.py).
        Antes de esta corrección, si el dataset traía ambas columnas por
        separado, el rename generaba una columna duplicada y df_base[c]
        devolvía un DataFrame en vez de una Series, crasheando en
        .str.strip() (ver test_dirty_refresh.py, workaround ya retirado)."""
        self.assertIn('N CONTRATO 3PC', self.modulo.df_base.columns)
        self.assertNotIn('N° CONTRATO 3PC', self.modulo.df_base.columns)


class TestContratosReconciliacionEstadoExpiredAlGuardar(unittest.TestCase):
    """Integración de ESTADO DEL CM = Expired con Revisión CM disparada
    desde ContratosModule._on_save_finished (edición manual) -- solo cuando
    el cambio realmente se guarda, no al elegir el valor en el combo."""

    @classmethod
    def setUpClass(cls):
        if QApplication.instance() is None:
            cls._app = QApplication([])

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.casos_path = os.path.join(self.temp_dir.name, "CasosBajaCM_Test.csv")

        self.df_master = pd.DataFrame({
            "VIN": ["VIN_EXPIRED_TEST"],
            "CHASIS": ["C1"],
            "RAZÓN SOCIAL": ["Empresa Test"],
            "N CONTRATO 3PC": ["3PC-1"],
            "FECHA INICIO CM": ["01-01-2025"],
            "FECHA TERMINO CM": ["01-01-2026"],
            "TIPO DE CM": ["Gold"],
            "ESTADO DEL CM": ["Running"],
            "PAGO CM": [15000],
        })
        self.dm = DataManager()
        self.dm.df_master = self.df_master.copy()

        self.modulo = ContratosModule()
        self.modulo.data_manager.df_master = self.df_master.copy()
        self.modulo.load_data()

        self._patches = [
            patch('src.ui.contratos.QMessageBox.information'),
            patch.object(config.settings, 'CASOS_BAJA_CM_PATH', self.casos_path),
        ]
        for p in self._patches:
            p.start()

    def tearDown(self):
        for p in self._patches:
            p.stop()
        self.temp_dir.cleanup()

    def test_running_a_expired_guardado_genera_caso_en_revision_cm(self):
        self.modulo.on_item_edited(row=0, col_name='ESTADO DEL CM', new_value='Expired')
        self.modulo._on_save_finished(True)

        df_casos = leer_casos_baja_cm(self.casos_path)
        caso = df_casos[df_casos['VIN'] == 'VIN_EXPIRED_TEST']
        self.assertEqual(len(caso), 1)
        self.assertEqual(caso.iloc[0]['ESTADO_CASO'], ESTADO_CASO_PENDIENTE)
        self.assertEqual(caso.iloc[0]['ORIGEN_CASO'], ORIGEN_CASO_ESTADO_EXPIRED)

    def test_edicion_no_guardada_no_genera_caso(self):
        self.modulo.on_item_edited(row=0, col_name='ESTADO DEL CM', new_value='Expired')

        df_casos = leer_casos_baja_cm(self.casos_path)
        self.assertTrue(df_casos.empty, "Fallo: una edición sin guardar no debe generar un caso.")

    def test_expired_a_running_guardado_resuelve_el_caso(self):
        self.modulo.on_item_edited(row=0, col_name='ESTADO DEL CM', new_value='Expired')
        self.modulo._on_save_finished(True)

        self.modulo.on_item_edited(row=0, col_name='ESTADO DEL CM', new_value='Running')
        self.modulo._on_save_finished(True)

        df_casos = leer_casos_baja_cm(self.casos_path)
        caso = df_casos[df_casos['VIN'] == 'VIN_EXPIRED_TEST']
        self.assertEqual(len(caso), 1)
        self.assertEqual(caso.iloc[0]['ESTADO_CASO'], ESTADO_CASO_REAPARECIDO)


class TestContratosReconciliacionEstadoExpiredAlImportar(unittest.TestCase):
    """Integración de ESTADO DEL CM = Expired con Revisión CM disparada
    desde ContratosModule._on_import_finished (importación mensual 3PC)."""

    @classmethod
    def setUpClass(cls):
        if QApplication.instance() is None:
            cls._app = QApplication([])

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.import_path = os.path.join(self.temp_dir.name, "CM_Mensual.xlsx")

        df_import = pd.DataFrame({
            "VIN No": ["VIN_EXPIRA_ESTE_MES"],
            "Contract No": ["3PC-999"],
            "Start Date": ["2026-01-01"],
            "Expiring Date": ["2026-06-01"],
            "Contract Status": ["Expired"],
            "Contract Type": ["Gold"],
            "Net Amount": [20000],
        })
        df_import.to_excel(self.import_path, index=False)

        self.df_master_original = pd.DataFrame({
            "VIN": ["VIN_EXPIRA_ESTE_MES"],
            "CHASIS": ["C1"],
            "N CONTRATO 3PC": ["3PC-111"],
            "FECHA INICIO CM": ["01-01-2025"],
            "FECHA TERMINO CM": ["01-01-2026"],
            "TIPO DE CM": ["Silver"],
            "ESTADO DEL CM": ["Running"],
            "PAGO CM": [15000],
        })

        self.dm = DataManager()
        self.dm.df_master = self.df_master_original.copy()

        self.modulo = ContratosModule()
        self.modulo.data_manager.df_master = self.df_master_original.copy()
        self.modulo.load_data()

        self.casos_path = os.path.join(self.temp_dir.name, "CasosBajaCM_Test.csv")

        self._patches = [
            patch('src.ui.contratos.QFileDialog.getOpenFileName', return_value=(self.import_path, '')),
            # ResumenImportacionDialog (30-09-2026) reemplazó al QMessageBox.question
            # de confirmación de import masivo: se mockea su exec() en vez del
            # QMessageBox, para no abrir un diálogo modal real en los tests.
            patch('src.ui.contratos.ResumenImportacionDialog.exec', return_value=QDialog.DialogCode.Accepted),
            patch('src.ui.contratos.QMessageBox.information'),
            patch.object(SaveContratosThread, 'start', lambda self_thread: None),
            patch.object(config.settings, 'CASOS_BAJA_CM_PATH', self.casos_path),
        ]
        for p in self._patches:
            p.start()

    def tearDown(self):
        for p in self._patches:
            p.stop()
        self.temp_dir.cleanup()

    def test_importar_3pc_con_estado_expired_genera_caso_en_revision_cm(self):
        self.modulo.importar_excel_reporte()
        self.modulo._on_import_finished(True)

        df_casos = leer_casos_baja_cm(self.casos_path)
        caso = df_casos[df_casos['VIN'] == 'VIN_EXPIRA_ESTE_MES']
        self.assertEqual(len(caso), 1)
        self.assertEqual(caso.iloc[0]['ESTADO_CASO'], ESTADO_CASO_PENDIENTE)
        self.assertEqual(caso.iloc[0]['ORIGEN_CASO'], ORIGEN_CASO_ESTADO_EXPIRED)


class TestContratosEdicionManualFlujoReal(unittest.TestCase):
    """Ejercita el flujo real de guardado de Contratos (edición manual):
    ContratosModule.on_item_edited -> confirm_save -> SaveContratosThread ->
    ExcelWriter.update_cells() real, contra un .xlsx y un LogApp.csv
    temporales, sin mockear ExcelWriter ni la escritura real (solo
    SaveContratosThread.start, para ejecutar run() en el mismo hilo, y
    _crear_backup, para no tocar data/backups real -- mismo patrón usado en
    TestConfirmarBajaCmFlujoRealDeGuardado). Nace de la auditoría de
    2026-09: los diccionarios que arma contratos.py ya incluían 'VIN'
    correctamente, pero ningún test anterior ejecutaba update_cells() de
    verdad para comprobarlo (todos mockeaban SaveContratosThread.start o
    llamaban directo a _on_save_finished(True))."""

    COLUMNAS_EXCEL = ['VIN', 'N CONTRATO 3PC', 'FECHA INICIO CM', 'FECHA TERMINO CM', 'TIPO DE CM', 'ESTADO DEL CM', 'PAGO CM']

    @classmethod
    def setUpClass(cls):
        if QApplication.instance() is None:
            cls._app = QApplication([])

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.excel_path = os.path.join(self.temp_dir.name, "PruebasFM_Test.xlsx")
        self.log_path = os.path.join(self.temp_dir.name, "LogApp_Test.csv")
        self.casos_path = os.path.join(self.temp_dir.name, "CasosBajaCM_Test.csv")

        self.filas_iniciales = [
            ['VIN_A', '3PC-A', '01-01-2025', '01-01-2026', 'Silver', 'Running', 15000],
            ['VIN_B', '3PC-B', '01-01-2025', '01-01-2026', 'Gold', 'Running', 20000],
            ['VIN_C', '3PC-C', '01-01-2025', '01-01-2026', 'Basico', 'Running', 10000],
            ['VIN_D', '3PC-D', '01-01-2025', '01-01-2026', 'Gold', 'Running', 99999],
        ]
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(self.COLUMNAS_EXCEL)
        for fila in self.filas_iniciales:
            ws.append(fila)
        wb.save(self.excel_path)
        wb.close()

        self.df_master_original = pd.DataFrame(self.filas_iniciales, columns=self.COLUMNAS_EXCEL)

        self.dm = DataManager()
        self.dm.df_master = self.df_master_original.copy()

        self._patches = [
            patch.object(config.settings, 'CASOS_BAJA_CM_PATH', self.casos_path),
            patch.object(SaveContratosThread, 'start', lambda self_thread: self_thread.run()),
            patch('src.core.excel_writer.ExcelWriter._crear_backup'),
        ]
        for p in self._patches:
            p.start()

        self.modulo = ContratosModule()
        self.modulo.data_manager.df_master = self.df_master_original.copy()
        self.modulo.load_data()

        self.modulo.writer.filepath = self.excel_path
        self.modulo.writer.log_filepath = ""
        self.modulo.writer.log_csv_filepath = self.log_path

    def tearDown(self):
        for p in self._patches:
            p.stop()
        self.temp_dir.cleanup()

    def _fila_de(self, vin):
        df_actual = self.modulo.model.get_df().reset_index(drop=True)
        matches = df_actual.index[df_actual['VIN'].astype(str).str.strip().str.upper() == vin]
        self.assertGreater(len(matches), 0, f"VIN {vin} no aparece en la tabla del módulo.")
        return int(matches[0])

    def _leer_celda_excel(self, vin, columna):
        wb = openpyxl.load_workbook(self.excel_path)
        ws = wb.active
        col_idx = self.COLUMNAS_EXCEL.index(columna)
        vin_idx = self.COLUMNAS_EXCEL.index('VIN')
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[vin_idx] == vin:
                wb.close()
                return row[col_idx]
        wb.close()
        self.fail(f"VIN {vin} no encontrado en el Excel temporal.")

    def _leer_log(self):
        if not os.path.exists(self.log_path):
            return pd.DataFrame(columns=["FECHA Y HORA", "MÓDULO", "IDENTIFICADOR", "CAMPO MODIFICADO", "VALOR ANTERIOR", "NUEVO VALOR"])
        return pd.read_csv(self.log_path, dtype=str, encoding='utf-8-sig')

    @patch('PyQt6.QtWidgets.QMessageBox.information')
    @patch('PyQt6.QtWidgets.QMessageBox.exec')
    def test_un_vin_flujo_real_completo(self, mock_exec, mock_info):
        mock_exec.return_value = QMessageBox.StandardButton.Save

        self.modulo.on_item_edited(row=self._fila_de('VIN_A'), col_name='N CONTRATO 3PC', new_value='3PC-A-NEW')
        self.modulo.confirm_save()

        self.assertEqual(self._leer_celda_excel('VIN_A', 'N CONTRATO 3PC'), '3PC-A-NEW')

        df_log = self._leer_log()
        self.assertEqual(len(df_log), 1, "Fallo: debe generarse exactamente una línea de log para el único cambio real.")
        fila = df_log.iloc[0]
        self.assertEqual(fila['MÓDULO'], 'Contratos CM')
        self.assertEqual(fila['IDENTIFICADOR'], 'VIN_A')
        self.assertEqual(fila['CAMPO MODIFICADO'], 'N CONTRATO 3PC')
        self.assertEqual(fila['VALOR ANTERIOR'], '3PC-A')
        self.assertEqual(fila['NUEVO VALOR'], '3PC-A-NEW')
        self.assertTrue(str(fila['FECHA Y HORA']).strip(), "Fallo: debe registrarse una fecha/hora real.")

    @patch('PyQt6.QtWidgets.QMessageBox.information')
    @patch('PyQt6.QtWidgets.QMessageBox.exec')
    def test_multiples_vin_una_sola_llamada_actualiza_y_loguea_todos(self, mock_exec, mock_info):
        mock_exec.return_value = QMessageBox.StandardButton.Save

        instancias_thread = []
        thread_original_init = SaveContratosThread.__init__

        def _init_espia(self_thread, *args, **kwargs):
            thread_original_init(self_thread, *args, **kwargs)
            instancias_thread.append(self_thread)

        with patch.object(SaveContratosThread, '__init__', _init_espia):
            # VIN_A: un cambio real + un "cambio" a un valor ya existente (no debe generar log).
            self.modulo.on_item_edited(row=self._fila_de('VIN_A'), col_name='N CONTRATO 3PC', new_value='3PC-A-NEW')
            self.modulo.on_item_edited(row=self._fila_de('VIN_A'), col_name='TIPO DE CM', new_value='Silver')

            # VIN_B: dos cambios reales.
            self.modulo.on_item_edited(row=self._fila_de('VIN_B'), col_name='TIPO DE CM', new_value='Platinum')
            self.modulo.on_item_edited(row=self._fila_de('VIN_B'), col_name='FECHA TERMINO CM', new_value='15-08-2027')

            # VIN_C: Terminated dispara la limpieza de 5 columnas + el propio ESTADO DEL CM = 6 cambios reales.
            self.modulo.on_item_edited(row=self._fila_de('VIN_C'), col_name='ESTADO DEL CM', new_value='Terminated')

            # VIN_D: no se toca.
            self.modulo.confirm_save()

        self.assertEqual(len(instancias_thread), 1, "Fallo: debe crearse un único SaveContratosThread para todo el lote.")
        self.assertEqual(set(instancias_thread[0].changes.keys()), {'VIN_A', 'VIN_B', 'VIN_C'})

        self.assertEqual(self._leer_celda_excel('VIN_A', 'N CONTRATO 3PC'), '3PC-A-NEW')
        self.assertEqual(self._leer_celda_excel('VIN_A', 'TIPO DE CM'), 'Silver')

        self.assertEqual(self._leer_celda_excel('VIN_B', 'TIPO DE CM'), 'Platinum')
        self.assertEqual(self._leer_celda_excel('VIN_B', 'FECHA TERMINO CM'), '15-08-2027')

        self.assertEqual(self._leer_celda_excel('VIN_C', 'ESTADO DEL CM'), 'Terminated')
        for col in ('N CONTRATO 3PC', 'FECHA INICIO CM', 'FECHA TERMINO CM', 'TIPO DE CM', 'PAGO CM'):
            self.assertEqual(self._leer_celda_excel('VIN_C', col), '-', f"Fallo: {col} no quedó limpio en VIN_C.")

        # VIN_D no seleccionado: intacto.
        self.assertEqual(self._leer_celda_excel('VIN_D', 'ESTADO DEL CM'), 'Running')
        self.assertEqual(self._leer_celda_excel('VIN_D', 'PAGO CM'), 99999)

        df_log = self._leer_log()
        self.assertEqual(len(df_log), 9, "Fallo: 1 (VIN_A) + 2 (VIN_B) + 6 (VIN_C) = 9 líneas de log.")
        self.assertEqual(len(df_log[df_log['IDENTIFICADOR'] == 'VIN_A']), 1)
        self.assertNotIn('TIPO DE CM', df_log[df_log['IDENTIFICADOR'] == 'VIN_A']['CAMPO MODIFICADO'].tolist(),
                          "Fallo: TIPO DE CM de VIN_A no cambió de valor real y no debía generar log.")
        self.assertEqual(len(df_log[df_log['IDENTIFICADOR'] == 'VIN_B']), 2)
        self.assertEqual(len(df_log[df_log['IDENTIFICADOR'] == 'VIN_C']), 6)
        self.assertNotIn('VIN_D', df_log['IDENTIFICADOR'].tolist())

    @patch('PyQt6.QtWidgets.QMessageBox.information')
    @patch('PyQt6.QtWidgets.QMessageBox.exec')
    def test_estructura_pending_changes_debe_incluir_vin_para_que_excel_writer_encuentre_la_fila(self, mock_exec, mock_info):
        """Protección explícita del bug real encontrado en Confirmar Baja CM
        (comercial.py): si alguien modifica on_item_edited y el diccionario
        deja de incluir 'VIN' (queda solo {'cambios': {...}}),
        ExcelWriter.update_cells() ya no puede ubicar la fila (usa
        data.get('VIN'), no la clave del diccionario) y la escritura falla
        por completo. Este test ejecuta el código real de construcción del
        diccionario (on_item_edited) y lo lleva hasta update_cells() real:
        si la estructura se rompe, esta prueba falla porque el Excel no
        cambia y LogApp.csv queda vacío -- no se conforma con inspeccionar
        pending_changes en memoria."""
        mock_exec.return_value = QMessageBox.StandardButton.Save

        self.modulo.on_item_edited(row=self._fila_de('VIN_A'), col_name='N CONTRATO 3PC', new_value='3PC-A-NEW')

        self.assertIn('VIN', self.modulo.pending_changes['VIN_A'],
                       "Fallo: pending_changes ya no incluye la clave 'VIN' requerida por ExcelWriter.update_cells().")

        self.modulo.confirm_save()

        self.assertEqual(self._leer_celda_excel('VIN_A', 'N CONTRATO 3PC'), '3PC-A-NEW',
                          "Fallo: el Excel real no cambió -- si esto falla, revisa que on_item_edited siga "
                          "incluyendo 'VIN' en pending_changes[vin].")
        self.assertEqual(len(self._leer_log()), 1)


class TestContratosImportacionMasivaFlujoReal(unittest.TestCase):
    """Ejercita el flujo real de la importación mensual 3PC
    (ContratosModule.importar_excel_reporte) hasta ExcelWriter.update_cells()
    real, contra un .xlsx y un LogApp.csv temporales -- sin mockear
    ExcelWriter ni la escritura real. Complementa
    TestContratosNoLimpiaAutomaticamenteVinAusente (que sí mockea
    SaveContratosThread.start y solo valida el DataFrame en memoria):
    aquí se comprueba que 'Contratos CM (Masivo)' realmente escribe en
    PruebasFM.xlsx y queda registrado en LogApp.csv."""

    COLUMNAS_EXCEL = ['VIN', 'N CONTRATO 3PC', 'FECHA INICIO CM', 'FECHA TERMINO CM', 'TIPO DE CM', 'ESTADO DEL CM', 'PAGO CM']

    @classmethod
    def setUpClass(cls):
        if QApplication.instance() is None:
            cls._app = QApplication([])

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.excel_path = os.path.join(self.temp_dir.name, "PruebasFM_Test.xlsx")
        self.log_path = os.path.join(self.temp_dir.name, "LogApp_Test.csv")
        self.casos_path = os.path.join(self.temp_dir.name, "CasosBajaCM_Test.csv")
        self.import_path = os.path.join(self.temp_dir.name, "CM_Mensual.xlsx")

        filas_iniciales = [
            # ESTADO DEL CM se mantiene igual en el import ('Running'): no debe generar log para esa columna.
            ['VIN_IMPORT_1', '3PC-OLD-1', '01-01-2020', '01-01-2021', 'Silver', 'Running', 1000],
            # ESTADO DEL CM sí cambia en el import ('Running' -> 'Expired'): las 6 columnas cambian.
            ['VIN_IMPORT_2', '3PC-OLD-2', '01-01-2020', '01-01-2021', 'Bronze', 'Running', 2000],
            ['VIN_NO_EN_IMPORT', '3PC-OLD-3', '01-01-2020', '01-01-2021', 'Gold', 'Running', 3000],
        ]
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(self.COLUMNAS_EXCEL)
        for fila in filas_iniciales:
            ws.append(fila)
        wb.save(self.excel_path)
        wb.close()

        self.df_master_original = pd.DataFrame(filas_iniciales, columns=self.COLUMNAS_EXCEL)

        df_import = pd.DataFrame({
            "VIN No": ["VIN_IMPORT_1", "VIN_IMPORT_2"],
            "Contract No": ["3PC-NEW-1", "3PC-NEW-2"],
            "Start Date": ["2026-01-01", "2026-02-01"],
            "Expiring Date": ["2027-01-01", "2027-02-01"],
            "Contract Status": ["Running", "Expired"],
            "Contract Type": ["Gold", "Platinum"],
            "Net Amount": [5000, 8000],
        })
        df_import.to_excel(self.import_path, index=False)

        self.dm = DataManager()
        self.dm.df_master = self.df_master_original.copy()

        self._patches = [
            patch.object(config.settings, 'CASOS_BAJA_CM_PATH', self.casos_path),
            patch.object(SaveContratosThread, 'start', lambda self_thread: self_thread.run()),
            patch('src.core.excel_writer.ExcelWriter._crear_backup'),
            patch('src.ui.contratos.QFileDialog.getOpenFileName', return_value=(self.import_path, '')),
            # ResumenImportacionDialog (30-09-2026) reemplazó al QMessageBox.question
            # de confirmación de import masivo: se mockea su exec() en vez del
            # QMessageBox, para no abrir un diálogo modal real en los tests.
            patch('src.ui.contratos.ResumenImportacionDialog.exec', return_value=QDialog.DialogCode.Accepted),
            patch('src.ui.contratos.QMessageBox.information'),
        ]
        for p in self._patches:
            p.start()

        self.modulo = ContratosModule()
        self.modulo.data_manager.df_master = self.df_master_original.copy()
        self.modulo.load_data()

        self.modulo.writer.filepath = self.excel_path
        self.modulo.writer.log_filepath = ""
        self.modulo.writer.log_csv_filepath = self.log_path

    def tearDown(self):
        for p in self._patches:
            p.stop()
        self.temp_dir.cleanup()

    def _leer_celda_excel(self, vin, columna):
        wb = openpyxl.load_workbook(self.excel_path)
        ws = wb.active
        col_idx = self.COLUMNAS_EXCEL.index(columna)
        vin_idx = self.COLUMNAS_EXCEL.index('VIN')
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[vin_idx] == vin:
                wb.close()
                return row[col_idx]
        wb.close()
        self.fail(f"VIN {vin} no encontrado en el Excel temporal.")

    def _leer_log(self):
        if not os.path.exists(self.log_path):
            return pd.DataFrame(columns=["FECHA Y HORA", "MÓDULO", "IDENTIFICADOR", "CAMPO MODIFICADO", "VALOR ANTERIOR", "NUEVO VALOR"])
        return pd.read_csv(self.log_path, dtype=str, encoding='utf-8-sig')

    def test_importacion_masiva_flujo_real_escribe_excel_y_loguea(self):
        self.modulo.importar_excel_reporte()

        self.assertEqual(self._leer_celda_excel('VIN_IMPORT_1', 'N CONTRATO 3PC'), '3PC-NEW-1')
        self.assertEqual(self._leer_celda_excel('VIN_IMPORT_1', 'FECHA INICIO CM'), '01-01-2026')
        self.assertEqual(self._leer_celda_excel('VIN_IMPORT_1', 'ESTADO DEL CM'), 'Running')

        self.assertEqual(self._leer_celda_excel('VIN_IMPORT_2', 'N CONTRATO 3PC'), '3PC-NEW-2')
        self.assertEqual(self._leer_celda_excel('VIN_IMPORT_2', 'ESTADO DEL CM'), 'Expired')

        # VIN ausente del import: la regla de negocio dice que no se toca (no confirma baja automática).
        self.assertEqual(self._leer_celda_excel('VIN_NO_EN_IMPORT', 'N CONTRATO 3PC'), '3PC-OLD-3')
        self.assertEqual(self._leer_celda_excel('VIN_NO_EN_IMPORT', 'ESTADO DEL CM'), 'Running')

        df_log = self._leer_log()
        self.assertEqual(len(df_log), 11, "Fallo: 5 (VIN_IMPORT_1, sin ESTADO DEL CM) + 6 (VIN_IMPORT_2) = 11 líneas.")
        self.assertNotIn('VIN_NO_EN_IMPORT', df_log['IDENTIFICADOR'].tolist(),
                          "Fallo: un VIN ausente del import no debe generar ningún registro de log.")

        self.assertTrue((df_log['MÓDULO'] == 'Contratos CM (Masivo)').all(),
                         "Fallo: el módulo de log de la importación masiva debe ser exactamente 'Contratos CM (Masivo)'.")

        log_vin1 = df_log[df_log['IDENTIFICADOR'] == 'VIN_IMPORT_1']
        self.assertEqual(len(log_vin1), 5)
        self.assertNotIn('ESTADO DEL CM', log_vin1['CAMPO MODIFICADO'].tolist(),
                          "Fallo: ESTADO DEL CM de VIN_IMPORT_1 no cambió de valor real y no debía generar log.")

        log_vin2 = df_log[df_log['IDENTIFICADOR'] == 'VIN_IMPORT_2']
        self.assertEqual(len(log_vin2), 6)

        fila_contrato_1 = log_vin1[log_vin1['CAMPO MODIFICADO'] == 'N CONTRATO 3PC'].iloc[0]
        self.assertEqual(fila_contrato_1['VALOR ANTERIOR'], '3PC-OLD-1')
        self.assertEqual(fila_contrato_1['NUEVO VALOR'], '3PC-NEW-1')


if __name__ == '__main__':
    unittest.main()

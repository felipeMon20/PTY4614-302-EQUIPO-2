import os
import tempfile
import unittest

import docx

from src.services.facturacion_manual_template import (
    construir_plantilla_base, generar_plantilla_base, PLACEHOLDER_PENDIENTE,
    FILAS_RESUMEN_FACTURACION, ENCABEZADOS_TABLA_FLOTA, TOKENS_FILA_FLOTA,
)


def _texto_completo(documento):
    return "\n".join(p.text for p in documento.paragraphs)


class TestConstruirPlantillaBase(unittest.TestCase):
    """Plantilla base de Facturación Manual (fase 'Crear plantilla base', ver
    Notas del proyecto.txt, sección 7): solo estructura visual, sin cálculos ni conexión
    a 'Facturación Manual - BBDD.xlsx'. Todo campo financiero debe quedar
    literalmente marcado como PENDIENTE CONTRAPARTE COMERCIAL."""

    def test_incluye_textos_fijos_reales(self):
        documento = construir_plantilla_base()
        texto = _texto_completo(documento)
        self.assertIn("Estimado(a) cliente:", texto)
        self.assertIn("El impuesto al valor agregado (IVA)", texto)
        self.assertIn("Ana Ejemplo Demo", texto)

    def test_incluye_placeholders_de_datos_confirmados(self):
        documento = construir_plantilla_base()
        texto = _texto_completo(documento)
        self.assertIn("{{RAZON_SOCIAL}}", texto)
        self.assertIn("{{RUT}}", texto)
        self.assertIn("{{FECHA_EMISION}}", texto)

    def test_tabla_resumen_facturacion_marca_pendientes(self):
        documento = construir_plantilla_base()
        tabla_resumen = documento.tables[0]
        filas = [[c.text for c in row.cells] for row in tabla_resumen.rows]
        self.assertEqual(filas, [list(f) for f in FILAS_RESUMEN_FACTURACION])

        campos_financieros = [f for e, f in FILAS_RESUMEN_FACTURACION if e != "Unidades conectadas"]
        self.assertTrue(all(v == PLACEHOLDER_PENDIENTE for v in campos_financieros))

    def test_tabla_flota_tiene_columnas_confirmadas(self):
        documento = construir_plantilla_base()
        tabla_flota = documento.tables[1]
        encabezados = [c.text for c in tabla_flota.rows[0].cells]
        self.assertEqual(encabezados, ENCABEZADOS_TABLA_FLOTA)

        fila_tokens = [c.text for c in tabla_flota.rows[1].cells]
        self.assertEqual(fila_tokens, TOKENS_FILA_FLOTA)

    def test_no_incluye_ningun_monto_o_numero_inventado(self):
        # Ningun campo financiero debe traer un valor numerico o de ejemplo:
        # todo lo que no sea texto fijo o token {{...}} debe ser el
        # placeholder literal.
        documento = construir_plantilla_base()
        tabla_resumen = documento.tables[0]
        for row in tabla_resumen.rows[1:]:  # fila 0 es "Unidades conectadas" (token, no dato)
            valor = row.cells[1].text
            self.assertEqual(valor, PLACEHOLDER_PENDIENTE)


class TestGenerarPlantillaBase(unittest.TestCase):
    def test_genera_archivo_docx_valido_en_disco(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            destino = os.path.join(tmp_dir, "sub", "plantilla_base.docx")
            resultado = generar_plantilla_base(destino)

            self.assertEqual(resultado, destino)
            self.assertTrue(os.path.exists(destino))

            documento_leido = docx.Document(destino)
            self.assertIn("Ana Ejemplo Demo", _texto_completo(documento_leido))
            self.assertEqual(len(documento_leido.tables), 2)


if __name__ == '__main__':
    unittest.main()

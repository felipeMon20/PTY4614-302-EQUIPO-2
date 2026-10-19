"""Presentación de los entregables de calidad (duoc/calidad/*.xlsx).

Se ejecuta como último paso de quality_sync.sync(): ordena y traduce lo que el
auditor acaba de escribir para que cualquier persona pueda abrir los archivos y
entender los resultados sin conocer el proyecto. No cambia las posiciones de
columna ni los marcadores que quality_sync lee en la siguiente corrida
('Clase.método — archivo.py' en el Reporte de Ejecución y 'Test automatizado:
archivo.py::Clase.' en Casos de Prueba); es idempotente.

Agrega una hoja "Resumen" al inicio de cada libro. La hoja de datos queda como
hoja activa, porque quality_sync y sus pruebas leen `wb.active`.
"""
from __future__ import annotations

import math
import os
import re
from collections import OrderedDict

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

AZUL = "1F4E78"
AZUL_CLARO = "DDEBF7"
GRIS = "F2F2F2"
VERDE_FONDO, VERDE_TEXTO = "E2F0D9", "375623"
ROJO_FONDO, ROJO_TEXTO = "F8D7DA", "842029"
AMBAR_FONDO, AMBAR_TEXTO = "FFF2CC", "7F6000"

_LINEA = Side(style="thin", color="BFBFBF")
BORDE = Border(left=_LINEA, right=_LINEA, top=_LINEA, bottom=_LINEA)

HOJA_REPORTE = "Reporte de Ejecución"
HOJA_CASOS = "Casos de Prueba"
HOJA_RESUMEN = "Resumen"

# Nombre legible de cada archivo de pruebas (grupo = nombre sin 'test_' ni '.py')
NOMBRES_MODULO = {
    "comercial": "Comercial (pantalla)",
    "comercial_services": "Comercial (cálculos)",
    "comercial_agregar": "Formulario de unidad",
    "operaciones_services": "Operaciones (cálculos)",
    "provisiones": "Provisiones (pantalla)",
    "provisiones_services": "Provisiones (cálculos)",
    "contratos": "Contratos CM (pantalla)",
    "cm_bajas_service": "Revisión de bajas CM",
    "contactos": "Contactos (pantalla)",
    "contactos_service": "Contactos (cálculos)",
    "contacto_correo_service": "Correo por contacto",
    "mail_dialog_correo_faltante": "Correos: correo faltante",
    "mail_registro": "Correos: historial de envíos",
    "mail_service": "Correos: destinatarios",
    "mail_sync_service": "Correos: cruce maestro",
    "general": "Vista General",
    "main_window": "Ventana principal",
    "configuracion": "Configuración",
    "settings": "Configuración (rutas)",
    "historial": "Historial y papelera",
    "backups": "Respaldos por escritura",
    "backup_manager": "Respaldos programados",
    "excel_writer": "Escritura de Excel",
    "excel_writer_csv_backend": "Escritura de CSV",
    "dataframe_formatter": "Formato y cálculo de fechas",
    "data_manager": "Gestor de datos",
    "vin_service": "Modelo según VIN",
    "theme_manager": "Tema visual",
    "table_sorting": "Tablas: orden y menús",
    "dirty_refresh": "Actualización de pantallas",
    "facturacion_manual_service": "Facturación Manual (cálculos)",
    "facturacion_manual_template": "Facturación Manual (plantilla)",
    "file_lock_monitor": "Archivo bloqueado",
    "app_lock_coordinator": "Bloqueo de la aplicación",
    "fase2_csv_integration": "Integración CSV (fase 2)",
    "fase3_1_csv_fixes": "Correcciones CSV (fase 3)",
    "quality_sync": "Sincronización de calidad",
}

ETIQUETAS_TIPO = {
    "Excel": "Excel/CSV", "interfaz/UI": "Interfaz", "unitaria": "Unitaria", "configuración": "Configuración",
    "integración": "Integración", "validación": "Validación", "backups": "Respaldos",
    "persistencia": "Persistencia", "manejo de errores": "Manejo de errores",
    "sincronización": "Sincronización", "regresión": "Regresión",
}


def nombre_modulo(grupo):
    grupo = str(grupo or "").strip()
    if grupo in NOMBRES_MODULO:
        return NOMBRES_MODULO[grupo]
    if grupo in NOMBRES_MODULO.values() or not grupo:
        return grupo
    return grupo.replace("_", " ").strip().capitalize()


def _tipos(texto):
    partes = [p.strip() for p in str(texto or "").split(",") if p.strip()]
    return ", ".join(ETIQUETAS_TIPO.get(p, p) for p in partes) if partes else "Sin clasificar"


def _fecha_legible(texto):
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})[T ](\d{2}:\d{2}:\d{2})", str(texto or ""))
    return f"{m.group(3)}-{m.group(2)}-{m.group(1)} {m.group(4)}" if m else str(texto or "")


def _filas_estimadas(texto, ancho):
    """Líneas que ocupará un texto con ajuste de línea en una columna de `ancho` caracteres."""
    if texto is None:
        return 1
    caracteres = max(1, int(ancho * 1.15))
    return sum(max(1, math.ceil(len(linea) / caracteres)) for linea in str(texto).split("\n"))


def _encabezado(celda):
    celda.font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    celda.fill = PatternFill("solid", fgColor=AZUL)
    celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    celda.border = BORDE


def _estado_estilo(celda, estado):
    estado = str(estado or "").strip().upper()
    fondo, texto = {"PASS": (VERDE_FONDO, VERDE_TEXTO), "FAIL": (ROJO_FONDO, ROJO_TEXTO),
                    "ERROR": (ROJO_FONDO, ROJO_TEXTO), "SKIPPED": (AMBAR_FONDO, AMBAR_TEXTO),
                    "OBSOLETA": ("D9D9D9", "595959")}.get(estado, (GRIS, "000000"))
    celda.fill = PatternFill("solid", fgColor=fondo)
    celda.font = Font(name="Calibri", bold=True, color=texto, size=10)
    celda.alignment = Alignment(horizontal="center", vertical="top")
    celda.border = BORDE


def _celda_texto(celda, negrita=False, centrado=False, size=10):
    celda.font = Font(name="Calibri", bold=negrita, size=size)
    celda.alignment = Alignment(horizontal="center" if centrado else "left", vertical="top", wrap_text=True)
    celda.border = BORDE


def _hoja_datos(wb, nombre):
    """La hoja de datos es la activa (o la que ya se llama así); nunca la 'Resumen'."""
    if nombre in wb.sheetnames:
        return wb[nombre]
    for ws in wb.worksheets:
        if ws.title != HOJA_RESUMEN:
            ws.title = nombre
            return ws
    raise RuntimeError("No se encontró la hoja de datos")


def _nueva_hoja_resumen(wb):
    if HOJA_RESUMEN in wb.sheetnames:
        del wb[HOJA_RESUMEN]
    ws = wb.create_sheet(HOJA_RESUMEN, 0)
    ws.sheet_view.showGridLines = False
    return ws


def _titulo(ws, fila, texto, subtitulo=None):
    ws.cell(fila, 1, texto).font = Font(name="Calibri", bold=True, size=18, color=AZUL)
    if subtitulo:
        ws.cell(fila + 1, 1, subtitulo).font = Font(name="Calibri", italic=True, size=10, color="595959")


def _seccion(ws, fila, texto, columnas=7):
    for c in range(1, columnas + 1):
        ws.cell(fila, c).fill = PatternFill("solid", fgColor=AZUL)
    ws.cell(fila, 1, texto).font = Font(name="Calibri", bold=True, size=12, color="FFFFFF")
    ws.cell(fila, 1).alignment = Alignment(vertical="center")
    ws.row_dimensions[fila].height = 20


def _parrafo(ws, fila, texto, columnas=7, negrita=False, alto=None):
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=columnas)
    c = ws.cell(fila, 1, texto)
    c.font = Font(name="Calibri", size=10, bold=negrita)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[fila].height = alto or max(15, 14 * _filas_estimadas(texto, 110))


def _tabla(ws, fila, encabezados, filas, anchos_centrados=()):
    for i, h in enumerate(encabezados, start=1):
        _encabezado(ws.cell(fila, i, h))
    for k, valores in enumerate(filas, start=1):
        for i, v in enumerate(valores, start=1):
            c = ws.cell(fila + k, i, v)
            _celda_texto(c, centrado=(i in anchos_centrados))
            if k % 2 == 0:
                c.fill = PatternFill("solid", fgColor=GRIS)
    return fila + len(filas) + 1


def _por_modulo(tests):
    resultado = OrderedDict()
    for t in sorted(tests, key=lambda x: (nombre_modulo(x.get("module_group")), x["file"])):
        r = resultado.setdefault(nombre_modulo(t.get("module_group")), {"total": 0, "PASS": 0, "FAIL": 0, "ERROR": 0, "SKIPPED": 0})
        r["total"] += 1
        r[t["status"]] = r.get(t["status"], 0) + 1
    return resultado


# ---------------------------------------------------------------------
# Reporte de Ejecución de Pruebas.xlsx
# ---------------------------------------------------------------------

ENCABEZADOS_REPORTE = [
    "Fecha y hora de ejecución", "N° de caso", "Módulo", "Tipo de prueba",
    "Resultado (PASS = aprobó · FAIL = falló · SKIPPED = omitida)", "Duración (ms)", "Defecto asociado",
    "Prueba automática (grupo.prueba — archivo)",
]
ANCHOS_REPORTE = {"A": 20, "B": 11, "C": 28, "D": 34, "E": 26, "F": 13, "G": 15, "H": 78}


def _formatear_reporte(ws, tests):
    from quality_sync import _identity_from_observaciones
    vigentes = {(x["file"], x["class"], x["method"]) for x in tests}
    ws.title = HOJA_REPORTE
    ultima = 1
    for fila in range(2, ws.max_row + 1):
        cp = ws.cell(fila, 2).value
        if isinstance(cp, str) and re.match(r"^CP-\d+$", cp):
            ultima = fila
    for i, h in enumerate(ENCABEZADOS_REPORTE, start=1):
        ws.cell(1, i, h)
        _encabezado(ws.cell(1, i))
    ws.row_dimensions[1].height = 52

    for fila in range(2, ultima + 1):
        ws.cell(fila, 1, _fecha_legible(ws.cell(fila, 1).value))
        ws.cell(fila, 3, nombre_modulo(ws.cell(fila, 3).value))
        ws.cell(fila, 4, _tipos(ws.cell(fila, 4).value))
        valor = ws.cell(fila, 6).value
        try:
            ws.cell(fila, 6, float(valor))
        except (TypeError, ValueError):
            pass
        identidad = _identity_from_observaciones(ws.cell(fila, 8).value)
        obsoleta = identidad is not None and identidad not in vigentes
        if obsoleta:
            ws.cell(fila, 5, "OBSOLETA")
            ws.cell(fila, 7, "Prueba eliminada o renombrada")
        if not ws.cell(fila, 7).value:
            ws.cell(fila, 7, "Sin defecto")
        for col in range(1, 9):
            c = ws.cell(fila, col)
            _celda_texto(c, centrado=col in (1, 2, 7))
        ws.cell(fila, 6).number_format = "#,##0.000"
        ws.cell(fila, 6).alignment = Alignment(horizontal="right", vertical="top")
        _estado_estilo(ws.cell(fila, 5), ws.cell(fila, 5).value)
        if obsoleta:
            for col in (1, 2, 3, 4, 6, 7, 8):
                ws.cell(fila, col).font = Font(name="Calibri", size=10, color="7F7F7F", italic=True)
        lineas = max(_filas_estimadas(ws.cell(fila, 8).value, ANCHOS_REPORTE["H"]),
                     _filas_estimadas(ws.cell(fila, 4).value, ANCHOS_REPORTE["D"]),
                     _filas_estimadas(ws.cell(fila, 3).value, ANCHOS_REPORTE["C"]))
        ws.row_dimensions[fila].height = max(18, 13.5 * lineas + 4)

    # Resumen al pie: se mantiene (el auditor lo actualiza por etiqueta) pero legible
    for fila in range(ultima + 1, ws.max_row + 1):
        etiqueta = ws.cell(fila, 1).value
        valor = ws.cell(fila, 2).value
        if etiqueta is None and valor is None:
            continue
        es_titulo = isinstance(etiqueta, str) and etiqueta.upper().startswith("RESUMEN") and valor is None
        for col in range(1, 9):
            ws.cell(fila, col).border = Border()
        if es_titulo:
            for col in range(1, 9):
                ws.cell(fila, col).fill = PatternFill("solid", fgColor=AZUL)
            c = ws.cell(fila, 1)
            c.font = Font(name="Calibri", bold=True, size=13, color="FFFFFF")
            c.alignment = Alignment(vertical="center")
            ws.row_dimensions[fila].height = 24
            continue
        ws.cell(fila, 1).font = Font(name="Calibri", bold=True, size=10)
        ws.cell(fila, 1).fill = PatternFill("solid", fgColor=AZUL_CLARO)
        ws.cell(fila, 1).alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(fila, 2).font = Font(name="Calibri", size=10)
        ws.cell(fila, 2).alignment = Alignment(vertical="top", wrap_text=False, horizontal="left")
        ws.row_dimensions[fila].height = max(16, 13.5 * _filas_estimadas(etiqueta, 20) + 3)

    for letra, ancho in ANCHOS_REPORTE.items():
        ws.column_dimensions[letra].width = ancho
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{ultima}"
    ws.sheet_properties.tabColor = AZUL
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)
    return ultima


def _contar_obsoletas(ws):
    return sum(1 for fila in range(2, ws.max_row + 1) if ws.cell(fila, 5).value == "OBSOLETA")


def _resumen_reporte(wb, data, obsoletas=0):
    ws = _nueva_hoja_resumen(wb)
    resumen, meta, metricas = data["summary"], data["metadata"], data["metrics"]
    resumen = {"tests_executed": 0, "pass": 0, "fail": 0, "error": 0, "skipped": 0, **resumen}
    metricas = {"pass_rate_pct": 0, "duration_total_ms": 0, **metricas}
    tests = data["tests"]
    fallas = resumen["fail"] + resumen["error"]
    for letra, ancho in {"A": 34, "B": 16, "C": 16, "D": 16, "E": 16, "F": 16, "G": 26}.items():
        ws.column_dimensions[letra].width = ancho

    _titulo(ws, 1, "Resultado de las pruebas automáticas",
            f"Gestor SSCC App · ejecución del {_fecha_legible(meta.get('generated_at'))} · generado por el auditor automático")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A4:G4")
    veredicto = ws.cell(4, 1, "RESULTADO: SIN FALLAS" if not fallas else f"RESULTADO: HAY {fallas} PRUEBA(S) CON FALLA")
    veredicto.font = Font(name="Calibri", bold=True, size=16, color=VERDE_TEXTO if not fallas else ROJO_TEXTO)
    veredicto.fill = PatternFill("solid", fgColor=VERDE_FONDO if not fallas else ROJO_FONDO)
    veredicto.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[4].height = 34
    _parrafo(ws, 5, (
        "Todas las pruebas que se pudieron ejecutar fueron aprobadas." if not fallas else
        "Hay pruebas que fallaron o terminaron con error; se listan más abajo y deben revisarse antes de entregar la versión."
    ) + (f" Además, {resumen['skipped']} prueba(s) quedaron omitidas: no se ejecutaron, por lo que no cuentan como aprobadas ni como falladas."
         if resumen.get("skipped") else ""), alto=32)

    _seccion(ws, 7, "Cifras principales")
    cifras = [
        ("Pruebas ejecutadas", resumen["tests_executed"]),
        ("Aprobadas (PASS)", resumen["pass"]),
        ("Falladas (FAIL)", resumen["fail"]),
        ("Con error (ERROR)", resumen["error"]),
        ("Omitidas (SKIPPED)", resumen["skipped"]),
        ("Porcentaje aprobado", f"{metricas['pass_rate_pct']}%"),
        ("Duración total", f"{metricas['duration_total_ms'] / 1000:.0f} segundos"),
        ("Archivos de pruebas", meta.get("files_discovered_for_execution", "-")),
        ("Grupos de pruebas (clases)", meta.get("classes_discovered", "-")),
    ]
    fila = 8
    for etiqueta, valor in cifras:
        a = ws.cell(fila, 1, etiqueta)
        a.font = Font(name="Calibri", bold=True, size=10)
        a.fill = PatternFill("solid", fgColor=AZUL_CLARO)
        a.border = BORDE
        b = ws.cell(fila, 2, valor)
        b.font = Font(name="Calibri", size=11, bold=True)
        b.alignment = Alignment(horizontal="center")
        b.border = BORDE
        if etiqueta.startswith("Falladas") or etiqueta.startswith("Con error"):
            if valor:
                b.fill = PatternFill("solid", fgColor=ROJO_FONDO)
        fila += 1

    fila += 1
    _seccion(ws, fila, "Cómo leer este archivo")
    guia = [
        "• Hoja «Reporte de Ejecución»: una fila por cada prueba automática, con su resultado y cuánto demoró.",
        "• Resultado PASS = la prueba se aprobó (el sistema hizo lo esperado). FAIL = la prueba falló. ERROR = la prueba no pudo terminar por un problema. SKIPPED = la prueba se omitió (no se ejecutó).",
        "• Con el filtro de la columna «Resultado» se puede ver solo lo que falló u omitió.",
        "• El archivo «Casos de Prueba» explica qué verifica cada grupo de pruebas, en lenguaje de negocio.",
    ]
    if obsoletas:
        guia.append(f"• Las {obsoletas} filas grises con resultado OBSOLETA son pruebas antiguas que ya no existen (se eliminaron o cambiaron de nombre). "
                    "Se conservan solo para mantener el historial de números de caso y no cuentan en las cifras de arriba.")
    fila += 1
    for linea in guia:
        _parrafo(ws, fila, linea)
        fila += 1

    fila += 1
    _seccion(ws, fila, "Resultado por módulo")
    modulos = _por_modulo(tests)
    filas = []
    for nombre, r in modulos.items():
        malo = r["FAIL"] + r["ERROR"]
        filas.append([nombre, r["total"], r["PASS"], r["FAIL"], r["ERROR"], r["SKIPPED"],
                      "Con fallas" if malo else ("Parcial (hay omitidas)" if r["SKIPPED"] else "Todo aprobado")])
    inicio = fila + 1
    fila = _tabla(ws, inicio, ["Módulo", "Pruebas", "Aprobadas", "Falladas", "Con error", "Omitidas", "Estado"], filas,
                  anchos_centrados=(2, 3, 4, 5, 6, 7))
    for k in range(len(filas)):
        c = ws.cell(inicio + 1 + k, 7)
        texto = c.value
        fondo, color = (ROJO_FONDO, ROJO_TEXTO) if texto == "Con fallas" else (
            (AMBAR_FONDO, AMBAR_TEXTO) if texto.startswith("Parcial") else (VERDE_FONDO, VERDE_TEXTO))
        c.fill = PatternFill("solid", fgColor=fondo)
        c.font = Font(name="Calibri", bold=True, size=10, color=color)

    problemas = [t for t in tests if t["status"] in ("FAIL", "ERROR")]
    omitidas = [t for t in tests if t["status"] == "SKIPPED"]
    for titulo, lista, motivo in (
        ("Pruebas con falla o error (requieren revisión)", problemas, "error"),
        ("Pruebas omitidas (no se ejecutaron)", omitidas, "skip_reason"),
    ):
        if not lista:
            continue
        fila += 1
        _seccion(ws, fila, titulo)
        filas = [[nombre_modulo(t.get("module_group")), f"{t['class']}.{t['method']}",
                  str(t.get(motivo) or "Sin detalle")[:300].replace("\n", " ")] for t in lista]
        for k, valores in enumerate(filas, start=1):
            ws.cell(fila + k, 1, valores[0])
            ws.cell(fila + k, 2, valores[1])
            ws.merge_cells(start_row=fila + k, start_column=2, end_row=fila + k, end_column=4)
            ws.cell(fila + k, 5, valores[2])
            ws.merge_cells(start_row=fila + k, start_column=5, end_row=fila + k, end_column=7)
            for c in range(1, 8):
                _celda_texto(ws.cell(fila + k, c))
            ws.row_dimensions[fila + k].height = max(30, 13.5 * _filas_estimadas(valores[2], 60) + 4)
        fila += len(filas) + 1
        if motivo == "skip_reason":
            _parrafo(ws, fila, "Una prueba omitida no es una falla: depende de datos de prueba que no estaban disponibles en ese momento. "
                               "Conviene corregir los datos de prueba para que se ejecuten.", alto=30)
            fila += 1
    ws.sheet_properties.tabColor = "70AD47" if not fallas else "C00000"
    return ws


def _formatear_reporte_libro(path, data):
    wb = openpyxl.load_workbook(path)
    datos = _hoja_datos(wb, HOJA_REPORTE)
    _formatear_reporte(datos, data["tests"])
    _resumen_reporte(wb, data, _contar_obsoletas(datos))
    wb.active = wb.worksheets.index(datos)
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws is datos
    wb.save(path)


# ---------------------------------------------------------------------
# Casos de Prueba.xlsx
# ---------------------------------------------------------------------

ENCABEZADOS_CASOS = [
    "N° de caso (del Reporte de Ejecución)", "Módulo o componente", "Qué se prueba", "Condiciones previas",
    "Pasos", "Resultado esperado", "Resultado obtenido", "Estado (PASS = aprobó · FAIL = falló)",
    "Comentarios y evidencia",
]
ANCHOS_CASOS = {"A": 24, "B": 34, "C": 60, "D": 42, "E": 52, "F": 46, "G": 40, "H": 16, "I": 48}
_PREFIJOS_PRECOND = ("Entorno sandbox", "Entorno de pruebas aislado", "Ninguna dependencia", "No necesita pantallas")


def _limpiar_descripcion(texto, metodos):
    texto = re.sub(r"\s*\n\s*", " ", str(texto or "")).strip()
    if not texto or texto == "NO DOCUMENTADO":
        nombres = ["verifica " + m.replace("_", " ") for m in metodos[:3]]
        texto = "; ".join(nombres) + ("…" if len(metodos) > 3 else "") if nombres else "Sin descripción"
        return texto[0].upper() + texto[1:]
    return texto if len(texto) <= 420 else texto[:417].rstrip() + "…"


def _formatear_casos(ws, tests, project_root):
    from quality_sync import (PRECOND_SANDBOX, _precondiciones_para, _componente_desde_imports,
                              _derive_component_label, _identity_from_comentarios)
    ws.title = HOJA_CASOS
    por_clase = {}
    for t in tests:
        por_clase.setdefault((t["file"], t["class"]), []).append(t)
    for i, h in enumerate(ENCABEZADOS_CASOS, start=1):
        ws.cell(1, i, h)
        _encabezado(ws.cell(1, i))
    ws.row_dimensions[1].height = 48

    ultima = 1
    for fila in range(2, ws.max_row + 1):
        if ws.cell(fila, 1).value is None:
            continue
        ultima = fila
        ident = _identity_from_comentarios(ws.cell(fila, 9).value)
        grupo = por_clase.get(ident, []) if ident else []
        etiqueta_vieja = ws.cell(fila, 2).value
        if isinstance(etiqueta_vieja, str):
            m = re.match(r"^(?P<n>.+?) \((?P<t>UI|Servicio|Core)\) - (?P<r>.+)$", etiqueta_vieja)
            if m:
                tipo_es = {"UI": "pantalla", "Servicio": "reglas y cálculos", "Core": "núcleo"}[m.group("t")]
                ws.cell(fila, 2, f"{m.group('n')} — {tipo_es} ({m.group('r')})")
        if grupo:
            etiqueta = ws.cell(fila, 2).value
            if isinstance(etiqueta, str) and ("NO DETERMINABLE" in etiqueta or "revisar manualmente" in etiqueta):
                nueva, _ = _derive_component_label(grupo[0].get("module_group", ""), project_root)
                nueva = _componente_desde_imports(grupo[0].get("relevant_imports"), project_root) or nueva
                ws.cell(fila, 2, nueva)
            ws.cell(fila, 3, _limpiar_descripcion(ws.cell(fila, 3).value, [t["method"] for t in sorted(grupo, key=lambda x: x["method"])]))
            precond = ws.cell(fila, 4).value
            if isinstance(precond, str) and precond.startswith(_PREFIJOS_PRECOND):
                ws.cell(fila, 4, _precondiciones_para(grupo[0].get("relevant_imports"), grupo[0].get("module_group", ""), grupo[0]["file"]))
        if ident and not grupo:
            ws.cell(fila, 8, "OBSOLETA")
        for col in range(1, 10):
            _celda_texto(ws.cell(fila, col), centrado=col == 1)
        _estado_estilo(ws.cell(fila, 8), ws.cell(fila, 8).value)
        if ident and not grupo:
            for col in (1, 2, 3, 4, 5, 6, 7, 9):
                ws.cell(fila, col).font = Font(name="Calibri", size=10, color="7F7F7F", italic=True)
        lineas = max(_filas_estimadas(ws.cell(fila, col).value, ancho)
                     for col, ancho in ((2, 34), (3, 60), (4, 42), (5, 52), (6, 46), (7, 40), (9, 48)))
        ws.row_dimensions[fila].height = min(300, max(30, 13.5 * lineas + 4))

    for letra, ancho in ANCHOS_CASOS.items():
        ws.column_dimensions[letra].width = ancho
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:I{ultima}"
    ws.sheet_properties.tabColor = AZUL
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)
    return ultima


def _resumen_casos(wb, data, filas_casos):
    ws = _nueva_hoja_resumen(wb)
    for letra, ancho in {"A": 34, "B": 16, "C": 16, "D": 16, "E": 22}.items():
        ws.column_dimensions[letra].width = ancho
    _titulo(ws, 1, "Casos de prueba",
            f"Gestor SSCC App · actualizado el {_fecha_legible(data['metadata'].get('generated_at'))} · generado por el auditor automático")
    ws.row_dimensions[1].height = 28
    _parrafo(ws, 4, (
        "Este archivo describe QUÉ se prueba. Cada fila es un grupo de pruebas automáticas (una «clase»): explica qué se verifica, "
        "cómo se ejecuta, qué resultado se espera y qué resultado se obtuvo en la última ejecución. "
        "El resultado de cada prueba individual está en «Reporte de Ejecución de Pruebas»."), columnas=5, alto=62)

    _seccion(ws, 6, "Cómo leer la hoja «Casos de Prueba»", columnas=5)
    guia = [
        "• N° de caso: números de las pruebas del grupo en el Reporte de Ejecución (CP-040_CP-045 significa del 40 al 45).",
        "• Qué se prueba: la regla de negocio o función que se comprueba, en palabras simples.",
        "• Estado: PASS = todas las pruebas del grupo se aprobaron; FAIL = alguna falló; OBSOLETA (en gris) = el grupo ya no existe en el código.",
        "• Con el filtro de cada columna se puede buscar por módulo o por estado.",
    ]
    fila = 7
    for linea in guia:
        _parrafo(ws, fila, linea, columnas=5)
        fila += 1

    por_modulo = OrderedDict()
    for t in data["tests"]:
        r = por_modulo.setdefault(nombre_modulo(t.get("module_group")), {"clases": set(), "pruebas": 0, "malas": 0})
        r["clases"].add((t["file"], t["class"]))
        r["pruebas"] += 1
        r["malas"] += t["status"] in ("FAIL", "ERROR")
    fila += 1
    _seccion(ws, fila, f"Cobertura por módulo ({filas_casos - 1} casos en total)", columnas=5)
    filas = [[n, len(r["clases"]), r["pruebas"], r["malas"], "Con fallas" if r["malas"] else "Todo aprobado"]
             for n, r in sorted(por_modulo.items())]
    inicio = fila + 1
    _tabla(ws, inicio, ["Módulo", "Casos (grupos)", "Pruebas", "Con falla", "Estado"], filas, anchos_centrados=(2, 3, 4, 5))
    for k in range(len(filas)):
        c = ws.cell(inicio + 1 + k, 5)
        bueno = c.value == "Todo aprobado"
        c.fill = PatternFill("solid", fgColor=VERDE_FONDO if bueno else ROJO_FONDO)
        c.font = Font(name="Calibri", bold=True, size=10, color=VERDE_TEXTO if bueno else ROJO_TEXTO)
    ws.sheet_properties.tabColor = "70AD47"


def _formatear_casos_libro(path, data, project_root):
    wb = openpyxl.load_workbook(path)
    datos = _hoja_datos(wb, HOJA_CASOS)
    ultima = _formatear_casos(datos, data["tests"], project_root)
    _resumen_casos(wb, data, ultima)
    wb.active = wb.worksheets.index(datos)
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws is datos
    wb.save(path)


def aplicar(calidad_dir, data, project_root):
    """Da formato y agrega la hoja Resumen a los dos libros. Devuelve un dict con el estado."""
    reporte = os.path.join(calidad_dir, "Reporte de Ejecución de Pruebas.xlsx")
    casos = os.path.join(calidad_dir, "Casos de Prueba.xlsx")
    _formatear_reporte_libro(reporte, data)
    _formatear_casos_libro(casos, data, project_root)
    return {"reporte": reporte, "casos": casos}

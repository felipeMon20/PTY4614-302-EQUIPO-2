import os
import csv
import unittest
import tempfile
import openpyxl
import pandas as pd
from datetime import datetime
from unittest.mock import patch
import config.settings
from src.core.excel_writer import ExcelWriter
from src.core.data_manager import DataManager
from src.services import comercial_services

def _derivar_csv_path_desde_xlsx(log_xlsx_path):
    """Helper de test: deriva la ruta LogApp.csv esperada a partir de una
    ruta LogApp.xlsx (mismo criterio que config/settings.py). No depende de
    codigo de produccion: get_log_csv_path() de excel_writer.py se elimino
    por no tener ningun uso productivo (ver auditoria de codigo muerto)."""
    return os.path.splitext(log_xlsx_path)[0] + ".csv"


class TestExcelWriter(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.writer = ExcelWriter()
        self.dm = DataManager()

    def tearDown(self):
        self.temp_dir.cleanup()

    @patch('src.core.excel_writer.ExcelWriter._crear_backup')
    def test_escenario_4_hoja_prov_inexistente(self, mock_backup):
        excel_path = os.path.join(self.temp_dir.name, "test_no_prov.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws.append(["COL1", "COL2"])
        ws.append(["val1", "val2"])
        wb.save(excel_path)
        wb.close()

        cambios = {
            "fila1": {
                "keys": {"COL1": "val1"},
                "cambios": {"COL2": "nuevo_val"}
            }
        }

        res = self.writer.update_cells(cambios, module_name="Test", custom_filepath=excel_path, sheet_name="Prov")

        self.assertFalse(res, "Fallo: update_cells debería retornar False si la hoja no existe.")

        wb_check = openpyxl.load_workbook(excel_path)
        self.assertNotIn("Prov", wb_check.sheetnames, "Fallo: update_cells no debería crear la hoja 'Prov' si no existe.")
        wb_check.close()

        self.assertFalse(self.dm.is_app_saving, "Fallo: is_app_saving debería quedar en False tras el error.")

    @patch('src.core.excel_writer.ExcelWriter._crear_backup')
    def test_escenario_5_ano_mes_como_datetime_real(self, mock_backup):
        excel_path = os.path.join(self.temp_dir.name, "test_datetime.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Prov"
        ws.append(["AÑO - MES", "VIN", "DYNAFLEET", "Plan"])
        
        fecha_real = datetime(2026, 8, 1)
        ws.append([fecha_real, "YV2TEST999", 100000, "BASICO"])
        wb.save(excel_path)
        wb.close()

        cambios = {
            "item1": {
                "keys": {"VIN": "YV2TEST999", "AÑO - MES": "2026 - Agosto"},
                "cambios": {"DYNAFLEET": 500000, "Plan": "FULL"}
            }
        }

        res = self.writer.update_cells(cambios, module_name="Test", custom_filepath=excel_path, sheet_name="Prov")
        self.assertTrue(res, "Fallo: ExcelWriter debería encontrar la fila cuando AÑO - MES es datetime y la key es string.")

        wb_check = openpyxl.load_workbook(excel_path)
        ws_check = wb_check["Prov"]
        self.assertEqual(ws_check.cell(row=2, column=3).value, 500000)
        self.assertEqual(ws_check.cell(row=2, column=4).value, "FULL")
        self.assertIsInstance(ws_check.cell(row=2, column=1).value, datetime)
        wb_check.close()

        self.assertFalse(self.dm.is_app_saving)

    # =========================================================
    # PUNTO A: LogApp.csv (LOG_CSV_PATH) como unico historial oficial,
    # desacoplado de LOG_EXCEL_PATH salvo para la migracion legada.
    # =========================================================
    def test_log_changes_crea_logapp_csv_directo_sin_xlsx(self):
        """Sin LogApp.xlsx previo, _log_changes escribe directo en LOG_CSV_PATH."""
        self.writer.log_filepath = os.path.join(self.temp_dir.name, "LogApp_Test.xlsx")  # no existe: no hay nada que migrar
        self.writer.log_csv_filepath = os.path.join(self.temp_dir.name, "LogApp_Test.csv")
        csv_path = self.writer._get_log_csv_path()
        self.assertEqual(csv_path, self.writer.log_csv_filepath)

        entradas = [["2026-08-31 10:00:00", "Comercial", "YV2TEST0001", "MODELO", "FM", "FMH"]]
        self.writer._log_changes(entradas)

        self.assertFalse(os.path.exists(self.writer.log_filepath),
                          "Fallo: no debe crearse ningun .xlsx; el historial vive solo en LogApp.csv.")
        self.assertTrue(os.path.exists(csv_path), "Fallo: LogApp.csv deberia haberse creado.")

        df = pd.read_csv(csv_path, dtype=str)
        self.assertEqual(len(df), 1)
        self.assertEqual(list(df.columns),
                          ["FECHA Y HORA", "MÓDULO", "IDENTIFICADOR", "CAMPO MODIFICADO", "VALOR ANTERIOR", "NUEVO VALOR"])
        self.assertEqual(df.iloc[0]["IDENTIFICADOR"], "YV2TEST0001")
        self.assertEqual(df.iloc[0]["NUEVO VALOR"], "FMH")

    def test_log_changes_sin_log_csv_path_configurado_no_falla(self):
        """Si LOG_CSV_PATH esta vacio (config sin normalizar todavia), _log_changes
        no debe lanzar excepcion ni intentar escribir en una ruta vacia."""
        self.writer.log_filepath = ""
        self.writer.log_csv_filepath = ""
        try:
            self.writer._log_changes([["2026-08-31 10:00:00", "Comercial", "YV2X", "X", "1", "2"]])
        except Exception as e:
            self.fail(f"_log_changes no debe lanzar excepcion con LOG_CSV_PATH vacio: {e}")

    def test_log_changes_multiples_llamadas_acumulan_sin_duplicar_encabezado(self):
        self.writer.log_filepath = os.path.join(self.temp_dir.name, "LogApp_multi.xlsx")
        self.writer.log_csv_filepath = os.path.join(self.temp_dir.name, "LogApp_multi.csv")
        csv_path = self.writer._get_log_csv_path()

        for i in range(3):
            self.writer._log_changes([[f"2026-08-31 10:0{i}:00", "Comercial", f"YV2TEST000{i}", "MODELO", "FM", "FMH"]])

        with open(csv_path, encoding='utf-8-sig') as f:
            lineas = f.readlines()
        self.assertEqual(len(lineas), 4, "Fallo: debe haber 1 encabezado + 3 filas, sin encabezados repetidos.")

        df = pd.read_csv(csv_path, dtype=str)
        self.assertEqual(len(df), 3)

    def test_log_changes_preserva_caracteres_utf8(self):
        self.writer.log_filepath = os.path.join(self.temp_dir.name, "LogApp_utf8.xlsx")
        self.writer.log_csv_filepath = os.path.join(self.temp_dir.name, "LogApp_utf8.csv")
        csv_path = self.writer._get_log_csv_path()

        self.writer._log_changes([["2026-08-31 10:00:00", "Comercial", "YV2ÑOÑO", "RAZÓN SOCIAL", "Ñandú S.A.", "Peña & Cía."]])

        df = pd.read_csv(csv_path, dtype=str, encoding='utf-8-sig')
        self.assertEqual(df.iloc[0]["IDENTIFICADOR"], "YV2ÑOÑO")
        self.assertEqual(df.iloc[0]["VALOR ANTERIOR"], "Ñandú S.A.")
        self.assertEqual(df.iloc[0]["NUEVO VALOR"], "Peña & Cía.")

    # =========================================================
    # leer_logs_actual(): relectura de LogApp.csv en el propio hilo de
    # guardado (SaveExcelThread), para poder propagar el resultado a
    # DataManager.set_logs_df() sin esperar a un load_data_async() completo
    # -- causa raiz del historial desactualizado hasta presionar "Recargar"
    # (15-09-2026).
    # =========================================================
    def test_leer_logs_actual_devuelve_las_entradas_recien_escritas(self):
        self.writer.log_csv_filepath = os.path.join(self.temp_dir.name, "LogApp_leer.csv")
        self.writer._log_changes([["2026-09-15 10:00:00", "Comercial", "YV2TEST0001", "MODELO", "FM", "FH"]])

        df = self.writer.leer_logs_actual()
        self.assertEqual(len(df), 1)
        self.assertEqual(df.iloc[0]["IDENTIFICADOR"], "YV2TEST0001")
        self.assertEqual(df.iloc[0]["VALOR ANTERIOR"], "FM")
        self.assertEqual(df.iloc[0]["NUEVO VALOR"], "FH")

    def test_leer_logs_actual_sin_archivo_devuelve_vacio_no_falla(self):
        self.writer.log_csv_filepath = os.path.join(self.temp_dir.name, "LogApp_no_existe.csv")
        df = self.writer.leer_logs_actual()
        self.assertTrue(df.empty)

    def test_leer_logs_actual_sin_ruta_configurada_devuelve_vacio_no_falla(self):
        self.writer.log_csv_filepath = ""
        df = self.writer.leer_logs_actual()
        self.assertTrue(df.empty)

    def test_leer_logs_actual_refleja_llamadas_sucesivas(self):
        """Dos guardados seguidos (dos _log_changes) -> dos relecturas
        sucesivas de leer_logs_actual() deben reflejar cada una el estado
        del CSV en ese momento, sin quedar pegadas al primer resultado."""
        self.writer.log_csv_filepath = os.path.join(self.temp_dir.name, "LogApp_sucesivas.csv")
        self.writer._log_changes([["2026-09-15 10:00:00", "Comercial", "YV2A", "MODELO", "FM", "FH"]])
        primera = self.writer.leer_logs_actual()
        self.assertEqual(len(primera), 1)

        self.writer._log_changes([["2026-09-15 10:05:00", "Comercial", "YV2B", "ESTADO COMERCIAL", "Revisar", "Vigente"]])
        segunda = self.writer.leer_logs_actual()
        self.assertEqual(len(segunda), 2)
        self.assertIn("YV2B", segunda["IDENTIFICADOR"].tolist())

    # ---------------------------------------------------------
    # Migracion unica: LogApp.xlsx + LogApp_incremental.csv (LOG_EXCEL_PATH)
    # -> LogApp.csv (LOG_CSV_PATH, destino explicito)
    # ---------------------------------------------------------
    def _crear_legacy_xlsx(self, path, header, filas):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Historial"
        ws.append(header)
        for fila in filas:
            ws.append(fila)
        wb.save(path)
        wb.close()

    def _crear_legacy_csv(self, path, header, filas):
        with open(path, 'w', newline='', encoding='utf-8-sig') as f:
            w = csv.writer(f)
            w.writerow(header)
            for fila in filas:
                w.writerow(fila)

    def test_migracion_combina_xlsx_y_csv_preservando_orden_sin_deduplicar(self):
        from src.core.excel_writer import migrar_historial_legacy

        legacy_xlsx = os.path.join(self.temp_dir.name, "LogApp_mig.xlsx")
        legacy_csv = os.path.join(self.temp_dir.name, "LogApp_mig_incremental.csv")
        csv_path = _derivar_csv_path_desde_xlsx(legacy_xlsx)  # destino: LOG_CSV_PATH normalizado desde el legado
        header = ["FECHA Y HORA", "MÓDULO", "VIN", "CAMPO MODIFICADO", "VALOR ANTERIOR", "NUEVO VALOR"]

        self._crear_legacy_xlsx(legacy_xlsx, header, [
            ["2026-08-01 09:00:00", "Comercial", "YV2A", "ESTADO", "Vigente", "Vencido"],
            ["2026-08-02 09:00:00", "Comercial", "YV2A", "ESTADO", "Vencido", "Vigente"],  # "duplicado" legitimo: no debe perderse
        ])
        self._crear_legacy_csv(legacy_csv, header, [
            ["2026-08-31 10:00:00", "Comercial", "YV2B", "MODELO", "FM", "FMH"],
        ])

        migrado = migrar_historial_legacy(legacy_xlsx, csv_path)
        self.assertTrue(migrado)
        self.assertTrue(os.path.exists(csv_path))

        # Los archivos legado no se tocan ni se borran
        self.assertTrue(os.path.exists(legacy_xlsx))
        wb_check = openpyxl.load_workbook(legacy_xlsx)
        self.assertEqual(wb_check.active.max_row, 3, "Fallo: LogApp.xlsx no debe modificarse por la migracion.")
        wb_check.close()
        self.assertTrue(os.path.exists(legacy_csv))

        df = pd.read_csv(csv_path, dtype=str)
        self.assertEqual(len(df), 3, "Fallo: deben migrarse las 2 filas del xlsx + 1 del csv, sin deduplicar.")
        self.assertEqual(list(df.columns)[2], "VIN", "Fallo: debe heredar el encabezado real del xlsx legado.")
        # Orden preservado: primero xlsx (mas antiguo), despues csv (mas reciente)
        self.assertEqual(df.iloc[0]["VIN"], "YV2A")
        self.assertEqual(df.iloc[0]["NUEVO VALOR"], "Vencido")
        self.assertEqual(df.iloc[1]["VIN"], "YV2A")
        self.assertEqual(df.iloc[1]["NUEVO VALOR"], "Vigente")
        self.assertEqual(df.iloc[2]["VIN"], "YV2B")

    def test_migracion_es_idempotente_no_se_repite(self):
        from src.core.excel_writer import migrar_historial_legacy

        legacy_xlsx = os.path.join(self.temp_dir.name, "LogApp_once.xlsx")
        csv_path = _derivar_csv_path_desde_xlsx(legacy_xlsx)
        header = ["FECHA Y HORA", "MÓDULO", "IDENTIFICADOR", "CAMPO MODIFICADO", "VALOR ANTERIOR", "NUEVO VALOR"]
        self._crear_legacy_xlsx(legacy_xlsx, header, [["2026-08-01 09:00:00", "Comercial", "YV2A", "X", "1", "2"]])

        primera = migrar_historial_legacy(legacy_xlsx, csv_path)
        self.assertTrue(primera)

        with open(csv_path, 'a', newline='', encoding='utf-8-sig') as f:
            csv.writer(f).writerow(["2026-09-01 10:00:00", "Comercial", "YV2C", "Y", "3", "4"])

        segunda = migrar_historial_legacy(legacy_xlsx, csv_path)
        self.assertFalse(segunda, "Fallo: la migracion no debe repetirse si LogApp.csv ya existe.")

        df = pd.read_csv(csv_path, dtype=str)
        self.assertEqual(len(df), 2, "Fallo: no debio volver a migrar ni perder la fila agregada despues.")

    def test_migracion_sin_archivos_legado_no_hace_nada(self):
        from src.core.excel_writer import migrar_historial_legacy

        ruta_inexistente = os.path.join(self.temp_dir.name, "LogApp_fresh.xlsx")
        csv_path = _derivar_csv_path_desde_xlsx(ruta_inexistente)
        resultado = migrar_historial_legacy(ruta_inexistente, csv_path)
        self.assertFalse(resultado)
        self.assertFalse(os.path.exists(csv_path))

    def test_migracion_sin_log_csv_path_no_hace_nada(self):
        from src.core.excel_writer import migrar_historial_legacy

        legacy_xlsx = os.path.join(self.temp_dir.name, "LogApp_sin_destino.xlsx")
        self._crear_legacy_xlsx(legacy_xlsx, ["FECHA Y HORA", "MÓDULO", "IDENTIFICADOR", "CAMPO MODIFICADO", "VALOR ANTERIOR", "NUEVO VALOR"],
                                 [["2026-08-01 09:00:00", "Comercial", "YV2A", "X", "1", "2"]])
        resultado = migrar_historial_legacy(legacy_xlsx, "")
        self.assertFalse(resultado, "Fallo: sin LOG_CSV_PATH configurado no hay a donde migrar.")

    def test_data_manager_lee_solo_logapp_csv_tras_migracion(self):
        """MainDataLoaderThread ya no debe depender de LogApp.xlsx + CSV incremental
        por separado: migra (si corresponde) y lee unicamente LOG_CSV_PATH."""
        from src.core.data_manager import MainDataLoaderThread

        main_excel_path = os.path.join(self.temp_dir.name, "PruebasFM_test.xlsx")
        df_base = pd.DataFrame({"VIN": ["YV2A", "YV2B"], "CHASIS": ["C1", "C2"]})
        df_base.to_excel(main_excel_path, index=False, sheet_name='BBDD (2)')

        legacy_xlsx = os.path.join(self.temp_dir.name, "LogApp_dm.xlsx")
        csv_path = _derivar_csv_path_desde_xlsx(legacy_xlsx)
        header = ["FECHA Y HORA", "MÓDULO", "IDENTIFICADOR", "CAMPO MODIFICADO", "VALOR ANTERIOR", "NUEVO VALOR"]
        self._crear_legacy_xlsx(legacy_xlsx, header, [["2026-08-01 09:00:00", "Comercial", "YV2VIEJO", "ESTADO COMERCIAL", "Vigente", "Vencido"]])

        with patch.object(config.settings, 'LOG_EXCEL_PATH', legacy_xlsx), \
             patch.object(config.settings, 'LOG_CSV_PATH', csv_path), \
             patch.object(config.settings, 'DELETED_LOG_EXCEL_PATH', os.path.join(self.temp_dir.name, "no_existe.xlsx")):

            self.writer.log_filepath = legacy_xlsx
            self.writer.log_csv_filepath = csv_path
            self.writer._log_changes([["2026-08-31 10:00:00", "Comercial", "YV2NUEVO", "MODELO", "FM", "FMH"]])

            resultado = {}
            def capturar(payload):
                # MainDataLoaderThread emite ahora 5 elementos: el 5to
                # (deleted_sheets) es un dict {modulo: DataFrame} con todas las
                # hojas de RegistroEliminado.xlsx (ver DataManager). No es lo
                # que este test valida, pero el desempaquetado debe reflejar
                # la forma real de la señal.
                df, df_logs, df_deleted, load_time, deleted_sheets = payload
                resultado['df_logs'] = df_logs

            thread = MainDataLoaderThread(main_excel_path)
            thread.data_loaded_signal.connect(capturar)
            thread.run()  # ejecucion sincronica directa (sin QThread real) para el test

        df_logs = resultado['df_logs']
        self.assertEqual(len(df_logs), 2, "Fallo: deben aparecer el registro migrado (ex-xlsx) y el nuevo, ambos desde LogApp.csv.")
        identificadores = df_logs["IDENTIFICADOR"].astype(str).tolist()
        self.assertIn("YV2VIEJO", identificadores)
        self.assertIn("YV2NUEVO", identificadores)

        # El xlsx legado sigue existiendo intacto (la migracion no lo borra)
        self.assertTrue(os.path.exists(legacy_xlsx))

    def test_main_data_loader_lee_todas_las_hojas_de_registro_eliminado(self):
        """MainDataLoaderThread debe exponer TODAS las hojas de
        RegistroEliminado.xlsx (una por modulo de origen), no solo la primera
        (comportamiento por defecto de pd.read_excel sin sheet_name). Antes de
        esta correccion, las eliminaciones de cualquier modulo que no fuera la
        primera hoja del archivo quedaban invisibles para la app."""
        from src.core.data_manager import MainDataLoaderThread

        main_excel_path = os.path.join(self.temp_dir.name, "PruebasFM_hojas.xlsx")
        df_base = pd.DataFrame({"VIN": ["YV2A"], "CHASIS": ["C1"]})
        df_base.to_excel(main_excel_path, index=False, sheet_name='BBDD (2)')

        deleted_path = os.path.join(self.temp_dir.name, "RegistroEliminado_hojas.xlsx")
        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "Comercial"
        ws1.append(["FECHA Y HORA DE ELIMINACIÓN", "MÓDULO", "VIN", "MODELO"])
        ws1.append(["2026-08-01 09:00:00", "Comercial", "YV2AAA", "FM"])
        ws2 = wb.create_sheet("Contactos")
        ws2.append(["FECHA Y HORA DE ELIMINACIÓN", "MÓDULO", "ID", "Nombre"])
        ws2.append(["2026-08-02 10:00:00", "Contactos", "458", "Ana"])
        wb.save(deleted_path)
        wb.close()

        with patch.object(config.settings, 'DELETED_LOG_EXCEL_PATH', deleted_path):
            resultado = {}
            def capturar(payload):
                df, df_logs, df_deleted, load_time, deleted_sheets = payload
                resultado['deleted_sheets'] = deleted_sheets

            thread = MainDataLoaderThread(main_excel_path)
            thread.data_loaded_signal.connect(capturar)
            thread.run()

        deleted_sheets = resultado['deleted_sheets']
        self.assertIn("Comercial", deleted_sheets)
        self.assertIn("Contactos", deleted_sheets)
        self.assertEqual(list(deleted_sheets["Comercial"].columns), ["FECHA Y HORA DE ELIMINACIÓN", "MÓDULO", "VIN", "MODELO"])
        self.assertEqual(list(deleted_sheets["Contactos"].columns), ["FECHA Y HORA DE ELIMINACIÓN", "MÓDULO", "ID", "Nombre"])
        self.assertEqual(len(deleted_sheets["Comercial"]), 1)
        self.assertEqual(len(deleted_sheets["Contactos"]), 1)
        self.assertEqual(deleted_sheets["Comercial"]["VIN"].iloc[0], "YV2AAA")


class TestUpdateCellsConfirmarBajaCm(unittest.TestCase):
    """Ejercita ExcelWriter.update_cells() REAL (sin mocks de SaveExcelThread
    ni de la propia ExcelWriter) con el mismo diccionario que arma
    comercial.py::_on_accion_cm_confirmar_baja_clicked para "Confirmar baja"
    de Gestionar CM.

    Nace de un bug real encontrado: `valid[vin]` se armaba como
    {'cambios': {...}}, sin la clave 'VIN'. update_cells() identifica la
    fila a editar con `data.get('VIN')` (no con la clave del diccionario,
    ver excel_writer.py), asi que sin esa clave nunca encontraba la fila,
    retornaba False de inmediato y jamas llegaba a escribir el Excel ni a
    llamar _log_changes(). Los tests de comercial.py no lo detectaban porque
    siempre reemplazaban SaveExcelThread por un doble que nunca invocaba
    update_cells() de verdad -- por eso pasaban 742 tests con la escritura
    real rota. Estas pruebas llaman a update_cells() directamente, sin
    ningun mock de SaveExcelThread ni de ExcelWriter (solo _crear_backup,
    para no tocar data/backups real -- mismo patron que el resto de esta
    clase de test, ver test_escenario_4/5 arriba)."""

    COLUMNAS = ['VIN', 'ESTADO DEL CM', 'ESTADO COMERCIAL', 'ETAPA GESTIÓN', 'N CONTRATO 3PC', 'PAGO CM']

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.writer = ExcelWriter()
        self.writer.log_filepath = ""
        self.writer.log_csv_filepath = os.path.join(self.temp_dir.name, "LogApp_confirmar_baja.csv")
        self.excel_path = os.path.join(self.temp_dir.name, "PruebasFM_confirmar_baja.xlsx")

    def tearDown(self):
        self.temp_dir.cleanup()

    def _crear_excel(self, filas):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(self.COLUMNAS)
        for fila in filas:
            ws.append(fila)
        wb.save(self.excel_path)
        wb.close()

    def _valid_para(self, vin):
        """Misma construccion que comercial.py::_on_accion_cm_confirmar_baja_clicked:
        {'VIN': vin, 'cambios': {...}}. El estado del CM pasa a Terminated y el
        contrato a BAJA sin tocar los datos del CM."""
        cambios = {'ESTADO DEL CM': 'Terminated', 'ESTADO COMERCIAL': 'BAJA', 'ETAPA GESTIÓN': 'OPERACIONES'}
        return {'VIN': vin, 'cambios': cambios}

    def _leer_celda(self, vin, columna):
        wb = openpyxl.load_workbook(self.excel_path)
        ws = wb.active
        col_idx = self.COLUMNAS.index(columna)
        vin_idx = self.COLUMNAS.index('VIN')
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[vin_idx] == vin:
                wb.close()
                return row[col_idx]
        wb.close()
        self.fail(f"VIN {vin} no encontrado en el Excel temporal.")

    def _leer_log(self):
        if not os.path.exists(self.writer.log_csv_filepath):
            return pd.DataFrame(columns=["FECHA Y HORA", "MÓDULO", "IDENTIFICADOR", "CAMPO MODIFICADO", "VALOR ANTERIOR", "NUEVO VALOR"])
        return pd.read_csv(self.writer.log_csv_filepath, dtype=str, encoding='utf-8-sig')

    @patch('src.core.excel_writer.ExcelWriter._crear_backup')
    def test_un_vin_escribe_excel_y_genera_log_por_cada_columna_que_cambia(self, mock_backup):
        self._crear_excel([
            ['VIN1', 'Expired', 'Vigente', 'COMERCIAL', '3PC-1', '15000'],
        ])
        valid = {'VIN1': self._valid_para('VIN1')}

        resultado = self.writer.update_cells(valid, module_name="Comercial", custom_filepath=self.excel_path)
        self.assertTrue(resultado, "Fallo: update_cells debe devolver True con el dict correcto (con 'VIN').")

        self.assertEqual(self._leer_celda('VIN1', 'ESTADO DEL CM'), 'Terminated')
        self.assertEqual(self._leer_celda('VIN1', 'ESTADO COMERCIAL'), 'BAJA')
        self.assertEqual(self._leer_celda('VIN1', 'ETAPA GESTIÓN'), 'OPERACIONES')
        self.assertEqual(self._leer_celda('VIN1', 'N CONTRATO 3PC'), '3PC-1', "Fallo: los datos del CM deben conservarse.")
        self.assertEqual(self._leer_celda('VIN1', 'PAGO CM'), '15000')

        df_log = self._leer_log()
        self.assertEqual(len(df_log), 3, "Fallo: deben generarse 3 líneas (ESTADO DEL CM, ESTADO COMERCIAL y ETAPA GESTIÓN).")

        fila_estado = df_log[df_log['CAMPO MODIFICADO'] == 'ESTADO DEL CM'].iloc[0]
        self.assertEqual(fila_estado['IDENTIFICADOR'], 'VIN1')
        self.assertEqual(fila_estado['MÓDULO'], 'Comercial')
        self.assertEqual(fila_estado['VALOR ANTERIOR'], 'Expired')
        self.assertEqual(fila_estado['NUEVO VALOR'], 'Terminated')

    @patch('src.core.excel_writer.ExcelWriter._crear_backup')
    def test_columna_ya_en_su_valor_final_no_genera_log_innecesario(self, mock_backup):
        """Si el contrato ya estaba en BAJA antes de confirmar la baja,
        update_cells() no debe generar una línea de log para esa columna
        (old_value == new_value, ver excel_writer.py update_cells())."""
        self._crear_excel([
            ['VIN1', 'Expired', 'BAJA', 'COMERCIAL', '3PC-1', '15000'],
        ])
        valid = {'VIN1': self._valid_para('VIN1')}

        self.writer.update_cells(valid, module_name="Comercial", custom_filepath=self.excel_path)

        df_log = self._leer_log()
        self.assertEqual(len(df_log), 2, "Fallo: ESTADO COMERCIAL ya estaba en BAJA; no debía generar una línea de log.")
        self.assertNotIn('ESTADO COMERCIAL', df_log['CAMPO MODIFICADO'].tolist())

    @patch('src.core.excel_writer.ExcelWriter._crear_backup')
    def test_multiples_vin_en_una_sola_llamada_escriben_y_loguean_todos(self, mock_backup):
        self._crear_excel([
            ['VIN1', 'Expired', 'Vigente', 'COMERCIAL', '3PC-1', '15000'],
            ['VIN2', 'Expired', 'Vigente', 'COMERCIAL', '3PC-2', '20000'],
            ['VIN3', 'Expired', 'Vigente', 'COMERCIAL', '3PC-3', '10000'],
            ['VIN_NO_SELECCIONADO', 'Expired', 'Vigente', 'COMERCIAL', '3PC-9', '99999'],
        ])
        valid = {vin: self._valid_para(vin) for vin in ('VIN1', 'VIN2', 'VIN3')}

        resultado = self.writer.update_cells(valid, module_name="Comercial", custom_filepath=self.excel_path)
        self.assertTrue(resultado)

        for vin in ('VIN1', 'VIN2', 'VIN3'):
            self.assertEqual(self._leer_celda(vin, 'ESTADO DEL CM'), 'Terminated')
            self.assertEqual(self._leer_celda(vin, 'ESTADO COMERCIAL'), 'BAJA')
            self.assertNotEqual(self._leer_celda(vin, 'N CONTRATO 3PC'), '-')

        # El VIN no seleccionado no debe tocarse.
        self.assertEqual(self._leer_celda('VIN_NO_SELECCIONADO', 'ESTADO DEL CM'), 'Expired')
        self.assertEqual(self._leer_celda('VIN_NO_SELECCIONADO', 'ESTADO COMERCIAL'), 'Vigente')
        self.assertEqual(self._leer_celda('VIN_NO_SELECCIONADO', 'PAGO CM'), '99999')

        df_log = self._leer_log()
        self.assertEqual(len(df_log), 9, "Fallo: 3 VIN x 3 columnas cada uno = 9 líneas de log.")
        for vin in ('VIN1', 'VIN2', 'VIN3'):
            self.assertEqual(len(df_log[df_log['IDENTIFICADOR'] == vin]), 3,
                              f"Fallo: {vin} debería generar exactamente 3 líneas de log.")
        self.assertNotIn('VIN_NO_SELECCIONADO', df_log['IDENTIFICADOR'].tolist())

    @patch('src.core.excel_writer.ExcelWriter._crear_backup')
    def test_bug_regresion_dict_sin_clave_vin_no_encuentra_la_fila(self, mock_backup):
        """Regresión del bug real encontrado en comercial.py: si `valid[vin]`
        no incluye la clave 'VIN' (solo 'cambios'), update_cells() no puede
        ubicar la fila (usa data.get('VIN'), no la clave del diccionario) y
        debe fallar sin escribir nada ni generar ningún log."""
        self._crear_excel([
            ['VIN1', 'Expired', 'Vigente', 'COMERCIAL', '3PC-1', '15000'],
        ])
        valid_roto = {'VIN1': {'cambios': {'ESTADO DEL CM': 'Terminated'}}}  # sin 'VIN'

        resultado = self.writer.update_cells(valid_roto, module_name="Comercial", custom_filepath=self.excel_path)
        self.assertFalse(resultado, "Fallo: sin la clave 'VIN', update_cells() no debería poder guardar.")

        self.assertEqual(self._leer_celda('VIN1', 'ESTADO DEL CM'), 'Expired',
                          "Fallo: el Excel no debería haberse modificado con un dict sin 'VIN'.")
        self.assertTrue(self._leer_log().empty, "Fallo: no debería generarse ningún log si no se encontró la fila.")

import os
import csv
import tempfile
import unittest
from unittest.mock import patch, MagicMock

import openpyxl
import pandas as pd

import config.settings
from src.services import mail_registro


class TestMailRegistro(unittest.TestCase):
    """Servicio de datos puro del historial de correos: append-only,
    lectura y migracion unica no destructiva desde Historial_Envios_VC.xlsx."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.excel_path = os.path.join(self.temp_dir.name, "PruebasFM.xlsx")
        with open(self.excel_path, 'w', encoding='utf-8') as f:
            f.write("dummy")
        self._patch = patch.object(config.settings, 'EXCEL_PATH', self.excel_path)
        self._patch.start()

    def tearDown(self):
        self._patch.stop()
        self.temp_dir.cleanup()

    def _registro_ejemplo(self, n=1, prefijo="ENV"):
        return [
            {
                "id_envio": f"{prefijo}-{i}", "fecha_hora": "2026-09-01 10:00:00", "tipo_correo": "BIENVENIDA",
                "razon_social": f"EMPRESA {i}", "chasis": f"CHS{i}", "destinatario": f"correo{i}@test.cl",
                "estado": "ENVIADO A OUTLOOK", "asunto": "Bienvenida",
            }
            for i in range(n)
        ]

    def test_registrar_envios_crea_csv_con_columnas_correctas(self):
        resultado = mail_registro.registrar_envios(self._registro_ejemplo(2))

        self.assertTrue(resultado)
        csv_path = mail_registro.get_historial_csv_path()
        self.assertTrue(os.path.exists(csv_path))

        df = pd.read_csv(csv_path, dtype=str)
        self.assertEqual(len(df), 2)
        self.assertEqual(list(df.columns), mail_registro._COLUMNAS)
        self.assertEqual(df.iloc[0]["destinatario"], "correo0@test.cl")

    def test_registrar_envios_lista_vacia_no_crea_archivo(self):
        resultado = mail_registro.registrar_envios([])
        self.assertTrue(resultado)
        self.assertFalse(os.path.exists(mail_registro.get_historial_csv_path()))

    def test_multiples_llamadas_acumulan_sin_duplicar_encabezado(self):
        mail_registro.registrar_envios(self._registro_ejemplo(1, prefijo="ENV-A"))
        mail_registro.registrar_envios(self._registro_ejemplo(1, prefijo="ENV-B"))

        csv_path = mail_registro.get_historial_csv_path()
        with open(csv_path, encoding='utf-8-sig') as f:
            lineas = f.readlines()
        self.assertEqual(len(lineas), 3, "Fallo: debe haber 1 encabezado + 2 filas, sin encabezados repetidos.")

        df = pd.read_csv(csv_path, dtype=str)
        self.assertEqual(len(df), 2)

    def test_fallo_de_escritura_devuelve_false_sin_lanzar_excepcion(self):
        with patch('src.services.mail_registro.open', side_effect=PermissionError("bloqueado"), create=True), \
                self.assertLogs(mail_registro.logger, level='ERROR'):
            resultado = mail_registro.registrar_envios(self._registro_ejemplo(1))
        self.assertFalse(resultado)

    def test_leer_historial_sin_archivo_devuelve_dataframe_vacio_con_columnas(self):
        df = mail_registro.leer_historial()
        self.assertTrue(df.empty)
        self.assertEqual(list(df.columns), mail_registro._COLUMNAS)

    def test_leer_historial_devuelve_lo_registrado(self):
        mail_registro.registrar_envios(self._registro_ejemplo(3))
        df = mail_registro.leer_historial()
        self.assertEqual(len(df), 3)

    def _crear_legacy_xlsx(self, filas):
        legacy_path = mail_registro.get_historial_legacy_path()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(mail_registro._COLUMNAS)
        for fila in filas:
            ws.append(fila)
        wb.save(legacy_path)
        wb.close()
        return legacy_path

    def test_migracion_legado_no_destructiva_y_verificada(self):
        legacy_path = self._crear_legacy_xlsx([
            ["ENV-VIEJO", "2026-08-01 09:00:00", "BIENVENIDA", "EMPRESA VIEJA", "CHS0", "viejo@test.cl", "ENVIADO A OUTLOOK", "Bienvenida"],
        ])

        migrado = mail_registro.migrar_historial_legacy()
        self.assertTrue(migrado)

        csv_path = mail_registro.get_historial_csv_path()
        df = pd.read_csv(csv_path, dtype=str)
        self.assertEqual(len(df), 1)
        self.assertEqual(df.iloc[0]["id_envio"], "ENV-VIEJO")

        # El xlsx legado no se toca ni se borra
        self.assertTrue(os.path.exists(legacy_path))
        wb_check = openpyxl.load_workbook(legacy_path)
        self.assertEqual(wb_check.active.max_row, 2, "Fallo: el .xlsx legado no debe modificarse por la migración.")
        wb_check.close()

    def test_migracion_es_idempotente(self):
        self._crear_legacy_xlsx([
            ["ENV-VIEJO", "2026-08-01 09:00:00", "BIENVENIDA", "EMPRESA VIEJA", "CHS0", "viejo@test.cl", "ENVIADO A OUTLOOK", "Bienvenida"],
        ])
        primera = mail_registro.migrar_historial_legacy()
        self.assertTrue(primera)

        csv_path = mail_registro.get_historial_csv_path()
        with open(csv_path, 'a', newline='', encoding='utf-8-sig') as f:
            csv.writer(f).writerow(["ENV-NUEVO", "2026-09-01 10:00:00", "BIENVENIDA", "EMPRESA NUEVA", "CHS1", "nuevo@test.cl", "ENVIADO A OUTLOOK", "Bienvenida"])

        segunda = mail_registro.migrar_historial_legacy()
        self.assertFalse(segunda, "Fallo: la migración no debe repetirse si el CSV ya existe.")

        df = pd.read_csv(csv_path, dtype=str)
        self.assertEqual(len(df), 2, "Fallo: no debió perderse la fila agregada después de la migración.")

    def test_registrar_envios_migra_automaticamente_antes_de_escribir(self):
        self._crear_legacy_xlsx([
            ["ENV-VIEJO", "2026-08-01 09:00:00", "BIENVENIDA", "EMPRESA VIEJA", "CHS0", "viejo@test.cl", "ENVIADO A OUTLOOK", "Bienvenida"],
        ])

        mail_registro.registrar_envios(self._registro_ejemplo(1, prefijo="ENV-NUEVO"))

        df = mail_registro.leer_historial()
        self.assertEqual(len(df), 2)
        self.assertIn("ENV-VIEJO", df["id_envio"].tolist())
        self.assertIn("ENV-NUEVO-0", df["id_envio"].tolist())


class TestOutlookSenderThreadUsaMailRegistro(unittest.TestCase):
    """OutlookSenderThread.run() debe delegar el registro a
    mail_registro.registrar_envios y propagar su resultado en
    finished_signal, para que la UI pueda avisar si el correo se envio
    pero el registro fallo (en vez de solo un print())."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.excel_path = os.path.join(self.temp_dir.name, "PruebasFM.xlsx")
        with open(self.excel_path, 'w', encoding='utf-8') as f:
            f.write("dummy")
        self._patch = patch.object(config.settings, 'EXCEL_PATH', self.excel_path)
        self._patch.start()

    def tearDown(self):
        self._patch.stop()
        self.temp_dir.cleanup()

    def _run_thread_sync(self, registro_mock_return):
        from src.ui.mail_dialog import OutlookSenderThread

        destinatarios = [{
            "correo": "test@empresa.cl", "flota": "EMPRESA TEST", "nombre": "Juan Perez",
            "chasis_list": ["CHS001"], "grupos_plan": {},
        }]

        thread = OutlookSenderThread(
            destinatarios, cc="", plantilla="Personalizado",
            asunto="Prueba", mensaje_base="Hola @NOMBRE de @FLOTA",
        )

        capturado = {}
        thread.finished_signal.connect(lambda e, f, r: capturado.update(exitos=e, fallos=f, registro_ok=r))

        fake_outlook = MagicMock()
        fake_mail = MagicMock()
        fake_outlook.CreateItem.return_value = fake_mail

        with patch('src.ui.mail_dialog.win32com.client.Dispatch', return_value=fake_outlook), \
             patch('src.ui.mail_dialog.pythoncom.CoInitialize'), \
             patch('src.ui.mail_dialog.pythoncom.CoUninitialize'), \
             patch('src.services.mail_registro.registrar_envios', return_value=registro_mock_return) as mock_registrar:
            thread.run()

        return capturado, mock_registrar, fake_mail

    def test_registro_exitoso_se_informa_en_finished_signal(self):
        capturado, mock_registrar, _ = self._run_thread_sync(registro_mock_return=True)

        self.assertEqual(capturado['exitos'], 1)
        self.assertEqual(capturado['fallos'], 0)
        self.assertTrue(capturado['registro_ok'])
        mock_registrar.assert_called_once()

    def test_fallo_de_registro_se_informa_en_finished_signal_aunque_outlook_funcione(self):
        capturado, mock_registrar, _ = self._run_thread_sync(registro_mock_return=False)

        self.assertEqual(capturado['exitos'], 1, "Fallo: el correo debe seguir considerándose enviado a Outlook.")
        self.assertEqual(capturado['fallos'], 0)
        self.assertFalse(capturado['registro_ok'],
                          "Fallo: finished_signal debe reflejar que el registro del historial falló.")

    def test_personalizado_asigna_html_body_con_variables_reemplazadas(self):
        # Regresion del bug detectado en la auditoria: enviar_correos() le pasa
        # "Personalizado" (corto) a OutlookSenderThread, pero run() comparaba
        # contra "Personalizado (Escribir desde cero)" y nunca coincidia, asi
        # que el cuerpo del correo personalizado nunca se asignaba.
        capturado, _, fake_mail = self._run_thread_sync(registro_mock_return=True)

        self.assertEqual(capturado['exitos'], 1)
        # Desde el 21-10-2026 el texto va dentro de la barra azul con el logo
        self.assertIn("Hola Juan de EMPRESA TEST", fake_mail.HTMLBody)
        self.assertIn("cid:logo_volvo_blanco", fake_mail.HTMLBody)


if __name__ == '__main__':
    unittest.main(verbosity=2)

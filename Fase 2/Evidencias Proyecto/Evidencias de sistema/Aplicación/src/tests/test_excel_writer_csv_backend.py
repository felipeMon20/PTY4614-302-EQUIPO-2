"""Pruebas del backend CSV de ExcelWriter (Fase 1, migracion PruebasFM.xlsx
-> .csv, 23-09-2026, ver excel_writer_csv_backend.py). Mismo patron que
test_excel_writer.py: llama a update_cells/add_new_row/delete_row REALES
(sin mocks de la propia ExcelWriter salvo _crear_backup, para no tocar
data/backups real), con custom_filepath apuntando a un CSV temporal -- nunca
toca PruebasFM.xlsx ni EXCEL_PATH real."""

import os
import tempfile
import unittest
from unittest.mock import patch

import openpyxl
import pandas as pd

from src.core.excel_writer import ExcelWriter
from src.core.data_manager import DataManager
from src.core.dataframe_formatter import DataFrameFormatter

COLUMNAS = [
    'VIN', 'CHASIS', 'RAZÓN SOCIAL', 'ESTADO COMERCIAL',
    'INICIO CONTRATO', 'TÉRMINO CONTRATO', 'FECHA DE BAJA',
    'INICIO', 'TERMINO', 'ESTADO DE ENTREGA',
    'FECHA INICIO CM', 'FECHA TERMINO CM',
]

FILAS = [
    # CHASIS con prefijo de letra (formato real confirmado en PruebasFM.xlsx,
    # ej. 'E920364'): un CHASIS puramente numerico ('0900054') dispara un
    # comportamiento de pandas.read_excel ajeno a ExcelWriter (infiere toda
    # la columna como int64 y pierde ceros iniciales SOLO si el 100% de los
    # valores de la columna son puramente numericos, algo que no ocurre en
    # los datos reales -- ver informe de esta fase para el detalle).
    ['YV2AAA0000000001', 'E000001', 'EMPRESA UNO SPA', 'Vigente',
     '01-01-2025', '01-01-2026', '-', '01-01-2025', '01-06-2025', '01-01-2025', '-', '-'],
    ['YV2BBB0000000002', 'E000002', 'EMPRESA DOS LTDA', 'Vigente',
     '15-03-2024', 'MENSUAL', '-', '-', '-', '-', '01-02-2024', '01-02-2025'],
    ['YV2CCC0000000003', 'E000003', 'EMPRESA TRES SA', 'Baja',
     '01-06-2023', '01-06-2024', '10-09-2026', '-', '-', '-', '-', '-'],
]


def _crear_csv(ruta):
    df = pd.DataFrame(FILAS, columns=COLUMNAS)
    df.to_csv(ruta, index=False, encoding='utf-8-sig')
    return df


def _crear_xlsx(ruta):
    df = pd.DataFrame(FILAS, columns=COLUMNAS)
    df.to_excel(ruta, index=False, sheet_name='BBDD (2)')
    return df


class _BaseCsvBackendTest(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.writer = ExcelWriter()
        self.writer.log_filepath = ""
        self.writer.log_csv_filepath = os.path.join(self.temp_dir.name, "LogApp_test.csv")
        self.writer.deleted_filepath = os.path.join(self.temp_dir.name, "RegistroEliminado_test.xlsx")
        self.csv_path = os.path.join(self.temp_dir.name, "PruebasFM_test.csv")
        _crear_csv(self.csv_path)
        self.dm = DataManager()

    def tearDown(self):
        self.temp_dir.cleanup()

    def _leer(self):
        return pd.read_csv(self.csv_path, dtype=str, keep_default_na=False)


# =========================================================
# A. update_cells
# =========================================================
class TestUpdateCellsCsv(_BaseCsvBackendTest):
    @patch('src.core.excel_writer.ExcelWriter._crear_backup')
    def test_modificar_una_columna_de_una_fila(self, mock_backup):
        res = self.writer.update_cells(
            {"1": {"VIN": "YV2AAA0000000001", "cambios": {"ESTADO COMERCIAL": "Baja"}}},
            module_name="Test", custom_filepath=self.csv_path,
        )
        self.assertTrue(res)
        df = self._leer()
        fila = df[df['VIN'] == 'YV2AAA0000000001'].iloc[0]
        self.assertEqual(fila['ESTADO COMERCIAL'], 'Baja')

    @patch('src.core.excel_writer.ExcelWriter._crear_backup')
    def test_modificar_varias_columnas_de_una_fila(self, mock_backup):
        res = self.writer.update_cells(
            {"1": {"VIN": "YV2AAA0000000001", "cambios": {
                "ESTADO COMERCIAL": "Revisar", "RAZÓN SOCIAL": "EMPRESA UNO RENOMBRADA",
            }}},
            module_name="Test", custom_filepath=self.csv_path,
        )
        self.assertTrue(res)
        fila = self._leer().set_index('VIN').loc['YV2AAA0000000001']
        self.assertEqual(fila['ESTADO COMERCIAL'], 'Revisar')
        self.assertEqual(fila['RAZÓN SOCIAL'], 'EMPRESA UNO RENOMBRADA')

    @patch('src.core.excel_writer.ExcelWriter._crear_backup')
    def test_modificar_varias_filas(self, mock_backup):
        res = self.writer.update_cells(
            {
                "1": {"VIN": "YV2AAA0000000001", "cambios": {"ESTADO COMERCIAL": "Baja"}},
                "2": {"VIN": "YV2BBB0000000002", "cambios": {"ESTADO COMERCIAL": "Revisar"}},
            },
            module_name="Test", custom_filepath=self.csv_path,
        )
        self.assertTrue(res)
        df = self._leer().set_index('VIN')
        self.assertEqual(df.loc['YV2AAA0000000001', 'ESTADO COMERCIAL'], 'Baja')
        self.assertEqual(df.loc['YV2BBB0000000002', 'ESTADO COMERCIAL'], 'Revisar')
        # La tercera fila no debe verse tocada.
        self.assertEqual(df.loc['YV2CCC0000000003', 'ESTADO COMERCIAL'], 'Baja')

    @patch('src.core.excel_writer.ExcelWriter._crear_backup')
    def test_vin_inexistente_retorna_false_y_no_modifica_nada(self, mock_backup):
        df_antes = self._leer()
        res = self.writer.update_cells(
            {"1": {"VIN": "VIN_QUE_NO_EXISTE", "cambios": {"ESTADO COMERCIAL": "Baja"}}},
            module_name="Test", custom_filepath=self.csv_path,
        )
        self.assertFalse(res)
        df_despues = self._leer()
        pd.testing.assert_frame_equal(df_antes, df_despues)

    @patch('src.core.excel_writer.ExcelWriter._crear_backup')
    def test_termino_contrato_nunca_se_escribe(self, mock_backup):
        """Misma regla que la rama .xlsx (excel_writer.py: 'if col_name ==
        TÉRMINO CONTRATO: continue'): se recalcula siempre en
        DataFrameFormatter.clean(), nunca se persiste directo."""
        res = self.writer.update_cells(
            {"1": {"VIN": "YV2BBB0000000002", "cambios": {"TÉRMINO CONTRATO": "31-12-2030"}}},
            module_name="Test", custom_filepath=self.csv_path,
        )
        self.assertTrue(res)
        fila = self._leer().set_index('VIN').loc['YV2BBB0000000002']
        self.assertEqual(fila['TÉRMINO CONTRATO'], 'MENSUAL')  # sin cambios

    @patch('src.core.excel_writer.ExcelWriter._crear_backup')
    def test_solo_cambian_las_celdas_esperadas(self, mock_backup):
        df_antes = self._leer()
        res = self.writer.update_cells(
            {"1": {"VIN": "YV2CCC0000000003", "cambios": {"ESTADO COMERCIAL": "Baja Solicitada"}}},
            module_name="Test", custom_filepath=self.csv_path,
        )
        self.assertTrue(res)
        df_despues = self._leer()
        diffs = (df_antes.reset_index(drop=True) != df_despues.reset_index(drop=True))
        columnas_con_diff = diffs.any(axis=0)
        self.assertEqual(list(columnas_con_diff[columnas_con_diff].index), ['ESTADO COMERCIAL'])
        self.assertEqual(diffs.values.sum(), 1)

    @patch('src.core.excel_writer.ExcelWriter._crear_backup')
    def test_is_app_saving_y_save_timestamp(self, mock_backup):
        self.assertFalse(self.dm.is_app_saving)
        ts_antes = self.dm.save_timestamp
        res = self.writer.update_cells(
            {"1": {"VIN": "YV2AAA0000000001", "cambios": {"ESTADO COMERCIAL": "Baja"}}},
            module_name="Test", custom_filepath=self.csv_path,
        )
        self.assertTrue(res)
        self.assertFalse(self.dm.is_app_saving)  # se resetea al terminar
        self.assertGreaterEqual(self.dm.save_timestamp, ts_antes)


# =========================================================
# B. add_new_row
# =========================================================
class TestAddNewRowCsv(_BaseCsvBackendTest):
    @patch('src.core.excel_writer.ExcelWriter._crear_backup')
    def test_agregar_fila_valida(self, mock_backup):
        nueva = {'VIN': 'YV2DDD0000000004', 'CHASIS': '0900074', 'RAZÓN SOCIAL': 'EMPRESA CUATRO',
                  'ESTADO COMERCIAL': 'Vigente'}
        res = self.writer.add_new_row(nueva, module_name="Test", custom_filepath=self.csv_path)
        self.assertTrue(res)

        df = self._leer()
        self.assertEqual(len(df), 4)  # verificar cantidad de filas
        self.assertIn('YV2DDD0000000004', df['VIN'].tolist())  # verificar VIN
        self.assertEqual(list(df.columns), COLUMNAS)  # verificar columnas (sin alterar)

        fila_nueva = df[df['VIN'] == 'YV2DDD0000000004'].iloc[0]
        self.assertEqual(fila_nueva['CHASIS'], '0900074')
        self.assertEqual(fila_nueva['RAZÓN SOCIAL'], 'EMPRESA CUATRO')

        # Filas existentes no alteradas.
        self.assertEqual(df.iloc[0]['VIN'], 'YV2AAA0000000001')
        self.assertEqual(df.iloc[1]['VIN'], 'YV2BBB0000000002')
        self.assertEqual(df.iloc[2]['VIN'], 'YV2CCC0000000003')

    @patch('src.core.excel_writer.ExcelWriter._crear_backup')
    def test_columna_faltante_en_dict_queda_en_blanco(self, mock_backup):
        """Mismo comportamiento que la rama .xlsx: una columna existente del
        dataset que NO viene en new_data_dict queda vacia, no falla."""
        nueva = {'VIN': 'YV2EEE0000000005', 'CHASIS': '0900075'}  # sin RAZÓN SOCIAL ni el resto
        res = self.writer.add_new_row(nueva, module_name="Test", custom_filepath=self.csv_path)
        self.assertTrue(res)

        fila = self._leer().set_index('VIN').loc['YV2EEE0000000005']
        self.assertEqual(fila['CHASIS'], '0900075')
        self.assertEqual(fila['RAZÓN SOCIAL'], '')
        self.assertEqual(fila['ESTADO COMERCIAL'], '')
        self.assertEqual(fila['TÉRMINO CONTRATO'], '')

    @patch('src.core.excel_writer.ExcelWriter._crear_backup')
    def test_columna_desconocida_en_dict_se_ignora(self, mock_backup):
        """Mismo comportamiento que la rama .xlsx: una clave de
        new_data_dict que no exista como columna del dataset se ignora
        silenciosamente -- NO se crea una columna nueva."""
        nueva = {'VIN': 'YV2FFF0000000006', 'CHASIS': '0000006', 'COLUMNA_INVENTADA': 'algo'}
        res = self.writer.add_new_row(nueva, module_name="Test", custom_filepath=self.csv_path)
        self.assertTrue(res)

        df = self._leer()
        self.assertEqual(list(df.columns), COLUMNAS)  # ninguna columna nueva
        self.assertNotIn('COLUMNA_INVENTADA', df.columns)

    @patch('src.core.excel_writer.ExcelWriter._crear_backup')
    def test_valor_falsy_no_se_escribe(self, mock_backup):
        """Mismo criterio que la rama .xlsx ('if col_name in header_map and
        value'): un valor falsy (aqui, string vacio) no sobreescribe el
        default en blanco -- comportamiento identico, no hay diferencia
        observable, pero se confirma que no lanza error."""
        nueva = {'VIN': 'YV2GGG0000000007', 'CHASIS': '0000007', 'RAZÓN SOCIAL': ''}
        res = self.writer.add_new_row(nueva, module_name="Test", custom_filepath=self.csv_path)
        self.assertTrue(res)
        fila = self._leer().set_index('VIN').loc['YV2GGG0000000007']
        self.assertEqual(fila['RAZÓN SOCIAL'], '')


# =========================================================
# C. delete_row
# =========================================================
class TestDeleteRowCsv(_BaseCsvBackendTest):
    @patch('src.core.excel_writer.ExcelWriter._crear_backup')
    def test_eliminar_vin_existente(self, mock_backup):
        res = self.writer.delete_row('YV2BBB0000000002', module_name="Test", custom_filepath=self.csv_path)
        self.assertTrue(res)
        df = self._leer()
        self.assertEqual(len(df), 2)
        self.assertNotIn('YV2BBB0000000002', df['VIN'].tolist())

    @patch('src.core.excel_writer.ExcelWriter._crear_backup')
    def test_filas_restantes_intactas(self, mock_backup):
        res = self.writer.delete_row('YV2BBB0000000002', module_name="Test", custom_filepath=self.csv_path)
        self.assertTrue(res)
        df = self._leer().set_index('VIN')
        self.assertEqual(df.loc['YV2AAA0000000001', 'RAZÓN SOCIAL'], 'EMPRESA UNO SPA')
        self.assertEqual(df.loc['YV2CCC0000000003', 'RAZÓN SOCIAL'], 'EMPRESA TRES SA')

    @patch('src.core.excel_writer.ExcelWriter._crear_backup')
    def test_vin_inexistente_retorna_false(self, mock_backup):
        df_antes = self._leer()
        res = self.writer.delete_row('VIN_QUE_NO_EXISTE', module_name="Test", custom_filepath=self.csv_path)
        self.assertFalse(res)
        pd.testing.assert_frame_equal(df_antes, self._leer())

    @patch('src.core.excel_writer.ExcelWriter._crear_backup')
    def test_logging_de_eliminacion_en_registro_eliminado(self, mock_backup):
        res = self.writer.delete_row('YV2CCC0000000003', module_name="Test", custom_filepath=self.csv_path)
        self.assertTrue(res)

        wb = openpyxl.load_workbook(self.writer.deleted_filepath)
        self.assertIn("Test", wb.sheetnames)
        ws = wb["Test"]
        encabezado = [c.value for c in ws[1]]
        self.assertEqual(encabezado, ["FECHA Y HORA DE ELIMINACIÓN", "MÓDULO"] + COLUMNAS)
        fila_registrada = [c.value for c in ws[2]]
        vin_idx = encabezado.index('VIN')
        self.assertEqual(fila_registrada[vin_idx], 'YV2CCC0000000003')
        wb.close()


# =========================================================
# D. Fechas tras update_cells
# =========================================================
class TestFechasTrasUpdateCsv(_BaseCsvBackendTest):
    COLS_FECHA = [
        'INICIO CONTRATO', 'TÉRMINO CONTRATO', 'FECHA DE BAJA', 'INICIO', 'TERMINO',
        'ESTADO DE ENTREGA', 'FECHA INICIO CM', 'FECHA TERMINO CM',
    ]

    @patch('src.core.excel_writer.ExcelWriter._crear_backup')
    def test_fechas_no_se_convierten_tras_editar_otras_columnas(self, mock_backup):
        res = self.writer.update_cells(
            {
                "1": {"VIN": "YV2AAA0000000001", "cambios": {
                    "FECHA DE BAJA": "05-05-2026", "ESTADO DE ENTREGA": "10-10-2025",
                }},
                "2": {"VIN": "YV2BBB0000000002", "cambios": {"INICIO": "20-01-2026"}},
            },
            module_name="Test", custom_filepath=self.csv_path,
        )
        self.assertTrue(res)
        df = self._leer()
        for col in self.COLS_FECHA:
            for val in df[col]:
                self.assertIsInstance(val, str, f"{col} contiene un valor no-str: {val!r}")
                self.assertNotRegex(val, r'^\d{4}-\d{2}-\d{2}', f"{col} quedo en formato YYYY-MM-DD: {val!r}")
                self.assertNotIn('00:00:00', val, f"{col} quedo con componente de hora: {val!r}")

        fila1 = df.set_index('VIN').loc['YV2AAA0000000001']
        self.assertEqual(fila1['FECHA DE BAJA'], '05-05-2026')
        self.assertEqual(fila1['ESTADO DE ENTREGA'], '10-10-2025')
        fila2 = df.set_index('VIN').loc['YV2BBB0000000002']
        self.assertEqual(fila2['INICIO'], '20-01-2026')
        self.assertEqual(fila2['TÉRMINO CONTRATO'], 'MENSUAL')  # intacto


# =========================================================
# E. Integridad
# =========================================================
class TestIntegridadCsv(_BaseCsvBackendTest):
    @patch('src.core.excel_writer.ExcelWriter._crear_backup')
    def test_columnas_y_orden_se_preservan_tras_update(self, mock_backup):
        self.writer.update_cells(
            {"1": {"VIN": "YV2AAA0000000001", "cambios": {"ESTADO COMERCIAL": "Baja"}}},
            module_name="Test", custom_filepath=self.csv_path,
        )
        self.assertEqual(list(self._leer().columns), COLUMNAS)

    @patch('src.core.excel_writer.ExcelWriter._crear_backup')
    def test_columnas_y_orden_se_preservan_tras_add(self, mock_backup):
        self.writer.add_new_row(
            {'VIN': 'YV2HHH0000000008', 'CHASIS': '0000008'},
            module_name="Test", custom_filepath=self.csv_path,
        )
        self.assertEqual(list(self._leer().columns), COLUMNAS)

    @patch('src.core.excel_writer.ExcelWriter._crear_backup')
    def test_sin_perdida_de_registros_tras_ciclo_completo(self, mock_backup):
        """add -> update -> delete de la fila agregada: al final deben
        quedar exactamente las 3 filas originales, sin registros
        adicionales ni faltantes."""
        self.writer.add_new_row({'VIN': 'YV2ZZZ0000000009', 'CHASIS': '0000009'},
                                 module_name="Test", custom_filepath=self.csv_path)
        self.assertEqual(len(self._leer()), 4)

        self.writer.update_cells(
            {"1": {"VIN": "YV2ZZZ0000000009", "cambios": {"ESTADO COMERCIAL": "Vigente"}}},
            module_name="Test", custom_filepath=self.csv_path,
        )
        self.assertEqual(len(self._leer()), 4)

        self.writer.delete_row('YV2ZZZ0000000009', module_name="Test", custom_filepath=self.csv_path)
        df_final = self._leer()
        self.assertEqual(len(df_final), 3)
        self.assertEqual(sorted(df_final['VIN'].tolist()),
                          ['YV2AAA0000000001', 'YV2BBB0000000002', 'YV2CCC0000000003'])


# =========================================================
# 10. Equivalencia funcional XLSX vs CSV
# =========================================================
class TestEquivalenciaXlsxVsCsv(unittest.TestCase):
    """Aplica la MISMA secuencia de operaciones sobre una copia .xlsx y una
    copia .csv de los mismos datos iniciales, y compara el resultado final.
    Ignora únicamente diferencias de dtype (ver Fase 0: pandas 3.0.5
    distingue 'object' vs 'str' aunque el contenido sea idéntico) -- todo lo
    demás debe coincidir exactamente."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.xlsx_path = os.path.join(self.temp_dir.name, "PruebasFM_equiv.xlsx")
        self.csv_path = os.path.join(self.temp_dir.name, "PruebasFM_equiv.csv")
        _crear_xlsx(self.xlsx_path)
        _crear_csv(self.csv_path)

        self.writer = ExcelWriter()
        self.writer.log_filepath = ""
        self.writer.log_csv_filepath = os.path.join(self.temp_dir.name, "LogApp_equiv.csv")
        self.writer.deleted_filepath = os.path.join(self.temp_dir.name, "RegistroEliminado_equiv.xlsx")

    def tearDown(self):
        self.temp_dir.cleanup()

    def _leer_xlsx(self):
        # Mismo pipeline que MainDataLoaderThread.run() (ver Fase 0): la
        # comparacion valida es POST-clean(), no la lectura cruda -- una
        # celda de Excel nunca tocada llega como NaN a pandas y una fila de
        # CSV nunca tocada llega como "" (ambas se normalizan a "" recien en
        # clean()/.fillna("")). Comparar antes de clean() mezclaria una
        # diferencia real con una simple diferencia de representacion que la
        # propia app ya resuelve en su pipeline normal.
        df = pd.read_excel(self.xlsx_path, sheet_name='BBDD (2)')
        return DataFrameFormatter.clean(df)

    def _leer_csv(self):
        df = pd.read_csv(self.csv_path, dtype=str, keep_default_na=False)
        return DataFrameFormatter.clean(df)

    @patch('src.core.excel_writer.ExcelWriter._crear_backup')
    def test_misma_secuencia_de_operaciones_produce_el_mismo_resultado(self, mock_backup):
        # 1) update_cells sobre ambos.
        cambios = {"1": {"VIN": "YV2AAA0000000001", "cambios": {
            "ESTADO COMERCIAL": "Baja", "FECHA DE BAJA": "20-09-2026",
        }}}
        self.assertTrue(self.writer.update_cells(cambios, module_name="Test", custom_filepath=self.xlsx_path))
        self.assertTrue(self.writer.update_cells(cambios, module_name="Test", custom_filepath=self.csv_path))

        # 2) add_new_row sobre ambos.
        nueva = {'VIN': 'YV2NEW0000000099', 'CHASIS': '0000099', 'RAZÓN SOCIAL': 'EMPRESA NUEVA',
                  'ESTADO COMERCIAL': 'Vigente'}
        self.assertTrue(self.writer.add_new_row(nueva, module_name="Test", custom_filepath=self.xlsx_path))
        self.assertTrue(self.writer.add_new_row(nueva, module_name="Test", custom_filepath=self.csv_path))

        # 3) delete_row sobre ambos.
        self.assertTrue(self.writer.delete_row('YV2BBB0000000002', module_name="Test", custom_filepath=self.xlsx_path))
        self.assertTrue(self.writer.delete_row('YV2BBB0000000002', module_name="Test", custom_filepath=self.csv_path))

        df_xlsx = self._leer_xlsx().sort_values('VIN').reset_index(drop=True)
        df_csv = self._leer_csv().sort_values('VIN').reset_index(drop=True)

        self.assertEqual(list(df_xlsx.columns), list(df_csv.columns))
        self.assertEqual(len(df_xlsx), len(df_csv))
        self.assertEqual(sorted(df_xlsx['VIN'].tolist()), sorted(df_csv['VIN'].tolist()))

        # Comparacion celda a celda como texto (ignora dtype, ver docstring).
        diffs = (df_xlsx.astype(str) != df_csv.astype(str))
        celdas_distintas = diffs.values.sum()
        if celdas_distintas > 0:
            columnas_con_diff = diffs.any(axis=0)
            detalle = []
            for col in columnas_con_diff[columnas_con_diff].index:
                for idx in df_xlsx.index[diffs[col]]:
                    detalle.append(f"fila={idx} col={col!r} xlsx={df_xlsx.at[idx, col]!r} csv={df_csv.at[idx, col]!r}")
            self.fail("Diferencia real de valor entre XLSX y CSV tras la misma secuencia de "
                      f"operaciones ({celdas_distintas} celdas):\n" + "\n".join(detalle))


if __name__ == '__main__':
    unittest.main()

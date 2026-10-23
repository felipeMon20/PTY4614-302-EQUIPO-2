# -*- coding: utf-8 -*-
"""Pruebas de src/tests/auditor/quality_sync.py.

IMPORTANTE: estas pruebas NUNCA operan sobre los archivos reales de
duoc/calidad. Cada test copia los 6 archivos reales a un directorio
temporal (setUp) y hace sync() unicamente sobre esas copias; el directorio
temporal se borra en tearDown. No se requiere el sandbox de
test_main.py::setup_sandbox (estas pruebas no tocan Excel de negocio ni
app_config.json), por lo que corren igual de forma aislada.
"""
import os
import shutil
import sys
import tempfile
import unittest
from datetime import datetime

BASE_DIR = os.path.abspath(os.path.dirname(__file__))
AUDITOR_DIR = os.path.join(BASE_DIR, "auditor")
PROJECT_ROOT = os.path.abspath(os.path.join(BASE_DIR, "..", ".."))
REAL_CALIDAD_DIR = os.path.join(PROJECT_ROOT, "duoc", "calidad")

if AUDITOR_DIR not in sys.path:
    sys.path.insert(0, AUDITOR_DIR)
import quality_sync  # noqa: E402

import openpyxl  # noqa: E402
import docx      # noqa: E402


CALIDAD_FILES = [
    "Reporte de Ejecución de Pruebas.xlsx",
    "Casos de Prueba.xlsx",
    "Plan de Pruebas.docx",
    "Plan de Gestión de Cambios.docx",
    "Request for Change (RFC).docx",
    "Reporte de Defectos.docx",
]


def _test_record(file, cls, method, status="PASS", duration_ms=1.0, tags=None,
                  docstring="Docstring de prueba.", imports=None):
    return {
        "file": file, "class": cls, "method": method,
        "module_group": file[len("test_"):-3] if file.startswith("test_") else file[:-3],
        "class_docstring": docstring,
        "relevant_imports": imports or [],
        "classification": {"tags": tags or ["unitaria"], "classification_source": "heuristic",
                            "confidence": "medium"},
        "status": status, "duration_ms": duration_ms, "error": None, "skip_reason": None,
    }


def _make_data(tests):
    total = len(tests)
    durations = [t["duration_ms"] for t in tests] or [0.0]
    pass_n = sum(1 for t in tests if t["status"] == "PASS")
    fail_n = sum(1 for t in tests if t["status"] == "FAIL")
    error_n = sum(1 for t in tests if t["status"] == "ERROR")
    skipped_n = sum(1 for t in tests if t["status"] == "SKIPPED")
    slowest = max(tests, key=lambda t: t["duration_ms"]) if tests else None
    fastest = min(tests, key=lambda t: t["duration_ms"]) if tests else None
    files = {t["file"] for t in tests}
    classes = {(t["file"], t["class"]) for t in tests}
    return {
        "tests": tests,
        "metadata": {
            "generated_at": datetime.now().isoformat(timespec="seconds"),
            "python_version": sys.version.split()[0],
            "pytest_installed": False,
            "execution_mechanism": "synthetic-test",
            "files_discovered_for_execution": len(files),
            "classes_discovered": len(classes),
            "tests_discovered_static": total,
        },
        "summary": {
            "tests_executed": total, "pass": pass_n, "fail": fail_n, "error": error_n,
            "skipped": skipped_n, "discovery_errors": 0,
        },
        "metrics": {
            "pass_rate_pct": round(pass_n / total * 100, 2) if total else 0,
            "fail_rate_pct": round(fail_n / total * 100, 2) if total else 0,
            "error_rate_pct": round(error_n / total * 100, 2) if total else 0,
            "duration_total_ms": round(sum(durations), 3),
            "duration_avg_ms": round(sum(durations) / len(durations), 3),
            "duration_median_ms": round(sorted(durations)[len(durations) // 2], 3),
            "slowest_test": {"test_id": f"{slowest['file']}.{slowest['class']}.{slowest['method']}",
                              "duration_ms": slowest["duration_ms"]} if slowest else None,
            "fastest_test": {"test_id": f"{fastest['file']}.{fastest['class']}.{fastest['method']}",
                              "duration_ms": fastest["duration_ms"]} if fastest else None,
        },
    }


class QualitySyncTestCase(unittest.TestCase):
    """Base: copia los 6 archivos reales de duoc/calidad a un tempdir antes
    de cada test; nunca opera sobre los archivos reales."""

    def setUp(self):
        self.tmpdir = tempfile.mkdtemp(prefix="quality_sync_test_")
        for name in CALIDAD_FILES:
            shutil.copy2(os.path.join(REAL_CALIDAD_DIR, name), os.path.join(self.tmpdir, name))

    def tearDown(self):
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    # -- helpers de lectura de las copias temporales --------------------
    def _reporte_ws(self):
        wb = openpyxl.load_workbook(os.path.join(self.tmpdir, "Reporte de Ejecución de Pruebas.xlsx"))
        return wb.active

    def _casos_ws(self):
        wb = openpyxl.load_workbook(os.path.join(self.tmpdir, "Casos de Prueba.xlsx"))
        return wb.active

    def _find_row_by_cp(self, ws, cp_id, col=2):
        for row in range(2, ws.max_row + 1):
            if ws.cell(row=row, column=col).value == cp_id:
                return row
        return None

    def _find_row_by_identity_in_comentarios(self, ws, file, cls):
        import re as _re
        pattern = _re.compile(
            rf"Test automatizado:\s*(?:[\w./-]*/)?{_re.escape(file)}::{_re.escape(cls)}\."
        )
        for row in range(2, ws.max_row + 1):
            v = ws.cell(row=row, column=9).value
            if isinstance(v, str) and pattern.search(v):
                return row
        return None

    def _find_row_by_observaciones(self, ws, file, cls, method):
        prefix = f"{cls}.{method} — {file}"
        for row in range(2, ws.max_row + 1):
            v = ws.cell(row=row, column=8).value
            if isinstance(v, str) and v.startswith(prefix):
                return row
        return None

    def _max_cp(self, ws):
        best = 0
        for row in range(2, ws.max_row + 1):
            v = ws.cell(row=row, column=2).value
            if isinstance(v, str) and v.startswith("CP-"):
                try:
                    best = max(best, int(v.split("-")[1]))
                except ValueError:
                    pass
        return best


class TestSincronizacionNormal(QualitySyncTestCase):
    """Caso 1: una corrida normal actualiza una fila existente y agrega una
    fila nueva, sin romper nada."""

    def test_actualiza_existente_y_agrega_nueva(self):
        existente = _test_record("test_comercial.py", "TestComercial", "test_01_carga_datos_falsos",
                                  status="PASS", duration_ms=99.9)
        nueva = _test_record("test_zzz_sync_normal.py", "TestZzzSyncNormal", "test_algo_nuevo")
        data = _make_data([existente, nueva])

        report = quality_sync.sync(data, PROJECT_ROOT, calidad_dir=self.tmpdir)

        self.assertTrue(report["ok"], report)
        self.assertFalse(report["skipped"])

        ws = self._reporte_ws()
        row = self._find_row_by_observaciones(ws, "test_comercial.py", "TestComercial",
                                                "test_01_carga_datos_falsos")
        self.assertIsNotNone(row)
        self.assertEqual(ws.cell(row=row, column=6).value, 99.9)

        new_row = None
        for r in range(2, ws.max_row + 1):
            v = ws.cell(row=r, column=8).value
            if isinstance(v, str) and "TestZzzSyncNormal.test_algo_nuevo" in v:
                new_row = r
                break
        self.assertIsNotNone(new_row, "la fila nueva debe existir en el Excel")


class TestIdempotencia(QualitySyncTestCase):
    """Caso 2: correr sync() dos veces con los mismos datos no duplica
    filas ni reasigna CP-XXX."""

    def test_segunda_corrida_no_duplica(self):
        t = _test_record("test_zzz_idempotencia.py", "TestZzzIdempotencia", "test_uno")
        data = _make_data([t])

        r1 = quality_sync.sync(data, PROJECT_ROOT, calidad_dir=self.tmpdir)
        cp_asignado = r1["documents"]["Reporte de Ejecución de Pruebas.xlsx"]["result"]["id_by_identity"][
            ("test_zzz_idempotencia.py", "TestZzzIdempotencia", "test_uno")
        ]
        ws = self._reporte_ws()
        filas_con_ese_test = sum(
            1 for r in range(2, ws.max_row + 1)
            if ws.cell(row=r, column=8).value == "TestZzzIdempotencia.test_uno — test_zzz_idempotencia.py"
        )
        self.assertEqual(filas_con_ese_test, 1)

        r2 = quality_sync.sync(data, PROJECT_ROOT, calidad_dir=self.tmpdir)
        cp_segunda = r2["documents"]["Reporte de Ejecución de Pruebas.xlsx"]["result"]["id_by_identity"][
            ("test_zzz_idempotencia.py", "TestZzzIdempotencia", "test_uno")
        ]
        self.assertEqual(cp_asignado, cp_segunda, "el CP no debe cambiar entre corridas")

        ws2 = self._reporte_ws()
        filas_con_ese_test_2 = sum(
            1 for r in range(2, ws2.max_row + 1)
            if ws2.cell(row=r, column=8).value == "TestZzzIdempotencia.test_uno — test_zzz_idempotencia.py"
        )
        self.assertEqual(filas_con_ese_test_2, 1, "no debe haber una segunda fila para el mismo test")

        casos_ws = self._casos_ws()
        filas_clase = sum(
            1 for r in range(2, casos_ws.max_row + 1)
            if self._find_row_by_identity_in_comentarios(
                casos_ws, "test_zzz_idempotencia.py", "TestZzzIdempotencia") == r
        )
        self.assertEqual(filas_clase, 1, "no debe haber dos filas para la misma clase en Casos de Prueba")


class TestNuevoTestNoRenumera(QualitySyncTestCase):
    """Caso 3: un test nuevo recibe el siguiente CP libre; los IDs ya
    publicados no se tocan."""

    def test_nuevo_test_recibe_cp_siguiente_sin_afectar_existentes(self):
        ws_antes = self._reporte_ws()
        max_antes = self._max_cp(ws_antes)
        fila_cp001_antes = [ws_antes.cell(row=2, column=c).value for c in range(1, 9)]

        nuevo = _test_record("test_zzz_nuevo.py", "TestZzzNuevo", "test_uno")
        data = _make_data([nuevo])
        report = quality_sync.sync(data, PROJECT_ROOT, calidad_dir=self.tmpdir)

        cp_nuevo = report["documents"]["Reporte de Ejecución de Pruebas.xlsx"]["result"]["id_by_identity"][
            ("test_zzz_nuevo.py", "TestZzzNuevo", "test_uno")
        ]
        self.assertEqual(cp_nuevo, f"CP-{max_antes + 1:03d}")

        ws_despues = self._reporte_ws()
        fila_cp001_despues = [ws_despues.cell(row=2, column=c).value for c in range(1, 9)]
        # CP-001 (fila 2) no estaba en el `data` de esta corrida sintetica:
        # su CP no debe cambiar aunque su Fecha/Estado tampoco se refresque.
        self.assertEqual(fila_cp001_antes[1], fila_cp001_despues[1], "CP-001 no debe renumerarse")


class TestMetodoNuevoEnClaseExistente(QualitySyncTestCase):
    """Caso 4: una clase que ya tiene fila en Casos de Prueba y gana un
    metodo nuevo actualiza SU MISMA fila, no crea una segunda."""

    def test_no_crea_fila_duplicada_para_la_clase(self):
        run_a = _make_data([
            _test_record("test_zzz_metodo_nuevo.py", "TestZzzMetodoNuevo", "test_a"),
            _test_record("test_zzz_metodo_nuevo.py", "TestZzzMetodoNuevo", "test_b"),
        ])
        r1 = quality_sync.sync(run_a, PROJECT_ROOT, calidad_dir=self.tmpdir)
        self.assertTrue(r1["ok"], r1)

        casos_ws = self._casos_ws()
        row_after_a = self._find_row_by_identity_in_comentarios(
            casos_ws, "test_zzz_metodo_nuevo.py", "TestZzzMetodoNuevo")
        self.assertIsNotNone(row_after_a)
        id_range_after_a = casos_ws.cell(row=row_after_a, column=1).value

        run_b = _make_data([
            _test_record("test_zzz_metodo_nuevo.py", "TestZzzMetodoNuevo", "test_a"),
            _test_record("test_zzz_metodo_nuevo.py", "TestZzzMetodoNuevo", "test_b"),
            _test_record("test_zzz_metodo_nuevo.py", "TestZzzMetodoNuevo", "test_c"),
        ])
        r2 = quality_sync.sync(run_b, PROJECT_ROOT, calidad_dir=self.tmpdir)
        self.assertTrue(r2["ok"], r2)

        casos_ws_2 = self._casos_ws()
        filas_de_esta_clase = [
            r for r in range(2, casos_ws_2.max_row + 1)
            if self._find_row_by_identity_in_comentarios(
                casos_ws_2, "test_zzz_metodo_nuevo.py", "TestZzzMetodoNuevo") == r
        ]
        self.assertEqual(len(filas_de_esta_clase), 1, "debe seguir habiendo UNA sola fila para la clase")
        id_range_after_b = casos_ws_2.cell(row=filas_de_esta_clase[0], column=1).value
        self.assertNotEqual(id_range_after_a, id_range_after_b,
                             "el rango de IDs debe reflejar el metodo nuevo")


class TestNuevoTestAlfabeticamenteEnMedioNoRenumera(QualitySyncTestCase):
    """Caso 5: insertar un test que cae alfabeticamente en medio de otros
    ya publicados NO debe cambiar el CP de los que ya existian."""

    def test_testb_y_testc_no_cambian_de_cp(self):
        run_1 = _make_data([
            _test_record("test_zzz_alpha.py", "TestA", "test_uno"),
            _test_record("test_zzz_alpha.py", "TestB", "test_uno"),
            _test_record("test_zzz_alpha.py", "TestC", "test_uno"),
        ])
        r1 = quality_sync.sync(run_1, PROJECT_ROOT, calidad_dir=self.tmpdir)
        ids1 = r1["documents"]["Reporte de Ejecución de Pruebas.xlsx"]["result"]["id_by_identity"]
        cp_b_antes = ids1[("test_zzz_alpha.py", "TestB", "test_uno")]
        cp_c_antes = ids1[("test_zzz_alpha.py", "TestC", "test_uno")]

        run_2 = _make_data([
            _test_record("test_zzz_alpha.py", "TestA", "test_uno"),
            _test_record("test_zzz_alpha.py", "TestAB", "test_uno"),  # cae entre TestA y TestB
            _test_record("test_zzz_alpha.py", "TestB", "test_uno"),
            _test_record("test_zzz_alpha.py", "TestC", "test_uno"),
        ])
        r2 = quality_sync.sync(run_2, PROJECT_ROOT, calidad_dir=self.tmpdir)
        ids2 = r2["documents"]["Reporte de Ejecución de Pruebas.xlsx"]["result"]["id_by_identity"]

        self.assertEqual(ids2[("test_zzz_alpha.py", "TestB", "test_uno")], cp_b_antes,
                          "TestB no debe cambiar de CP")
        self.assertEqual(ids2[("test_zzz_alpha.py", "TestC", "test_uno")], cp_c_antes,
                          "TestC no debe cambiar de CP")
        cp_ab = ids2[("test_zzz_alpha.py", "TestAB", "test_uno")]
        self.assertNotIn(cp_ab, (cp_b_antes, cp_c_antes))
        # el nuevo CP debe ser el siguiente libre (al final), no uno insertado
        # entre el de TestA y el de TestB.
        cp_a = ids2[("test_zzz_alpha.py", "TestA", "test_uno")]
        numero_ab = int(cp_ab.split("-")[1])
        numero_a = int(cp_a.split("-")[1])
        self.assertGreater(numero_ab, numero_a)
        self.assertGreater(numero_ab, int(cp_c_antes.split("-")[1]))


class TestFailNoRompeSincronizacion(QualitySyncTestCase):
    """Caso 6: un FAIL debe reflejarse correctamente sin romper sync()."""

    def test_fail_se_registra_y_no_rompe(self):
        t = _test_record("test_zzz_fail.py", "TestZzzFail", "test_que_falla", status="FAIL")
        data = _make_data([t])
        report = quality_sync.sync(data, PROJECT_ROOT, calidad_dir=self.tmpdir)
        self.assertTrue(report["ok"], report)

        casos_ws = self._casos_ws()
        row = self._find_row_by_identity_in_comentarios(casos_ws, "test_zzz_fail.py", "TestZzzFail")
        self.assertIsNotNone(row)
        self.assertEqual(casos_ws.cell(row=row, column=8).value, "FAIL")


class TestErrorConocidoNoDuplicaDefecto(QualitySyncTestCase):
    """Caso 7: un ERROR que corresponde a BUG-001 (ya documentado) no debe
    generar un aviso de defecto nuevo, y Reporte de Defectos.docx no se
    modifica."""

    def test_error_ya_documentado_como_bug_001(self):
        defectos_path = os.path.join(self.tmpdir, "Reporte de Defectos.docx")
        texto_antes = "\n".join(p.text for p in docx.Document(defectos_path).paragraphs)

        t = _test_record(
            "test_historial.py", "TestVaciarPapeleraConBackupYMultiplesHojas",
            "test_archivo_queda_valido_y_se_puede_recargar", status="ERROR",
        )
        data = _make_data([t])
        report = quality_sync.sync(data, PROJECT_ROOT, calidad_dir=self.tmpdir)
        self.assertTrue(report["ok"], report)

        eventos = "\n".join(report["events"])
        self.assertIn("ERROR conocido: BUG-001", eventos)
        self.assertNotIn("NUEVO FAIL/ERROR DETECTADO", eventos)

        texto_despues = "\n".join(p.text for p in docx.Document(defectos_path).paragraphs)
        self.assertEqual(texto_antes, texto_despues, "Reporte de Defectos.docx no debe modificarse")


class TestErrorDesconocidoQuedaComoAviso(QualitySyncTestCase):
    """Caso 8: un ERROR/FAIL que no coincide con ningun BUG existente solo
    genera un aviso [QUALITY]; nunca se escribe en el docx."""

    def test_error_desconocido_solo_genera_aviso(self):
        defectos_path = os.path.join(self.tmpdir, "Reporte de Defectos.docx")
        texto_antes = "\n".join(p.text for p in docx.Document(defectos_path).paragraphs)

        t = _test_record("test_zzz_defecto_nuevo.py", "TestZzzDefectoNuevo",
                          "test_algo_totalmente_nuevo_sin_precedente", status="ERROR")
        data = _make_data([t])
        report = quality_sync.sync(data, PROJECT_ROOT, calidad_dir=self.tmpdir)
        self.assertTrue(report["ok"], report)

        eventos = "\n".join(report["events"])
        self.assertIn("NUEVO FAIL/ERROR DETECTADO", eventos)
        self.assertIn("test_zzz_defecto_nuevo.py", eventos)

        texto_despues = "\n".join(p.text for p in docx.Document(defectos_path).paragraphs)
        self.assertEqual(texto_antes, texto_despues, "Reporte de Defectos.docx no debe modificarse")


class TestErrorDeEscrituraNoRompeElResto(QualitySyncTestCase):
    """Caso 9: si un documento no se puede escribir (archivo corrupto), los
    demas documentos deben seguir sincronizandose y el resultado debe
    reportar el fallo con claridad, sin lanzar excepcion hacia afuera."""

    def test_documento_corrupto_no_bloquea_a_los_demas(self):
        casos_path = os.path.join(self.tmpdir, "Casos de Prueba.xlsx")
        with open(casos_path, "wb") as f:
            f.write(b"esto no es un archivo xlsx valido")

        t = _test_record("test_comercial.py", "TestComercial", "test_01_carga_datos_falsos")
        data = _make_data([t])
        report = quality_sync.sync(data, PROJECT_ROOT, calidad_dir=self.tmpdir)

        self.assertFalse(report["ok"])
        self.assertFalse(report["documents"]["Casos de Prueba.xlsx"]["ok"])
        self.assertIn("error", report["documents"]["Casos de Prueba.xlsx"])
        # el resto de los documentos SI debieron sincronizarse:
        self.assertTrue(report["documents"]["Reporte de Ejecución de Pruebas.xlsx"]["ok"])
        self.assertTrue(report["documents"]["Plan de Pruebas.docx"]["ok"])
        self.assertTrue(report["documents"]["Plan de Gestión de Cambios.docx"]["ok"])
        self.assertTrue(report["documents"]["Reporte de Defectos.docx (solo lectura)"]["ok"])

    def test_calidad_dir_inexistente_no_lanza_excepcion(self):
        data = _make_data([_test_record("test_zzz_x.py", "TestZzzX", "test_uno")])
        report = quality_sync.sync(data, PROJECT_ROOT,
                                    calidad_dir=os.path.join(self.tmpdir, "no_existe"))
        self.assertFalse(report["ok"])
        self.assertIn("Reporte de Ejecución de Pruebas.xlsx", report["documents"])


class TestVolskipQualitySyncOmite(QualitySyncTestCase):
    """El escape hatch VOLVO_SKIP_QUALITY_SYNC debe omitir todo sin tocar
    ningun archivo."""

    def test_skip_omite_todo(self):
        antes = {}
        for name in CALIDAD_FILES:
            with open(os.path.join(self.tmpdir, name), "rb") as f:
                antes[name] = f.read()

        os.environ["VOLVO_SKIP_QUALITY_SYNC"] = "1"
        try:
            data = _make_data([_test_record("test_zzz_skip.py", "TestZzzSkip", "test_uno")])
            report = quality_sync.sync(data, PROJECT_ROOT, calidad_dir=self.tmpdir)
        finally:
            del os.environ["VOLVO_SKIP_QUALITY_SYNC"]

        self.assertTrue(report["skipped"])
        self.assertTrue(report["ok"])
        for name in CALIDAD_FILES:
            with open(os.path.join(self.tmpdir, name), "rb") as f:
                self.assertEqual(antes[name], f.read(), f"{name} no debe modificarse con el skip activo")


if __name__ == "__main__":
    unittest.main(verbosity=2)

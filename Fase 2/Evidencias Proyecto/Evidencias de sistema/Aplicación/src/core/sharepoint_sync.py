import os
import glob
import time
import shutil
import pandas as pd
import openpyxl
import webbrowser
from datetime import date
from openpyxl.formula.translate import Translator
from PyQt6.QtCore import QThread, pyqtSignal, QStandardPaths
from config.settings import PROVISION_EXCEL_PATH
from src.core.excel_writer import ExcelWriter

class SharePointSyncWorker(QThread):
    log_signal = pyqtSignal(str)          
    preview_signal = pyqtSignal(object)   
    success_signal = pyqtSignal(str)      
    error_signal = pyqtSignal(str)        

    def __init__(self, sharepoint_url):
        super().__init__()
        self.sharepoint_url = sharepoint_url
        self.download_dir = QStandardPaths.writableLocation(QStandardPaths.StandardLocation.DownloadLocation)
        self.temp_dir = os.path.join(os.getcwd(), 'data', 'temp')
        self.new_data_df = None
        self.is_injecting = False 

    def run(self):
        if not self.is_injecting:
            self._fase_1_extraccion_y_transformacion()
        else:
            self._fase_2_inyeccion()

    def _fase_1_extraccion_y_transformacion(self):
        try:
            if not os.path.exists(self.temp_dir):
                self.log_signal.emit("[SISTEMA] Carpeta 'temp' no encontrada. Creando directorio...")
                os.makedirs(self.temp_dir)
            
            self._limpiar_archivos_viejos(self.download_dir)
            self._limpiar_archivos_viejos(self.temp_dir)

            self.log_signal.emit("[SISTEMA] Solicitando descarga directa desde SharePoint...")
            webbrowser.open(self.sharepoint_url)
            
            self.log_signal.emit("[SISTEMA] Atrapando descarga...")
            
            downloaded_file = None
            timeout = 60
            elapsed = 0
            
            while elapsed < timeout:
                search_pattern = os.path.join(self.download_dir, "Provision Contratos-RpySS*.xlsx")
                files = glob.glob(search_pattern)
                valid_files = [f for f in files if not f.endswith(('.crdownload', '.part', '.tmp'))]
                
                if valid_files:
                    downloaded_file = max(valid_files, key=os.path.getctime)
                    time.sleep(2) 
                    break
                
                time.sleep(1)
                elapsed += 1
                if elapsed % 5 == 0:
                    self.log_signal.emit(f"[ESPERANDO] Buscando archivo en Descargas... ({elapsed}s/60s)")

            if not downloaded_file:
                raise Exception("Tiempo de espera agotado (60 segundos).\n\nPosibles causas:\n1. Tu navegador te está pidiendo iniciar sesión en Microsoft 365.\n2. Tu internet está un poco lento.\n\nPor favor, revisa tu navegador web, asegúrate de tener la sesión iniciada y vuelve a intentarlo.")

            self.log_signal.emit(f"[EXITO] Archivo atrapado: {os.path.basename(downloaded_file)}")

            temp_file_path = os.path.join(self.temp_dir, "temp_sharepoint.xlsx")
            shutil.copy(downloaded_file, temp_file_path)

            hoja = self._hoja_del_anio(temp_file_path)
            self.log_signal.emit(f"[PROCESANDO] Leyendo hoja '{hoja}' y adaptando formato...")
            df_sp = pd.read_excel(temp_file_path, sheet_name=hoja)
            
            required_cols = ['AÑO', 'MES', 'CHASIS', 'CLIENTE', 'Ejecutivo', 'OBSERVACIONES', 'Dynafleet']
            for col in required_cols:
                if col not in df_sp.columns:
                    raise Exception(f"Falta la columna '{col}' en el archivo de SharePoint.")

            # Traductor de fechas
            meses_map = {
                '1': 'Enero', '2': 'Febrero', '3': 'Marzo', '4': 'Abril', '5': 'Mayo', '6': 'Junio',
                '7': 'Julio', '8': 'Agosto', '9': 'Septiembre', '10': 'Octubre', '11': 'Noviembre', '12': 'Diciembre',
                '01': 'Enero', '02': 'Febrero', '03': 'Marzo', '04': 'Abril', '05': 'Mayo', '06': 'Junio',
                '07': 'Julio', '08': 'Agosto', '09': 'Septiembre'
            }

            def formatear_fecha(row):
                año = str(row['AÑO']).replace('.0', '').strip()
                mes_num_o_texto = str(row['MES']).strip()
                mes_texto = meses_map.get(mes_num_o_texto, mes_num_o_texto.capitalize())
                return f"{año} - {mes_texto}"

            df_sp['AÑO - MES'] = df_sp.apply(formatear_fecha, axis=1)
            
            df_sp.rename(columns={'CHASIS': 'VIN', 'Ejecutivo': 'EJECUTIVO', 'Dynafleet': 'DYNAFLEET'}, inplace=True)
            df_sp['EJECUTIVO'] = df_sp['EJECUTIVO'].astype(str).str.upper()
            
            cols_finales = ['AÑO - MES', 'VIN', 'CLIENTE', 'EJECUTIVO', 'OBSERVACIONES', 'DYNAFLEET']
            df_sp = df_sp[cols_finales].copy()
            df_sp.dropna(subset=['VIN'], inplace=True)
            
            # Limpiamos espacios en blanco del VIN para evitar fallos de lectura
            df_sp['VIN'] = df_sp['VIN'].astype(str).str.strip()

            # --- NUEVA LÓGICA DE DEDUPLICACIÓN (VIN + MES) ---
            self.log_signal.emit("[COMPARANDO] Filtrando contratos por VIN y Mes...")
            
            # 1. Traemos la columna VIN y el Mes del Excel local
            df_local = pd.read_excel(PROVISION_EXCEL_PATH, sheet_name="Prov", usecols=['AÑO - MES', 'VIN'])
            
            # 2. Creamos una "Llave Única" fusionando el VIN y el Mes (ej: "2026 - Julio_93KPZ...")
            df_local['LLAVE_UNICA'] = df_local['AÑO - MES'].astype(str).str.strip() + "_" + df_local['VIN'].astype(str).str.strip()
            df_sp['LLAVE_UNICA'] = df_sp['AÑO - MES'].astype(str).str.strip() + "_" + df_sp['VIN'].astype(str).str.strip()
            
            # 3. Guardamos todas las llaves que ya existen en una lista rápida (set)
            llaves_existentes = set(df_local['LLAVE_UNICA'].dropna())
            
            # 4. Dejamos pasar únicamente las filas cuya llave (Mes + VIN) NO esté en esa lista
            df_nuevos = df_sp[~df_sp['LLAVE_UNICA'].isin(llaves_existentes)].copy()
            
            # 5. Borramos la llave temporal porque ya no la necesitamos
            df_nuevos.drop(columns=['LLAVE_UNICA'], inplace=True)
            # ----------------------------------------------------

            if df_nuevos.empty:
                self.log_signal.emit("[RESULTADO] Búsqueda finalizada: No hay contratos nuevos para este mes.")
                self.success_signal.emit("ZERO")
                return

            self.new_data_df = df_nuevos
            self.log_signal.emit(f"[RESULTADO] ¡Se encontraron {len(df_nuevos)} registros de cobro totalmente nuevos!")
            self.preview_signal.emit(df_nuevos)

        except Exception as e:
            self.error_signal.emit(str(e))

    @staticmethod
    def _hoja_del_anio(ruta):
        """Hoja del año en curso; si aún no existe, la de año más reciente."""
        hojas = pd.ExcelFile(ruta).sheet_names
        anio = str(date.today().year)
        if anio in hojas:
            return anio
        anios = sorted((h for h in hojas if h.strip().isdigit()), key=lambda h: int(h.strip()))
        if anios:
            return anios[-1]
        raise Exception(f"El archivo de SharePoint no tiene una hoja con el año (por ejemplo '{anio}').")

    def _fase_2_inyeccion(self):
        try:
            self.log_signal.emit("[SISTEMA] Creando copia de seguridad previa a la inyección...")
            writer_backup = ExcelWriter()
            writer_backup._crear_backup(PROVISION_EXCEL_PATH)
            if not writer_backup.last_backup_ok:
                self.error_signal.emit("No se pudo crear la copia de seguridad de provisionFelipe.xlsx. La operación fue cancelada y no se realizó ningún cambio.")
                return

            self.log_signal.emit("[INYECCION] Abriendo tu Excel provisionFelipe...")
            wb = openpyxl.load_workbook(PROVISION_EXCEL_PATH)
            ws = wb["Prov"]
            
            header_map = {str(cell.value).strip(): col_idx for col_idx, cell in enumerate(ws[1], start=1) if cell.value}

            vin_col = header_map.get('VIN', 2)
            last_row = 1
            for row in range(2, ws.max_row + 2):
                if ws.cell(row=row, column=vin_col).value is None:
                    last_row = row - 1
                    break

            start_inject_row = last_row + 1
            self.log_signal.emit(f"[INYECCION] Escribiendo nuevos datos desde la fila {start_inject_row}...")

            for _, row_data in self.new_data_df.iterrows():
                for col_name, val in row_data.items():
                    if col_name in header_map:
                        ws.cell(row=start_inject_row, column=header_map[col_name]).value = val
                
                if 'Cantidad de Meses' in header_map: ws.cell(row=start_inject_row, column=header_map['Cantidad de Meses']).value = 0
                if 'Plan' in header_map: ws.cell(row=start_inject_row, column=header_map['Plan']).value = ""
                
                calc_cols = ['Monto Provisionado', 'Valor Plan', 'VAR']
                for calc_col in calc_cols:
                    if calc_col in header_map:
                        col_idx = header_map[calc_col]
                        source_cell = ws.cell(row=start_inject_row - 1, column=col_idx)
                        target_cell = ws.cell(row=start_inject_row, column=col_idx)
                        
                        if source_cell.data_type == 'f':
                            target_cell.value = Translator(source_cell.value, origin=source_cell.coordinate).translate_formula(target_cell.coordinate)

                start_inject_row += 1

            self.log_signal.emit("[GUARDANDO] Actualizando provisionFelipe.xlsx...")
            wb.save(PROVISION_EXCEL_PATH)
            
            self._limpiar_archivos_viejos(self.temp_dir)
            self.success_signal.emit("SUCCESS")

        except Exception as e:
            self.error_signal.emit(f"Error de inyección: {str(e)}")

    def _limpiar_archivos_viejos(self, folder_path):
        if os.path.exists(folder_path):
            for f in glob.glob(os.path.join(folder_path, "Provision Contratos-RpySS*.xlsx")):
                try: os.remove(f)
                except: pass
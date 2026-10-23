from datetime import datetime
import pandas as pd
import qtawesome as qta
import os
from PyQt6.QtWidgets import (QWidget, QVBoxLayout, QHBoxLayout, QFrame,
                             QLabel, QPushButton, QMessageBox, QComboBox, QCheckBox, QListWidget, QListWidgetItem,
                             QStackedWidget, QProgressBar,
                             QDialog, QTextEdit, QLineEdit, QInputDialog, QApplication)
from PyQt6.QtCore import Qt, QThread, pyqtSignal
from src.core.data_manager import DataManager
from src.ui.custom_table import VolvoTableWidget, PandasTableModel
from src.ui.ingresos_base import ModeloTablaIngresos
from src.ui import admin_estilo, admin_listas, comercial_estilo as estilo_com
from src.theme.theme_manager import ThemeManager 
from src.theme.colors import *
from config.settings import PROVISION_EXCEL_PATH, INSTALACIONES_PATH, POWERBI_EXCEL_PATH, CONTACTOS_EXCEL_PATH, RM_JAVI_PATH
from src.core.excel_writer import ExcelWriter
from src.services import powerbi_cache_service, provisiones_services, comercial_services
from src.ui.guardado_overlay import mostrar_guardando, terminar_guardando

# =========================================================
# WORKER: El encargado de cruzar Provisiones + Comercial + PowerBI + Contactos
# =========================================================
def _formatear_meses(valor):
    """Meses sin decimales ("60.0" -> "60"). Vacío si no hay dato o es 0
    (registro de Provisiones aún sin cantidad de meses); '-' y MENSUAL se
    conservan."""
    texto = str(valor).strip()
    if texto.lower() in ('', 'nan', 'none', '<na>'):
        return ""
    if texto.upper() in ('-', 'MENSUAL'):
        return texto.upper()
    try:
        numero = float(texto.replace(',', '.'))
    except ValueError:
        return texto
    return str(int(round(numero))) if numero > 0 else ""


def _elegir_contacto(contactos):
    """(nombre, correo, teléfono) de UN solo administrador de flota del CDB: el
    más reciente (Fecha de extracción) entre los que traen nombre, correo y
    teléfono; si ninguno está completo, el que tenga más datos y, a igualdad,
    el más reciente."""
    if contactos is None or contactos.empty:
        return "", "", ""

    def limpio(valor, telefono=False):
        texto = str(valor).strip()
        if telefono:
            texto = texto.split('.')[0]
        return "" if texto.lower() in ('nan', 'none', '<na>', '-', '0', '') else texto

    candidatos = []
    for _, fila in contactos.iterrows():
        nombre = " ".join(p for p in (limpio(fila.get('Nombre', '')), limpio(fila.get('Apellidos', ''))) if p)
        correo = limpio(fila.get('Correo electrónico', ''))
        telefono = limpio(fila.get('Número móvil', ''), telefono=True)
        fecha = pd.to_datetime(fila.get('Fecha de extracción'), errors='coerce')
        completos = sum(bool(v) for v in (nombre, correo, telefono))
        candidatos.append((completos, fecha if pd.notna(fecha) else pd.Timestamp.min, nombre, correo, telefono))
    _c, _f, nombre, correo, telefono = max(candidatos, key=lambda c: (c[0], c[1]))
    return nombre, correo, telefono


def _meses_sin_decimales(df):
    """Columna MESES de una hoja leída de Excel: 60.0 -> '60', vacío si no hay dato."""
    if df is not None and 'MESES' in df.columns:
        df = df.copy()
        df['MESES'] = df['MESES'].map(_formatear_meses)
    return df


def _mostrar_velo(widget, mensaje):
    # El guardado corre en el hilo principal: se pinta el velo antes de bloquear
    velo = mostrar_guardando(widget, mensaje)
    velo.repaint()
    QApplication.processEvents()


class InstalacionesSyncWorker(QThread):
    log_signal = pyqtSignal(str)          
    preview_signal = pyqtSignal(object, dict) 
    error_signal = pyqtSignal(str)        

    def __init__(self, df_comercial, vines_objetivo=None, vines_extra=None, usar_provision=True):
        super().__init__()
        self.df_comercial = df_comercial
        self.vines_objetivo = {str(v).strip().upper() for v in vines_objetivo} if vines_objetivo else None
        # VIN del RM que se suman a los de la provisión del último mes (o van
        # solos si `usar_provision` es False)
        self.vines_extra = {str(v).strip().upper() for v in vines_extra} if vines_extra else set()
        self.usar_provision = usar_provision
        self.vines_sin_provision = set()
        self.df_cruzado = None

    def _leer_rm_excel(self, vines, limpiar):
        """VIN -> {RUT, CDB, RAZÓN SOCIAL} del RM de Excel, solo para los VIN
        de Comercial que no traen RUT o CDB. Vacío si no hay ruta o no se lee."""
        faltantes = set()
        comercial = self.df_comercial
        for vin, rut, cdb in zip(comercial['VIN_CLEAN'], comercial.get('RUT', ''), comercial.get('CDB', '')):
            if vin in vines and not (limpiar(rut) and limpiar(cdb)):
                faltantes.add(vin)
        sin_comercial = vines - set(comercial['VIN_CLEAN'])
        faltantes |= sin_comercial
        if not faltantes or not RM_JAVI_PATH or not os.path.exists(RM_JAVI_PATH):
            return {}
        try:
            self.log_signal.emit("[LEYENDO] Completando RUT y CDB desde el RM de Excel...")
            df = pd.read_excel(RM_JAVI_PATH, sheet_name="BBDD", usecols=['VIN', 'RUT', 'CDB', 'RAZÓN SOCIAL'], dtype=str)
        except Exception as e:
            self.log_signal.emit(f"[AVISO] No se pudo leer el RM de Excel: {e}")
            return {}
        df['VIN'] = df['VIN'].astype(str).str.replace(r'\s+', '', regex=True).str.upper()
        df = df[df['VIN'].isin(faltantes)].drop_duplicates('VIN', keep='last')
        return {fila['VIN']: {'RUT': limpiar(fila['RUT']), 'CDB': limpiar(fila['CDB']).replace('.0', ''),
                              'RAZÓN SOCIAL': limpiar(fila['RAZÓN SOCIAL'])} for _, fila in df.iterrows()}

    def run(self):
        try:
            self.log_signal.emit("[SISTEMA] Iniciando creación de Instalaciones...")
            
            self.log_signal.emit("[LEYENDO] Extrayendo VINs nuevos de Provisiones...")
            df_prov = pd.read_excel(PROVISION_EXCEL_PATH, sheet_name="Prov", usecols=['VIN', 'Cantidad de Meses', 'AÑO - MES', 'CLIENTE', 'OBSERVACIONES'])
            df_prov.dropna(subset=['VIN'], inplace=True)
            df_prov['VIN'] = df_prov['VIN'].astype(str).str.replace(r'\s+', '', regex=True).str.upper()
            
            meses_map = {
                'Enero': 1, 'Febrero': 2, 'Marzo': 3, 'Abril': 4, 'Mayo': 5, 'Junio': 6,
                'Julio': 7, 'Agosto': 8, 'Septiembre': 9, 'Octubre': 10, 'Noviembre': 11, 'Diciembre': 12
            }
            df_prov['AÑO - MES'] = df_prov['AÑO - MES'].astype(str).str.strip()
            df_prov_todo = df_prov.copy()
            fechas_unicas = [f for f in df_prov['AÑO - MES'].unique() if "-" in f]

            def _orden_periodo(f):
                partes = str(f).split("-")
                try:
                    return (int(partes[0].strip()), meses_map.get(partes[1].strip().capitalize(), 0))
                except (ValueError, IndexError):
                    return (0, 0)

            if self.vines_objetivo:
                # VIN concretos: se usa su provisión más reciente, sin filtrar por mes
                df_prov = df_prov[df_prov['VIN'].isin(self.vines_objetivo)].copy()
                df_prov['_orden'] = df_prov['AÑO - MES'].map(_orden_periodo)
                df_prov = df_prov.sort_values('_orden').drop_duplicates('VIN', keep='last').drop(columns='_orden')
                self.log_signal.emit(f"[FILTRO] Provisión más reciente de {len(df_prov)} VIN solicitados")
            elif fechas_unicas:
                fechas_ordenadas = sorted(
                    fechas_unicas, 
                    key=lambda f: (
                        int(f.split("-")[0].strip()), 
                        meses_map.get(f.split("-")[1].strip().capitalize(), 0)
                    )
                )
                mes_objetivo = fechas_ordenadas[-1]
                self.log_signal.emit(f"[FILTRO] Cruzando exclusivamente con los VIN de la Provisión: {mes_objetivo}")
                df_prov = df_prov[df_prov['AÑO - MES'] == mes_objetivo].copy()
            else:
                self.log_signal.emit("[ALERTA] No se detectó formato de fecha en Provisiones. Procesando todos los VINs.")
                
            if not self.usar_provision and not self.vines_objetivo:
                df_prov = df_prov.iloc[0:0].copy()

            if self.vines_extra and not self.vines_objetivo:
                nuevos = self.vines_extra - set(df_prov['VIN'])
                if nuevos:
                    previos = df_prov_todo[df_prov_todo['VIN'].isin(nuevos)].copy()
                    previos['_orden'] = previos['AÑO - MES'].map(_orden_periodo)
                    previos = previos.sort_values('_orden').drop_duplicates('VIN', keep='last').drop(columns='_orden')
                    self.vines_sin_provision = nuevos - set(previos['VIN'])
                    sin_prov = pd.DataFrame({'VIN': sorted(self.vines_sin_provision)})
                    for col in ('Cantidad de Meses', 'AÑO - MES', 'CLIENTE', 'OBSERVACIONES'):
                        sin_prov[col] = ""
                    df_prov = pd.concat([df_prov, previos, sin_prov], ignore_index=True)
                    self.log_signal.emit(f"[FILTRO] {len(nuevos)} VIN del RM sumados"
                                         + (f" ({len(self.vines_sin_provision)} sin provisión: sin meses)" if self.vines_sin_provision else ""))

            self.log_signal.emit(f"[INFO] Total VINs a procesar: {len(df_prov)}")

            self.df_comercial['VIN_CLEAN'] = self.df_comercial['VIN'].astype(str).str.replace(r'\s+', '', regex=True).str.upper()

            self.log_signal.emit("[LEYENDO] Conectando con BBDD CLIENTE CM.xlsx (David, cache CSV interno)...")
            df_bbdd = pd.DataFrame()
            if POWERBI_EXCEL_PATH and os.path.exists(POWERBI_EXCEL_PATH):
                df_bbdd = powerbi_cache_service.leer_bbdd_clientes_cm(POWERBI_EXCEL_PATH)
                if 'CDB' in df_bbdd.columns:
                    df_bbdd['CDB_CLEAN'] = df_bbdd['CDB'].astype(str).str.replace(r'\.0$', '', regex=True).str.strip()
                    df_bbdd = df_bbdd.drop_duplicates(subset=['CDB_CLEAN'])
            else:
                self.log_signal.emit("[AVISO] BBDD CLIENTE CM no configurada o no encontrada.")

            self.log_signal.emit("[LEYENDO] Conectando con Contactos.xlsx (Maxi)...")
            df_contactos = pd.DataFrame()
            if CONTACTOS_EXCEL_PATH and os.path.exists(CONTACTOS_EXCEL_PATH):
                df_contactos = pd.read_excel(CONTACTOS_EXCEL_PATH)
                if 'CDB' in df_contactos.columns:
                    df_contactos['CDB_CLEAN'] = df_contactos['CDB'].astype(str).str.replace(r'\.0$', '', regex=True).str.strip()
            else:
                self.log_signal.emit("[AVISO] Contactos.xlsx no configurado o no encontrado.")

            self.log_signal.emit("[CRUZANDO] Recopilando todos los datos mediante VIN solicitados...")
            
            registros = []
            dict_actualizacion_comercial = {}
            
            hoy = datetime.now()
            fecha_inicio_str = f"01-{hoy.strftime('%m-%Y')}"
            
            from dateutil.relativedelta import relativedelta

            def _limpio(valor):
                v = str(valor).strip()
                return "" if v.lower() in ['nan', 'none', '<na>', '-'] else v

            rm_excel = self._leer_rm_excel(set(df_prov['VIN']), _limpio)

            for _, row_prov in df_prov.iterrows():
                vin = row_prov['VIN']
                meses = _formatear_meses(row_prov['Cantidad de Meses'])
                if not meses:
                    # Provisión sin meses: se toman de la observación solo si habla de
                    # un plan Volvo Connect (Blue/Gold/Silver son de CM y no cuentan)
                    meses = provisiones_services.meses_desde_observacion(row_prov.get('OBSERVACIONES', '')) or ""
                cliente_prov = _limpio(row_prov.get('CLIENTE', ''))

                chasis = ""
                rut = ""
                cdb = ""
                razon_comercial = ""

                com_match = self.df_comercial[self.df_comercial['VIN_CLEAN'] == vin]
                if not com_match.empty:
                    if 'ID' in com_match.columns:
                        com_match = com_match.assign(_ID=pd.to_numeric(com_match['ID'], errors='coerce')).sort_values('_ID', na_position='first')
                    match = com_match.iloc[-1]
                    chasis = str(match.get('CHASIS', '')).strip()
                    rut = _limpio(match.get('RUT', ''))
                    cdb = _limpio(match.get('CDB', '')).replace('.0', '')
                    razon_comercial = _limpio(match.get('RAZÓN SOCIAL', ''))
                else:
                    chasis = vin[-7:] if len(vin) >= 7 else vin

                # Comercial puede ir atrás del RM de Excel (aún sin sincronizar):
                # RUT y CDB se completan desde el RM por VIN
                if not (rut and cdb) and vin in rm_excel:
                    rut = rut or rm_excel[vin]['RUT']
                    cdb = cdb or rm_excel[vin]['CDB']
                    razon_comercial = razon_comercial or rm_excel[vin]['RAZÓN SOCIAL']

                # Sin el VIN en Comercial (o con datos incompletos) se busca la
                # empresa por el cliente de Provisiones, solo si es inequívoca
                if cliente_prov and not (rut and cdb and razon_comercial):
                    candidatos = provisiones_services.buscar_candidatos_por_cliente(self.df_comercial, cliente_prov)
                    combos = {(_limpio(c.get('RAZÓN SOCIAL')), _limpio(c.get('RUT')), _limpio(c.get('CDB')).replace('.0', ''))
                              for c in candidatos if _limpio(c.get('RUT')) or _limpio(c.get('CDB'))}
                    if len(combos) == 1:
                        razon_c, rut_c, cdb_c = next(iter(combos))
                        razon_comercial = razon_comercial or razon_c
                        rut = rut or rut_c
                        cdb = cdb or cdb_c

                razon_social = ""
                direccion = ""
                comuna = ""
                director = {'nombre': '', 'correo': '', 'telefono': ''}

                if cdb and not df_bbdd.empty:
                    bbdd_match = df_bbdd[df_bbdd['CDB_CLEAN'] == cdb]
                    if not bbdd_match.empty:
                        b_match = bbdd_match.iloc[0]
                        razon_social = _limpio(b_match.get('Nombre', ''))
                        direccion = str(b_match.get('Dirección', '')).strip()
                        comuna = str(b_match.get('Comuna', '')).strip()
                        rut = rut or _limpio(b_match.get('RUT', ''))
                        director = {'nombre': razon_social,
                                    'correo': _limpio(b_match.get('Correo electrónico', '')),
                                    'telefono': _limpio(str(b_match.get('Teléfono', '')).split('.')[0])}
                razon_social = razon_comercial or razon_social or cliente_prov

                admin_flota = ""
                correo = ""
                telefono = ""

                if cdb and not df_contactos.empty:
                    admin_flota, correo, telefono = _elegir_contacto(df_contactos[df_contactos['CDB_CLEAN'] == cdb])

                # Lo que el administrador de flota no trae se completa con el
                # director de empresa (BBDD CLIENTES CM, un registro por CDB)
                if not admin_flota and (director['correo'] or director['telefono']):
                    admin_flota = "Director de empresa"
                correo = correo or director['correo']
                telefono = telefono or director['telefono']

                nombre_contrato = f"{chasis} {razon_social}".strip()
                
                fecha_termino_str = "Pendiente"
                if meses.upper() in ['-', 'MENSUAL']:
                    fecha_termino_str = "MENSUAL"
                else:
                    try:
                        meses_int = int(float(meses))
                        fecha_ini_dt = pd.to_datetime(fecha_inicio_str, format="%d-%m-%Y")
                        fecha_fin_dt = fecha_ini_dt + relativedelta(months=meses_int)
                        fecha_termino_str = fecha_fin_dt.strftime("%d-%m-%Y")
                    except Exception:
                        pass
                        
                # --- CAMBIO DE NOMBRE DE COLUMNAS AQUÍ ---
                if vin not in self.vines_sin_provision:
                    dict_actualizacion_comercial[vin] = {
                        'VIN': vin,
                        'cambios': {
                            'INICIO': fecha_inicio_str,
                            'TERMINO': fecha_termino_str
                        }
                    }
                
                def _c(val):
                    v = str(val).strip()
                    return "" if v.lower() in ['nan', 'none', '<na>', '0', '0.0'] else v

                registros.append({
                    'RUT': _c(rut),
                    'VIN': _c(vin),
                    'CHASSIS': _c(chasis),
                    'NOMBRE CONTRATO': _c(nombre_contrato),
                    'RAZON SOCIAL': _c(razon_social),
                    'DIRECCIÓN': _c(direccion),
                    'COMUNA': _c(comuna),
                    'ADMINISTRADOR DE FLOTA': _c(admin_flota),
                    'CORREO ADMINISTRADOR': _c(correo),
                    'TELÉFONO ADMINISTRADOR': _c(telefono),
                    'INICIO': fecha_inicio_str,
                    'MESES': meses,
                    'TERMINO': fecha_termino_str
                })
                
            df_resultado = pd.DataFrame(registros)
            
            if df_resultado.empty:
                self.log_signal.emit("[ALERTA] Cruce finalizó con 0 resultados.")
            else:
                df_resultado = df_resultado.drop_duplicates(subset=['VIN'])
                self.log_signal.emit(f"[ÉXITO] Cruce completado. {len(df_resultado)} registros listos para previsualizar.")
            
            self.df_cruzado = df_resultado
            self.preview_signal.emit(df_resultado, dict_actualizacion_comercial)

        except Exception as e:
            self.error_signal.emit(f"Error fatal: {str(e)}")

# =========================================================
# LA VENTANA EMERGENTE DE SINCRONIZACIÓN
# =========================================================
class SyncInstalacionesDialog(QDialog):
    """Generar el mes en tres pasos: elegir los VIN, cruzar los datos (con el
    avance por etapas) y revisar el resultado antes de guardar."""
    ETAPAS = (
        ("Provisiones", "Lee los VIN y los meses del plan"),
        ("Director de empresa", "Dirección, comuna, correo y teléfono por CDB"),
        ("Contactos", "Un administrador de flota por empresa"),
        ("Comercial y RM", "Completa RUT, CDB y razón social"),
    )
    PASOS = ("Origen de los VIN", "Cruce de datos", "Revisión")
    _TAGS = {"SISTEMA": "gris", "INFO": "gris", "LEYENDO": "azul", "FILTRO": "violeta", "CRUZANDO": "violeta",
             "AVISO": "ambar", "ALERTA": "ambar", "ÉXITO": "verde"}
    MAX_RESULTADOS = 60

    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Generar mes de instalaciones")
        self.resize(980, 660)
        self.setStyleSheet(f"QDialog {{ background-color: {estilo_com.FONDO}; }}" + estilo_com.QSS_FUENTE_MODULO)
        self.worker = None
        self.df_para_guardar = None
        self.dict_comercial = None
        self.data_manager = DataManager()
        self.vines_extra = []
        self._rm = None
        self._etapa_actual = -1
        self._avisos = set()
        self._setup_ui()

    # ------------------------------------------------------------ interfaz
    def _setup_ui(self):
        raiz = QVBoxLayout(self)
        raiz.setContentsMargins(16, 16, 16, 16)
        tarjeta = QFrame()
        tarjeta.setObjectName("tarjetaSync")
        tarjeta.setStyleSheet(estilo_com.qss_tarjeta("tarjetaSync"))
        raiz.addWidget(tarjeta)
        lay = QVBoxLayout(tarjeta)
        lay.setContentsMargins(0, 0, 0, 0)
        lay.setSpacing(0)

        cabecera = QFrame()
        cabecera.setStyleSheet(f"QFrame {{ background: transparent; border: none; border-bottom: 1px solid {estilo_com.LINEA}; }}")
        fila = QHBoxLayout(cabecera)
        fila.setContentsMargins(20, 14, 20, 14)
        fila.addWidget(admin_listas.texto("Generar mes de instalaciones", tam=16, negrita=True))
        fila.addStretch()
        self._lbl_pasos = []
        for i, nombre in enumerate(self.PASOS):
            lbl = QLabel()
            lbl.setTextFormat(Qt.TextFormat.RichText)
            self._lbl_pasos.append(lbl)
            fila.addWidget(lbl)
        lay.addWidget(cabecera)

        self.pila = QStackedWidget()
        self.pila.addWidget(self._pagina_origen())
        self.pila.addWidget(self._pagina_cruce())
        self.pila.addWidget(self._pagina_revision())
        lay.addWidget(self.pila, 1)

        pie = QFrame()
        pie.setStyleSheet(f"QFrame {{ background: {estilo_com.SUPERFICIE_2}; border: none; border-top: 1px solid {estilo_com.LINEA}; "
                          "border-bottom-left-radius: 13px; border-bottom-right-radius: 13px; }}")
        fila_pie = QHBoxLayout(pie)
        fila_pie.setContentsMargins(20, 12, 20, 12)
        self.lbl_resumen = admin_listas.texto("", tam=12)
        fila_pie.addWidget(self.lbl_resumen, 1)
        self.btn_cancel = QPushButton("Cancelar")
        ThemeManager.apply_secondary_button(self.btn_cancel)
        self.btn_cancel.clicked.connect(self.close)
        self.btn_back = QPushButton("Volver")
        ThemeManager.apply_secondary_button(self.btn_back)
        self.btn_back.clicked.connect(lambda: self._ir_a(0))
        self.btn_start = QPushButton("Iniciar cruce")
        ThemeManager.apply_primary_button(self.btn_start)
        self.btn_start.clicked.connect(self.start_sync)
        self.btn_inject = QPushButton("Aprobar y guardar")
        ThemeManager.apply_success_button(self.btn_inject)
        self.btn_inject.clicked.connect(self.save_data)
        for b in (self.btn_cancel, self.btn_back, self.btn_start, self.btn_inject):
            fila_pie.addWidget(b)
        lay.addWidget(pie)
        self._ir_a(0)

    def _caja(self, titulo):
        caja = QFrame()
        caja.setObjectName("cajaSync")
        caja.setStyleSheet(f"QFrame#cajaSync {{ background: {estilo_com.SUPERFICIE_2}; border: 1px solid {estilo_com.LINEA}; border-radius: 12px; }}")
        lay = QVBoxLayout(caja)
        lay.setContentsMargins(14, 12, 14, 14)
        lay.setSpacing(10)
        lay.addWidget(admin_listas.texto(titulo, tam=13, negrita=True))
        return caja, lay

    def _pagina_origen(self):
        pagina = QWidget()
        lay = QHBoxLayout(pagina)
        lay.setContentsMargins(20, 18, 20, 18)
        lay.setSpacing(14)

        caja_prov, l1 = self._caja("Provisiones del último mes")
        self.chk_prov = QCheckBox("Incluir los VIN de la provisión más reciente")
        self.chk_prov.setChecked(True)
        self.chk_prov.setStyleSheet(f"QCheckBox {{ color: {estilo_com.TINTA}; font-weight: 700; background: transparent; border: none; }}")
        self.chk_prov.toggled.connect(self._actualizar_resumen_origen)
        l1.addWidget(self.chk_prov)
        descripcion = admin_listas.texto("Es lo que hace hoy «Generar mes»: toma el último mes de la hoja Prov.",
                                         color=estilo_com.APAGADO, tam=12)
        descripcion.setWordWrap(True)
        l1.addWidget(descripcion)
        l1.addStretch()
        lay.addWidget(caja_prov, 2)

        caja_rm, l2 = self._caja("Sumar VIN del RM")
        self.txt_buscar = QLineEdit()
        self.txt_buscar.setPlaceholderText("Buscar por VIN, chasis, razón social o RUT")
        self.txt_buscar.textChanged.connect(self._buscar_vin)
        l2.addWidget(self.txt_buscar)
        estilo_lista = (f"QListWidget {{ background: {estilo_com.SUPERFICIE}; border: 1px solid {estilo_com.LINEA}; border-radius: 10px; outline: none; }}"
                        f"QListWidget::item {{ padding: 6px 10px; border-bottom: 1px solid {estilo_com.LINEA_SUAVE}; color: {estilo_com.TINTA}; }}"
                        f"QListWidget::item:hover {{ background: {estilo_com.PRIMARIO_SUAVE}; }}")
        self.lista_rm = QListWidget()
        self.lista_rm.setStyleSheet(estilo_lista)
        self.lista_rm.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        self.lista_rm.setTextElideMode(Qt.TextElideMode.ElideRight)
        self.lista_rm.itemClicked.connect(self._agregar_vin)
        l2.addWidget(self.lista_rm, 1)
        fila_sel = QHBoxLayout()
        self.lbl_agregados = admin_listas.texto("Ningún VIN del RM agregado.", color=estilo_com.APAGADO, tam=12, negrita=True)
        fila_sel.addWidget(self.lbl_agregados, 1)
        btn_quitar = QPushButton("Quitar seleccionado")
        ThemeManager.apply_secondary_button(btn_quitar)
        btn_quitar.clicked.connect(self._quitar_vin)
        fila_sel.addWidget(btn_quitar)
        l2.addLayout(fila_sel)
        self.lista_sel = QListWidget()
        self.lista_sel.setStyleSheet(estilo_lista)
        self.lista_sel.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        self.lista_sel.setTextElideMode(Qt.TextElideMode.ElideRight)
        self.lista_sel.setMaximumHeight(120)
        self.lista_sel.itemDoubleClicked.connect(lambda _i: self._quitar_vin())
        l2.addWidget(self.lista_sel)
        lay.addWidget(caja_rm, 3)
        return pagina

    def _pagina_cruce(self):
        pagina = QWidget()
        lay = QVBoxLayout(pagina)
        lay.setContentsMargins(20, 18, 20, 18)
        lay.setSpacing(10)
        self._filas_etapa = []
        for nombre, detalle in self.ETAPAS:
            marco = QFrame()
            marco.setObjectName("filaEtapa")
            fila = QHBoxLayout(marco)
            fila.setContentsMargins(12, 9, 14, 9)
            icono = QLabel()
            icono.setFixedWidth(20)
            textos = QVBoxLayout()
            textos.setSpacing(0)
            textos.addWidget(admin_listas.texto(nombre, tam=13, negrita=True))
            textos.addWidget(admin_listas.texto(detalle, color=estilo_com.APAGADO, tam=11))
            estado = admin_listas.texto("En espera", color=estilo_com.APAGADO, tam=11, negrita=True)
            fila.addWidget(icono)
            fila.addLayout(textos, 1)
            fila.addWidget(estado)
            lay.addWidget(marco)
            self._filas_etapa.append((marco, icono, estado))
        self.barra_avance = QProgressBar()
        self.barra_avance.setTextVisible(False)
        self.barra_avance.setFixedHeight(6)
        self.barra_avance.setStyleSheet(
            f"QProgressBar {{ background: {estilo_com.LINEA_SUAVE}; border: none; border-radius: 3px; }}"
            f"QProgressBar::chunk {{ background: {estilo_com.AZUL}; border-radius: 3px; }}")
        lay.addWidget(self.barra_avance)
        self.registro = QTextEdit()
        self.registro.setReadOnly(True)
        self.registro.setStyleSheet(
            f"QTextEdit {{ background: {estilo_com.SUPERFICIE}; color: {estilo_com.TINTA}; border: 1px solid {estilo_com.LINEA}; "
            "border-radius: 10px; padding: 8px; font-family: Consolas, 'Courier New', monospace; font-size: 11px; }")
        lay.addWidget(self.registro, 1)
        self._pintar_etapas()
        return pagina

    def _pagina_revision(self):
        pagina = QWidget()
        lay = QVBoxLayout(pagina)
        lay.setContentsMargins(20, 18, 20, 18)
        lay.setSpacing(10)
        self.fila_resumen = QHBoxLayout()
        self.fila_resumen.setSpacing(8)
        lay.addLayout(self.fila_resumen)
        self.preview_table = VolvoTableWidget()
        self.preview_table.menu_type = "general"  # solo lectura: sin "Eliminar Camión"
        self.preview_model = PandasTableModel()
        self.preview_table.setModel(self.preview_model)
        lay.addWidget(self.preview_table, 1)
        return pagina

    # ------------------------------------------------------------ pasos
    def _ir_a(self, paso):
        self.pila.setCurrentIndex(paso)
        for i, lbl in enumerate(self._lbl_pasos):
            hecho, actual = i < paso, i == paso
            color = estilo_com.VERDE if hecho else (estilo_com.TINTA if actual else estilo_com.APAGADO)
            fondo = estilo_com.VERDE if hecho else ("#1B2B4E" if actual else "#E3E8F0")
            marca = "✓" if hecho else str(i + 1)
            lbl.setText(f'<span style="background:{fondo}; color:{"#FFFFFF" if hecho or actual else estilo_com.APAGADO};"> &nbsp;{marca}&nbsp; </span>'
                        f' <span style="color:{color}; font-weight:700;">{self.PASOS[i]}</span>&nbsp;&nbsp;')
        self.btn_start.setVisible(paso == 0)
        self.btn_back.setVisible(paso == 2)
        self.btn_inject.setVisible(paso == 2)
        if paso == 0:
            self._actualizar_resumen_origen()
        elif paso == 1:
            self.lbl_resumen.setText("Cruzando los datos…")
        else:
            self.lbl_resumen.setText("Se guardará como hoja nueva y actualizará INICIO y TERMINO en Comercial.")

    # ------------------------------------------------------- selección de VIN
    def _cargar_rm(self):
        if self._rm is not None:
            return self._rm
        df = self.data_manager.get_comercial_df()
        if df is None or df.empty or 'VIN' not in df.columns:
            self._rm = pd.DataFrame(columns=['VIN', 'CHASIS', 'RAZÓN SOCIAL', 'RUT', '_buscar'])
            return self._rm
        columnas = [c for c in ('VIN', 'CHASIS', 'RAZÓN SOCIAL', 'RUT') if c in df.columns]
        rm = df[columnas].copy().fillna("").astype(str)
        rm['VIN'] = rm['VIN'].str.replace(r'\s+', '', regex=True).str.upper()
        for c in ('CHASIS', 'RAZÓN SOCIAL', 'RUT'):
            if c not in rm.columns:
                rm[c] = ""
        rm = rm[rm['VIN'] != ""].drop_duplicates('VIN')
        rm['_buscar'] = (rm['VIN'] + " " + rm['CHASIS'] + " " + rm['RAZÓN SOCIAL'] + " " + rm['RUT']).str.lower()
        self._rm = rm.reset_index(drop=True)
        return self._rm

    def _buscar_vin(self, texto):
        self.lista_rm.clear()
        texto = texto.strip().lower()
        if len(texto) < 3:
            item = QListWidgetItem("Escribe al menos 3 caracteres para buscar en el RM.")
            item.setFlags(Qt.ItemFlag.NoItemFlags)
            self.lista_rm.addItem(item)
            return
        rm = self._cargar_rm()
        coincidencias = rm[rm['_buscar'].str.contains(texto, regex=False)]
        for _, f in coincidencias.head(self.MAX_RESULTADOS).iterrows():
            agregado = f['VIN'] in self.vines_extra
            item = QListWidgetItem(f"{'✓ ' if agregado else ''}{f['VIN']}   ·   {f['CHASIS']}   ·   {f['RAZÓN SOCIAL']}")
            item.setData(Qt.ItemDataRole.UserRole, f['VIN'])
            if agregado:
                item.setFlags(Qt.ItemFlag.NoItemFlags)
            self.lista_rm.addItem(item)
        if coincidencias.empty:
            item = QListWidgetItem("Ningún VIN del RM coincide.")
            item.setFlags(Qt.ItemFlag.NoItemFlags)
            self.lista_rm.addItem(item)
        elif len(coincidencias) > self.MAX_RESULTADOS:
            item = QListWidgetItem(f"Hay {len(coincidencias)} coincidencias: se muestran {self.MAX_RESULTADOS}. Afina la búsqueda.")
            item.setFlags(Qt.ItemFlag.NoItemFlags)
            self.lista_rm.addItem(item)

    def _agregar_vin(self, item):
        vin = item.data(Qt.ItemDataRole.UserRole)
        if not vin or vin in self.vines_extra:
            return
        self.vines_extra.append(vin)
        self._repintar_seleccion()
        self._buscar_vin(self.txt_buscar.text())

    def _quitar_vin(self):
        item = self.lista_sel.currentItem()
        if item is None:
            return
        vin = item.data(Qt.ItemDataRole.UserRole)
        if vin in self.vines_extra:
            self.vines_extra.remove(vin)
        self._repintar_seleccion()
        self._buscar_vin(self.txt_buscar.text())

    def _repintar_seleccion(self):
        self.lista_sel.clear()
        rm = self._cargar_rm().set_index('VIN')
        for vin in self.vines_extra:
            detalle = rm.loc[vin, 'RAZÓN SOCIAL'] if vin in rm.index else ""
            item = QListWidgetItem(f"{vin}   ·   {detalle}")
            item.setData(Qt.ItemDataRole.UserRole, vin)
            self.lista_sel.addItem(item)
        n = len(self.vines_extra)
        self.lbl_agregados.setText(admin_listas.plural(n, "VIN agregado") + " del RM" if n else "Ningún VIN del RM agregado.")
        self._actualizar_resumen_origen()

    def _actualizar_resumen_origen(self, *_):
        if self.pila.currentIndex() != 0:
            return
        partes = []
        if self.chk_prov.isChecked():
            partes.append("VIN de la provisión más reciente")
        if self.vines_extra:
            partes.append(admin_listas.plural(len(self.vines_extra), "VIN") + " del RM")
        self.lbl_resumen.setText(" + ".join(partes) if partes else "Elige al menos un origen de VIN.")
        self.btn_start.setEnabled(bool(partes))

    # ------------------------------------------------------------ cruce
    def _pintar_etapas(self, actual=-1, avisos=(), terminado=False):
        for i, (marco, icono, estado) in enumerate(self._filas_etapa):
            if terminado or i < actual:
                tono, texto, glifo = (("ambar", "Con avisos", "!") if i in avisos else ("verde", "Listo", "✓"))
            elif i == actual:
                tono, texto, glifo = "azul", "En curso", "●"
            else:
                tono, texto, glifo = "gris", "En espera", "○"
            fondo, color = admin_listas.TONOS[tono]
            icono.setText(glifo)
            icono.setStyleSheet(f"color: {color}; font-size: 14px; font-weight: 800; background: transparent; border: none;")
            estado.setText(texto)
            estado.setStyleSheet(f"color: {color}; font-size: 11px; font-weight: 700; background: transparent; border: none;")
            marco.setStyleSheet(f"QFrame#filaEtapa {{ background: {estilo_com.SUPERFICIE}; border: 1px solid {estilo_com.LINEA}; border-radius: 10px; }}")
        total = len(self._filas_etapa)
        self.barra_avance.setValue(100 if terminado else int(max(actual, 0) / total * 100 + (10 if actual >= 0 else 0)))

    def write_log(self, text):
        etiqueta, _, resto = text.partition("]")
        etiqueta = etiqueta.lstrip("[").strip().upper()
        if resto and etiqueta in self._TAGS:
            fondo, color = admin_listas.TONOS[self._TAGS[etiqueta]]
            self.registro.append(
                f'<span style="background:{fondo}; color:{color}; font-weight:700;">&nbsp;{etiqueta}&nbsp;</span> {resto.strip()}')
        else:
            self.registro.append(text)
        self.registro.verticalScrollBar().setValue(self.registro.verticalScrollBar().maximum())

        etapa = None
        if "Extrayendo VINs" in text:
            etapa = 0
        elif "BBDD CLIENTE" in text:
            etapa = 1
        elif "Contactos.xlsx" in text and "Conectando" in text:
            etapa = 2
        elif etiqueta == "CRUZANDO":
            etapa = 3
        if etapa is not None and etapa > self._etapa_actual:
            self._etapa_actual = etapa
        if etiqueta in ("AVISO", "ALERTA") and self._etapa_actual >= 0:
            self._avisos.add(self._etapa_actual)
        self._pintar_etapas(self._etapa_actual, self._avisos, terminado=etiqueta == "ÉXITO")

    def start_sync(self):
        self._etapa_actual = -1
        self._avisos = set()
        self.registro.clear()
        self._pintar_etapas()
        self._ir_a(1)
        self.btn_start.setEnabled(False)

        df_comercial = self.data_manager.get_comercial_df()
        self.worker = InstalacionesSyncWorker(
            df_comercial, vines_extra=self.vines_extra, usar_provision=self.chk_prov.isChecked())
        self.worker.log_signal.connect(self.write_log)
        self.worker.preview_signal.connect(self.show_preview)
        self.worker.error_signal.connect(self.handle_error)
        self.worker.start()

    def show_preview(self, df_resultado, dict_comercial):
        self.df_para_guardar = df_resultado
        self.dict_comercial = dict_comercial
        self.preview_model.set_data(df_resultado)

        while self.fila_resumen.count():
            item = self.fila_resumen.takeAt(0)
            if item.widget() is not None:
                item.widget().deleteLater()
        total = len(df_resultado)
        sin_correo = int((df_resultado['CORREO ADMINISTRADOR'].astype(str).str.strip() == "").sum()) if total and 'CORREO ADMINISTRADOR' in df_resultado.columns else 0
        sin_meses = int((df_resultado['MESES'].astype(str).str.strip() == "").sum()) if total and 'MESES' in df_resultado.columns else 0
        self.fila_resumen.addWidget(admin_listas.pildora(admin_listas.plural(total, "instalación", "instalaciones"), "azul"))
        self.fila_resumen.addWidget(admin_listas.pildora(f"{sin_correo} sin correo", "ambar" if sin_correo else "verde"))
        self.fila_resumen.addWidget(admin_listas.pildora(f"{sin_meses} sin meses", "ambar" if sin_meses else "verde"))
        sin_provision = len(getattr(self.worker, 'vines_sin_provision', ()) or ())
        if sin_provision:
            self.fila_resumen.addWidget(admin_listas.pildora(f"{sin_provision} sin provisión (no se tocan las fechas en Comercial)", "gris"))
        self.fila_resumen.addStretch()
        self._ir_a(2)

    def handle_error(self, error_msg):
        QMessageBox.critical(self, "Error Fatal", error_msg)
        self.btn_start.setEnabled(True)
        self._ir_a(0)

    def save_data(self):
        if self.df_para_guardar is None or self.df_para_guardar.empty:
            QMessageBox.warning(self, "Error", "No hay datos para guardar.")
            return

        nombre_hoja, ok = QInputDialog.getText(self, "Nombre del Mes", "Ingresa el nombre para esta nueva hoja en Instalaciones (ej: 2026 - Julio):")

        if not ok or not nombre_hoja.strip():
            return

        nombre_hoja = nombre_hoja.strip()

        try:
            import openpyxl
            from openpyxl.utils.dataframe import dataframe_to_rows

            if os.path.exists(INSTALACIONES_PATH):
                wb = openpyxl.load_workbook(INSTALACIONES_PATH)
                if nombre_hoja in wb.sheetnames:
                    reply = QMessageBox.question(self, "Sobrescribir", f"La hoja '{nombre_hoja}' ya existe.\n¿Deseas reemplazarla por completo?", QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No)
                    if reply == QMessageBox.StandardButton.No:
                        return

            _mostrar_velo(self, f"Guardando '{nombre_hoja}' en Instalaciones…")
            writer_backup = ExcelWriter()
            writer_backup._crear_backup(INSTALACIONES_PATH)
            if not writer_backup.last_backup_ok:
                terminar_guardando(self, exito=False)
                QMessageBox.critical(self, "Error", "No se pudo crear la copia de seguridad de Instalaciones.xlsx. La operación fue cancelada y no se realizó ningún cambio.")
                return

            if not os.path.exists(INSTALACIONES_PATH):
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = nombre_hoja
            else:
                if nombre_hoja in wb.sheetnames:
                    wb.remove(wb[nombre_hoja])
                ws = wb.create_sheet(nombre_hoja)

            for r in dataframe_to_rows(self.df_para_guardar, index=False, header=True):
                ws.append(r)

            wb.save(INSTALACIONES_PATH)

        # --- INYECCIÓN HACIA COMERCIAL (PruebasFM) ---
            if self.dict_comercial:
                writer = ExcelWriter()
                writer.update_cells(self.dict_comercial, module_name="Auto-Instalaciones")
                self.data_manager.load_data_async()

            mensaje_exito = f"Datos guardados exitosamente en la hoja '{nombre_hoja}' de Instalaciones.\nFechas de facturación actualizadas en Comercial."
            if not writer_backup.last_backup_ok:
                mensaje_exito += "\n\n⚠️ Aviso: no se pudo crear la copia de seguridad previa de Instalaciones.xlsx."
            terminar_guardando(self, texto_exito="Mes guardado")
            if not writer_backup.last_backup_ok:
                QMessageBox.warning(self, "Aviso", mensaje_exito)
            self.accept()

        except PermissionError:
            terminar_guardando(self, exito=False)
            QMessageBox.critical(self, "Error", "El archivo Instalaciones.xlsx está abierto. Ciérralo y vuelve a intentar.")
        except Exception as e:
            terminar_guardando(self, exito=False)
            QMessageBox.critical(self, "Error al guardar", str(e))


# =========================================================
# LA PESTAÑA PRINCIPAL
# =========================================================
class InstalacionesModule(QWidget):
    def __init__(self):
        super().__init__()
        self.data_manager = DataManager()
        self.df_base = pd.DataFrame()

        self.search_input = QLineEdit()
        self.search_input.hide()

        self._setup_ui()
        self.load_local_data()

    def _setup_ui(self):
        main_layout = admin_estilo.raiz_pantalla(self)

        self._barra = admin_estilo.barra_opciones('fa5s.satellite-dish', "Instalaciones", "")
        main_layout.addWidget(self._barra.marco)

        self.combo_mes = estilo_com.ComboTarjeta("Hoja (mes)")
        self.combo_mes.addItem("Sin datos...")
        self.combo_mes.currentTextChanged.connect(self.change_sheet)
        admin_estilo.combo_barra(self.combo_mes, 190)
        self._barra.zona_filtros.addWidget(admin_estilo.etiqueta_barra("Hoja"))
        self._barra.zona_filtros.addWidget(self.combo_mes)

        self.btn_generar = admin_estilo.boton_capsula(self._barra, 'fa5s.cogs', "Generar mes", QPushButton())
        self.btn_generar.clicked.connect(self.open_sync_dialog)
        self.btn_renombrar = admin_estilo.boton_capsula(self._barra, 'fa5s.edit', "Renombrar hoja", QPushButton())
        self.btn_renombrar.clicked.connect(self.rename_current_sheet)
        self.btn_eliminar = admin_estilo.boton_capsula(self._barra, 'fa5s.trash-alt', "Eliminar hoja", QPushButton())
        self.btn_eliminar.clicked.connect(self.delete_current_sheet)
        self.btn_export_excel = admin_estilo.boton_capsula(self._barra, 'fa5s.file-excel', "Exportar a Excel", QPushButton(), grupo='generales')
        self.btn_export_excel.clicked.connect(lambda: self.table.export_to_excel("Instalaciones_Claudia"))

        marco, capas = admin_estilo.tarjeta_contenido("tarjetaInstalaciones")
        self.table = VolvoTableWidget()
        self.table.menu_type = "general"  # Instalaciones no tiene "Eliminar Camión"
        self.table.sort_changed.connect(self.filter_table)
        self.model = ModeloTablaIngresos(alineacion=Qt.AlignmentFlag.AlignCenter | Qt.AlignmentFlag.AlignVCenter)

        df_vacio = pd.DataFrame(columns=[
            'RUT', 'VIN', 'CHASSIS', 'NOMBRE CONTRATO', 'RAZON SOCIAL',
            'DIRECCIÓN', 'COMUNA', 'ADMINISTRADOR DE FLOTA',
            'CORREO ADMINISTRADOR', 'TELÉFONO ADMINISTRADOR', 'INICIO', 'MESES', 'TERMINO'
        ])
        self.model.set_data(df_vacio)
        self.table.set_model_and_delegate(self.model)
        admin_estilo.estilar_tabla(self.table)
        capas.addWidget(self.table, 1)
        main_layout.addWidget(marco, 1)

    _ANCHOS_COLUMNAS = {
        'RUT': 120, 'VIN': 180, 'CHASSIS': 100, 'NOMBRE CONTRATO': 160, 'RAZON SOCIAL': 220, 'DIRECCIÓN': 180,
        'COMUNA': 120, 'ADMINISTRADOR DE FLOTA': 180, 'CORREO ADMINISTRADOR': 220, 'TELÉFONO ADMINISTRADOR': 170,
        'INICIO': 110, 'MESES': 80, 'TERMINO': 110,
    }

    def _actualizar_kpis(self, df):
        """Subtítulo de la barra según las filas visibles (con búsqueda)."""
        total = len(df)
        clientes = int(df['RAZON SOCIAL'].astype(str).str.strip().replace("", pd.NA).dropna().nunique()) \
            if not df.empty and 'RAZON SOCIAL' in df.columns else 0
        hoja = self.combo_mes.currentText()
        if hoja and hoja != "Sin datos...":
            self._barra.lbl_subtitulo.setText(f"{hoja} · {total} instalaciones · {clientes} clientes")
        else:
            self._barra.lbl_subtitulo.setText("Sin datos")

    def rename_current_sheet(self):
        current_sheet = self.combo_mes.currentText()
        if not current_sheet or current_sheet == "Sin datos...": return
        
        new_name, ok = QInputDialog.getText(self, "Renombrar Hoja", "Ingresa el nuevo nombre:", text=current_sheet)
        
        if ok and new_name.strip() and new_name.strip() != current_sheet:
            new_name = new_name.strip()
            try:
                import openpyxl
                wb = openpyxl.load_workbook(INSTALACIONES_PATH)
                if new_name in wb.sheetnames:
                    QMessageBox.warning(self, "Error", "Ya existe una hoja con ese nombre.")
                    return

                _mostrar_velo(self, "Renombrando la hoja…")
                writer_backup = ExcelWriter()
                writer_backup._crear_backup(INSTALACIONES_PATH)
                if not writer_backup.last_backup_ok:
                    terminar_guardando(self, exito=False)
                    QMessageBox.critical(self, "Error", "No se pudo crear la copia de seguridad de Instalaciones.xlsx. La operación fue cancelada y no se realizó ningún cambio.")
                    return

                wb[current_sheet].title = new_name
                wb.save(INSTALACIONES_PATH)

                mensaje_exito = f"La hoja se renombró a '{new_name}'."
                if not writer_backup.last_backup_ok:
                    mensaje_exito += "\n\n⚠️ Aviso: no se pudo crear la copia de seguridad previa."
                terminar_guardando(self, texto_exito="Hoja renombrada")
                if not writer_backup.last_backup_ok:
                    QMessageBox.warning(self, "Aviso", mensaje_exito)
                self.load_local_data()
                self.combo_mes.setCurrentText(new_name)

            except PermissionError:
                terminar_guardando(self, exito=False)
                QMessageBox.critical(self, "Error", "El archivo Instalaciones.xlsx está abierto. Ciérralo.")
            except Exception as e:
                terminar_guardando(self, exito=False)
                QMessageBox.critical(self, "Error", str(e))

    def delete_current_sheet(self):
        current_sheet = self.combo_mes.currentText()
        if not current_sheet or current_sheet == "Sin datos...": return
        
        reply = QMessageBox.question(self, "Eliminar Hoja",
                                     f"¿Estás completamente seguro de eliminar la hoja '{current_sheet}'?\nSe creará una copia de seguridad antes de eliminarla, pero la restauración es manual desde Configuración.",
                                     QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No)

        if reply == QMessageBox.StandardButton.Yes:
            try:
                import openpyxl
                wb = openpyxl.load_workbook(INSTALACIONES_PATH)

                if len(wb.sheetnames) <= 1:
                    QMessageBox.warning(self, "Acción denegada", "No puedes eliminar la única hoja del archivo.\nSi quieres vaciarlo, puedes renombrarla o borrar el archivo desde la carpeta.")
                    return

                _mostrar_velo(self, "Eliminando la hoja…")
                writer_backup = ExcelWriter()
                writer_backup._crear_backup(INSTALACIONES_PATH)
                if not writer_backup.last_backup_ok:
                    terminar_guardando(self, exito=False)
                    QMessageBox.critical(self, "Error", "No se pudo crear la copia de seguridad de Instalaciones.xlsx. La operación fue cancelada y no se realizó ningún cambio.")
                    return

                del wb[current_sheet]
                wb.save(INSTALACIONES_PATH)

                mensaje_exito = f"La hoja '{current_sheet}' fue eliminada."
                if not writer_backup.last_backup_ok:
                    mensaje_exito += "\n\n⚠️ Aviso: no se pudo crear la copia de seguridad previa."
                terminar_guardando(self, texto_exito="Hoja eliminada")
                if not writer_backup.last_backup_ok:
                    QMessageBox.warning(self, "Aviso", mensaje_exito)
                self.load_local_data()

            except PermissionError:
                terminar_guardando(self, exito=False)
                QMessageBox.critical(self, "Error", "El archivo Instalaciones.xlsx está abierto. Ciérralo.")
            except Exception as e:
                terminar_guardando(self, exito=False)
                QMessageBox.critical(self, "Error", str(e))

    def open_sync_dialog(self):
        dialog = SyncInstalacionesDialog(self)
        if dialog.exec() == QDialog.DialogCode.Accepted:
            self.load_local_data()
            
    def load_local_data(self):
        try:
            if not os.path.exists(INSTALACIONES_PATH):
                self.combo_mes.blockSignals(True)
                self.combo_mes.clear()
                self.combo_mes.addItem("Sin datos...")
                self.combo_mes.blockSignals(False)
                self.model.set_data(pd.DataFrame())
                self._actualizar_kpis(pd.DataFrame())
                return

            # Un solo ExcelFile: read_excel volvía a abrir y parsear el libro completo
            with pd.ExcelFile(INSTALACIONES_PATH) as xls:
                hojas = xls.sheet_names

                if not hojas:
                    return

                self.combo_mes.blockSignals(True)
                self.combo_mes.clear()
                self.combo_mes.addItems(hojas)
                self.combo_mes.setCurrentText(hojas[-1])
                self.combo_mes.blockSignals(False)

                df_hoja = xls.parse(hojas[-1])
            self.df_base = _meses_sin_decimales(df_hoja.fillna(""))
            self.filter_table()
            
        except Exception as e:
            print(f"Error cargando datos locales de instalaciones: {e}")

    def change_sheet(self, sheet_name):
        if not sheet_name or sheet_name == "Sin datos..." or not os.path.exists(INSTALACIONES_PATH):
            return
            
        try:
            self.df_base = _meses_sin_decimales(pd.read_excel(INSTALACIONES_PATH, sheet_name=sheet_name).fillna(""))
            self.filter_table()
        except Exception as e:
            print(f"Error cambiando a la hoja {sheet_name}: {e}")

    def filter_table(self):
        if self.df_base.empty:
            self.model.set_data(pd.DataFrame())
            self._actualizar_kpis(pd.DataFrame())
            return

        search_text = self.search_input.text().lower()
        df_display = self.df_base.copy()
        
        if search_text:
            columnas_busqueda = ['VIN', 'CHASSIS', 'RUT', 'NOMBRE CONTRATO', 'RAZON SOCIAL']
            mask = comercial_services.construir_mascara_busqueda(
                df_display, search_text.strip(), columnas_busqueda)
            df_display = df_display[mask]

        # Ordenamiento tipo Excel elegido por el usuario (no toca el Excel de Instalaciones)
        df_display = self.table.apply_sort(df_display)

        self.model.set_data(df_display)
        for col, ancho in self._ANCHOS_COLUMNAS.items():
            if col in df_display.columns:
                self.table.setColumnWidth(df_display.columns.get_loc(col), ancho)
        self._actualizar_kpis(df_display)

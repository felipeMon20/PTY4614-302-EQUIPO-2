import os
import sys
import shutil
from datetime import datetime
import openpyxl
import qtawesome as qta
from PyQt6.QtWidgets import (QWidget, QVBoxLayout, QHBoxLayout, QLabel, 
                             QLineEdit, QPushButton, QFileDialog, QMessageBox, 
                             QFrame, QGridLayout, QScrollArea, QDialog, QTextEdit, QApplication, QStackedWidget)
from urllib.parse import urlparse
from PyQt6.QtCore import Qt, QSize, QProcess, pyqtSignal
from PyQt6.QtGui import QColor
from config.settings import load_config, save_config, EXCEL_PATH
from config import rutas
from config import settings as ajustes
from src.core.data_manager import DataManager
from src.core.excel_writer import ExcelWriter
from src.theme.theme_manager import ThemeManager
from src.theme.colors import *
from src.theme.ui_scale import NIVELES_ESCALA, get_nivel_escala, set_nivel_escala
from src.ui import admin_estilo
from src.ui import comercial_estilo as estilo_com
from src.ui.guia_conexion import GuiaConexionDialog
from src.ui.ayuda_botones import crear_pagina_ayuda
from src.theme.tema import GestorTema, CLARO, OSCURO

VOLVO = "#003057"
VOLVO_HOVER = "#00446F"
VOLVO_SUAVE = "#E4EDF6"
_SIN_FONDO = "border: none; background: transparent;"


class TarjetaOpcion(QFrame):
    """Tarjeta seleccionable de Apariencia: vista en miniatura, título y detalle."""
    clicked = pyqtSignal()

    def __init__(self, titulo, detalle, vista, parent=None):
        super().__init__(parent)
        self.setObjectName("tarjetaOpcion")
        self.setCursor(Qt.CursorShape.PointingHandCursor)
        lay = QVBoxLayout(self)
        lay.setContentsMargins(14, 14, 14, 12)
        self.setMinimumHeight(176)
        lay.setSpacing(8)
        vista.setAttribute(Qt.WidgetAttribute.WA_TransparentForMouseEvents)
        lay.addWidget(vista)
        fila = QHBoxLayout()
        self.lbl_titulo = QLabel(titulo)
        self.lbl_titulo.setStyleSheet(f"font-size: 13px; font-weight: 600; color: {estilo_com.TINTA}; {_SIN_FONDO}")
        self.lbl_actual = QLabel("Actual")
        self.lbl_actual.setStyleSheet(
            f"font-size: 11px; font-weight: 600; color: {VOLVO}; background: {VOLVO_SUAVE}; border-radius: 9px; padding: 1px 9px; border: none;")
        fila.addWidget(self.lbl_titulo)
        fila.addStretch()
        fila.addWidget(self.lbl_actual)
        lay.addLayout(fila)
        lbl_detalle = QLabel(detalle)
        lbl_detalle.setWordWrap(True)
        lbl_detalle.setStyleSheet(f"font-size: 12px; color: {estilo_com.APAGADO}; {_SIN_FONDO}")
        lay.addWidget(lbl_detalle)
        for hijo in (self.lbl_titulo, self.lbl_actual, lbl_detalle):
            hijo.setAttribute(Qt.WidgetAttribute.WA_TransparentForMouseEvents)
        self.marcar(False)

    def marcar(self, activa):
        borde = f"2px solid {VOLVO}" if activa else f"1px solid {estilo_com.LINEA}"
        self.setStyleSheet(f"QFrame#tarjetaOpcion {{ background: {estilo_com.SUPERFICIE}; border: {borde}; border-radius: 12px; }}")
        self.lbl_actual.setVisible(activa)

    def mousePressEvent(self, event):
        if event.button() == Qt.MouseButton.LeftButton:
            self.clicked.emit()
        super().mousePressEvent(event)


class ConfiguracionModule(QWidget):
    _CLAVES_RUTAS = ("EXCEL_PATH", "PROVISION_EXCEL_PATH", "POWERBI_EXCEL_PATH", "CONTACTOS_EXCEL_PATH",
                     "INSTALACIONES_PATH", "CORREOS_EXCEL_PATH", "LOG_CSV_PATH", "DELETED_LOG_EXCEL_PATH",
                     "CASOS_BAJA_CM_PATH", "RM_JAVI_PATH", "EJECUTIVOS_PATH", "METAS_VC_PATH")

    def __init__(self):
        super().__init__()
        self.config = load_config()
        self.data_manager = DataManager()
        self.last_errors = [] 
        
        self.path_inputs = {}
        self.url_inputs = {}
        # Rutas sin campo propio (ej. facturacionManual.xlsx) que la autodetección sí completa
        self._rutas_extra = {k: self.config.get(k, "") for k in ("FACTURACION_MANUAL_PATH",)}
        
        self.data_manager.data_ready.connect(self.update_health_panel)
        self._setup_ui()

    def _setup_ui(self):
        raiz = admin_estilo.raiz_pantalla(self)

        cabecera = QFrame()
        cabecera.setObjectName("cabeceraConfig")
        cabecera.setStyleSheet(
            f"QFrame#cabeceraConfig {{ background: {estilo_com.SUPERFICIE}; border: 1px solid {estilo_com.LINEA}; "
            f"border-top: 3px solid {VOLVO}; border-radius: 14px; }}")
        fila_cab = QHBoxLayout(cabecera)
        fila_cab.setContentsMargins(16, 12, 16, 12)
        fila_cab.setSpacing(14)
        icono_cab = QLabel()
        icono_cab.setFixedSize(40, 40)
        icono_cab.setAlignment(Qt.AlignmentFlag.AlignCenter)
        icono_cab.setPixmap(qta.icon('fa5s.cog', color='white').pixmap(QSize(18, 18)))
        icono_cab.setStyleSheet(f"background: {VOLVO}; border-radius: 11px; border: none;")
        bloque = QVBoxLayout()
        bloque.setSpacing(0)
        titulo_cab = QLabel("Configuración")
        titulo_cab.setStyleSheet(f"font-size: 20px; font-weight: 700; color: {estilo_com.TINTA}; border: none; background: transparent;")
        sub_cab = QLabel("Carpeta de datos, enlaces y apariencia de la aplicación")
        sub_cab.setStyleSheet(f"font-size: 12px; color: {estilo_com.APAGADO}; border: none; background: transparent;")
        bloque.addWidget(titulo_cab)
        bloque.addWidget(sub_cab)
        self.lbl_estado_barra = QLabel("")
        btn_save = self._boton(" Guardar y reiniciar", 'fa5s.save', "primario")
        btn_save.clicked.connect(self._save_changes)
        fila_cab.addWidget(icono_cab)
        fila_cab.addLayout(bloque, 1)
        fila_cab.addWidget(self.lbl_estado_barra)
        fila_cab.addWidget(btn_save)
        raiz.addWidget(cabecera)

        self.card_db = self._create_status_card("Base de datos", "fa5s.database", PRIMARY, "Cargando...")
        self.card_rows = self._create_status_card("Total contratos", "fa5s.truck", "#8B5CF6", "Cargando...")
        self.card_sync = self._create_status_card("Última modificación", "fa5s.clock", WARNING, "Cargando...")
        self.card_ram = self._create_status_card("Memoria RAM", "fa5s.microchip", "#10B981", "Cargando...")
        self.card_load_time = self._create_status_card("Tiempo de lectura", "fa5s.tachometer-alt", "#F59E0B", "Cargando...")
        self.card_logs = self._create_status_card("Estado de logs", "fa5s.heartbeat", TEXT_SECONDARY, "Cargando...")
        tira_estado = QFrame()
        tira_estado.setObjectName("tiraEstadoConfig")
        tira_estado.setStyleSheet(
            f"QFrame#tiraEstadoConfig {{ background: {estilo_com.SUPERFICIE_2}; border: 1px solid {estilo_com.LINEA_SUAVE}; "
            "border-radius: 12px; }")
        fila_kpis = QHBoxLayout(tira_estado)
        fila_kpis.setContentsMargins(10, 6, 10, 6)
        fila_kpis.setSpacing(6)
        for tarjeta in (self.card_db, self.card_rows, self.card_sync,
                        self.card_ram, self.card_load_time, self.card_logs):
            fila_kpis.addWidget(tarjeta['frame'], 1)

        contenedor_tabs, self.tabbar, self._badges = admin_estilo.subpestanas(
            ["Archivos", "Enlaces", "Apariencia", "Avanzado", "Ayuda"])
        for i in (2, 3, 4):
            self._badges[i].setVisible(False)
        admin_estilo.actualizar_contador(self.tabbar, self._badges, 1, 3)
        raiz.addWidget(contenedor_tabs)

        # Las rutas ya no se editan una a una: viven en campos que completa la carpeta de datos
        for config_key in self._CLAVES_RUTAS:
            self.path_inputs[config_key] = QLineEdit(self.config.get(config_key, ""))

        # Contenedores primero: mover una página ya armada a otro contenedor
        # le reaplica los estilos a todos sus widgets, así se hace una sola vez
        self.paginas = QStackedWidget()
        marco, capas = admin_estilo.tarjeta_contenido("tarjetaConfiguracion")
        capas.addWidget(self.paginas)
        raiz.addWidget(marco, 1)
        raiz.addWidget(tira_estado)
        for pagina in (self._crear_grupo_carpeta_datos(), self._crear_pagina_enlaces(),
                       self._crear_grupo_escala(), self._crear_pagina_avanzado()):
            self._con_desplazamiento(pagina, self.paginas)
        # Ayuda es la página más grande (~600 widgets): se arma la primera vez que se abre
        self._area_ayuda = self._con_desplazamiento(None, self.paginas)
        self.tabbar.currentChanged.connect(self._asegurar_pagina_ayuda)
        self.tabbar.currentChanged.connect(self.paginas.setCurrentIndex)

        if not ajustes.MODO_PROGRAMADOR:
            self.tabbar.setTabVisible(3, False)
        for txt in self.path_inputs.values():
            txt.textChanged.connect(lambda *_: self._refrescar_estado_rutas())
        self._refrescar_estado_rutas()
        self.update_health_panel()

    @staticmethod
    def _pagina():
        pagina = QWidget()
        pagina.setObjectName("paginaConfiguracion")
        pagina.setStyleSheet("QWidget#paginaConfiguracion { background: transparent; }")
        lay = QVBoxLayout(pagina)
        lay.setContentsMargins(20, 18, 20, 18)
        lay.setSpacing(14)
        return pagina, lay

    @staticmethod
    def _con_desplazamiento(pagina, pila):
        """Cada pestaña mantiene su altura natural y se desplaza si no cabe."""
        area = QScrollArea()
        area.setWidgetResizable(True)
        area.setFrameShape(QFrame.Shape.NoFrame)
        area.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        area.setStyleSheet(estilo_com.QSS_SCROLL_PANEL)
        area.viewport().setAutoFillBackground(False)
        pila.addWidget(area)
        if pagina is not None:
            area.setWidget(pagina)
        return area

    def _asegurar_pagina_ayuda(self, indice):
        if self.paginas.widget(indice) is self._area_ayuda and self._area_ayuda.widget() is None:
            contenedor = QWidget()
            contenedor.setObjectName("paginaConfiguracion")
            contenedor.setStyleSheet("QWidget#paginaConfiguracion { background: transparent; }")
            lay = QVBoxLayout(contenedor)
            lay.setContentsMargins(20, 18, 20, 0)
            lay.setSpacing(0)
            lay.addWidget(self._crear_grupo_actualizaciones())
            lay.addWidget(crear_pagina_ayuda(self._titulo_pagina))
            self._area_ayuda.setWidget(contenedor)

    def _crear_grupo_actualizaciones(self):
        from src.core.actualizador import ControlActualizacion
        self._control_act = ControlActualizacion.instancia()
        marco, caja = self._marco_lista("listaActualizaciones")
        caja.setContentsMargins(18, 14, 18, 14)
        caja.setSpacing(8)
        self._titulo_pagina(caja, "Actualizaciones", "Versión instalada y última versión publicada por el equipo.")
        self.lbl_act_instalada = QLabel(self._control_act.version_instalada)
        self.lbl_act_publicada = QLabel()
        for etiqueta, valor in (("Versión instalada", self.lbl_act_instalada), ("Última publicada", self.lbl_act_publicada)):
            fila = QHBoxLayout()
            lbl = QLabel(etiqueta)
            lbl.setStyleSheet(f"font-size: 13px; color: {estilo_com.APAGADO}; {_SIN_FONDO}")
            valor.setStyleSheet(f"font-size: 13px; font-weight: 600; color: {estilo_com.TINTA}; {_SIN_FONDO}")
            fila.addWidget(lbl)
            fila.addStretch()
            fila.addWidget(valor)
            caja.addLayout(fila)
        self.lbl_act_notas = QLabel()
        self.lbl_act_notas.setWordWrap(True)
        self.lbl_act_notas.setStyleSheet(f"font-size: 12px; color: {estilo_com.APAGADO}; {_SIN_FONDO}")
        caja.addWidget(self.lbl_act_notas)
        botones = QHBoxLayout()
        self.btn_act_instalar = self._boton(" Actualizar ahora", 'fa5s.download', "primario")
        self.btn_act_instalar.clicked.connect(self._instalar_actualizacion)
        self.btn_act_buscar = self._boton(" Buscar actualizaciones", 'fa5s.sync-alt')
        self.btn_act_buscar.clicked.connect(lambda *_: self._control_act.buscar())
        self.lbl_act_revision = QLabel()
        self.lbl_act_revision.setStyleSheet(f"font-size: 11px; color: {estilo_com.APAGADO}; {_SIN_FONDO}")
        botones.addWidget(self.btn_act_instalar)
        botones.addWidget(self.btn_act_buscar)
        botones.addStretch()
        botones.addWidget(self.lbl_act_revision)
        caja.addLayout(botones)
        self._control_act.cambio.connect(self._refrescar_actualizaciones)
        self._refrescar_actualizaciones()
        return marco

    def _refrescar_actualizaciones(self):
        from src.services.actualizador_service import puede_instalar
        control = self._control_act
        man = control.manifiesto
        disponible = control.disponible()
        if control.revision_en_curso:
            publicada = "Buscando…"
        elif man is None:
            publicada = "No se pudo leer la carpeta de actualizaciones"
        else:
            publicada = f"{man['version']} · " + ("disponible" if disponible else "estás al día")
        self.lbl_act_publicada.setText(publicada)
        self.lbl_act_notas.setText(man["notas"] if disponible and man["notas"] else "")
        self.lbl_act_notas.setVisible(bool(self.lbl_act_notas.text()))
        self.btn_act_instalar.setEnabled(disponible and not control.revision_en_curso)
        self.btn_act_instalar.setToolTip("" if puede_instalar() else "La instalación automática solo funciona en el ejecutable")
        self.btn_act_buscar.setEnabled(not control.revision_en_curso)
        if control.ultima_revision is not None:
            self.lbl_act_revision.setText(f"Última revisión: {control.ultima_revision:%d-%m-%Y %H:%M}")

    def _instalar_actualizacion(self):
        from src.ui.dialogo_actualizacion import abrir_dialogo_actualizacion
        abrir_dialogo_actualizacion(self.window())

    @staticmethod
    def _titulo_pagina(lay, titulo, info):
        lbl = QLabel(titulo)
        lbl.setStyleSheet(f"font-size: 15px; font-weight: 600; color: {estilo_com.TINTA};")
        sub = QLabel(info)
        sub.setWordWrap(True)
        sub.setStyleSheet(f"color: {estilo_com.APAGADO}; font-size: 12px;")
        lay.addWidget(lbl)
        lay.addWidget(sub)

    @staticmethod
    def _boton(texto, icono, tipo="secundario"):
        """Botón de Configuración: 'primario' (azul Volvo), 'secundario' (borde fino) o 'texto'."""
        btn = QPushButton(texto)
        btn.setCursor(Qt.CursorShape.PointingHandCursor)
        if tipo == "primario":
            btn.setIcon(qta.icon(icono, color='white'))
            qss = (f"QPushButton {{ background: {VOLVO}; color: white; border: 1px solid {VOLVO}; }}"
                   f"QPushButton:hover {{ background: {VOLVO_HOVER}; border-color: {VOLVO_HOVER}; }}")
        elif tipo == "texto":
            btn.setIcon(qta.icon(icono, color=VOLVO))
            qss = (f"QPushButton {{ background: transparent; color: {VOLVO}; border: 1px solid transparent; }}"
                   f"QPushButton:hover {{ background: {VOLVO_SUAVE}; }}")
        else:
            btn.setIcon(qta.icon(icono, color=estilo_com.APAGADO))
            qss = (f"QPushButton {{ background: {estilo_com.SUPERFICIE}; color: {estilo_com.TINTA}; border: 1px solid {estilo_com.LINEA}; }}"
                   f"QPushButton:hover {{ background: {estilo_com.SUPERFICIE_2}; border-color: {estilo_com.APAGADO}; }}")
        btn.setStyleSheet(qss.replace("QPushButton {", "QPushButton { border-radius: 9px; padding: 8px 14px; "
                                      "font-size: 13px; font-weight: 600;", 1))
        return btn

    _ENLACES = (("Provisiones", "PROVISIONES_SHAREPOINT_URL", "fa5s.table"),
                ("Reporte Notas de Venta", "REPORTE_NOTAS_VENTAS_URL", "fa5s.file-alt"))

    @staticmethod
    def _marco_lista(nombre):
        marco = QFrame()
        marco.setObjectName(nombre)
        marco.setStyleSheet(f"QFrame#{nombre} {{ background: {estilo_com.SUPERFICIE}; border: 1px solid {estilo_com.LINEA}; border-radius: 12px; }}")
        lay = QVBoxLayout(marco)
        lay.setContentsMargins(0, 0, 0, 0)
        lay.setSpacing(0)
        return marco, lay

    @staticmethod
    def _linea():
        linea = QFrame()
        linea.setFixedHeight(1)
        linea.setStyleSheet(f"background: {estilo_com.LINEA_SUAVE}; border: none;")
        return linea

    @staticmethod
    def _resumen_url(url):
        if not url:
            return "Sin enlace configurado"
        partes_url = urlparse(url)
        if not partes_url.netloc:
            return url
        tramos = [t for t in partes_url.path.split("/") if t]
        texto = partes_url.netloc
        if tramos:
            texto += " › " + " › ".join(tramos[:2]) + (" › …" if len(tramos) > 2 else "")
        return texto

    def _crear_pagina_enlaces(self):
        pagina, lay = self._pagina()
        self._titulo_pagina(lay, "Enlaces web",
                            "Direcciones de SharePoint y Power BI que usa la aplicación. Pulsa «Editar» solo si cambió alguna.")
        marco, caja = self._marco_lista("listaEnlaces")
        self._filas_enlace = {}
        for i, (nombre, clave, icono) in enumerate(self._ENLACES):
            if i:
                caja.addWidget(self._linea())
            fila = QWidget()
            col = QVBoxLayout(fila)
            col.setContentsMargins(16, 12, 16, 12)
            col.setSpacing(10)
            principal = QHBoxLayout()
            principal.setSpacing(14)
            logo = QLabel()
            logo.setFixedSize(38, 38)
            logo.setAlignment(Qt.AlignmentFlag.AlignCenter)
            logo.setPixmap(qta.icon(icono, color=VOLVO).pixmap(QSize(16, 16)))
            logo.setStyleSheet(f"background: {VOLVO_SUAVE}; border-radius: 10px; border: none;")
            textos = QVBoxLayout()
            textos.setSpacing(0)
            lbl_nombre = QLabel(nombre)
            lbl_nombre.setStyleSheet(f"font-size: 13px; font-weight: 600; color: {estilo_com.TINTA}; {_SIN_FONDO}")
            lbl_dominio = QLabel("")
            lbl_dominio.setStyleSheet(f"font-size: 12px; color: {estilo_com.APAGADO}; {_SIN_FONDO}")
            textos.addWidget(lbl_nombre)
            textos.addWidget(lbl_dominio)
            pildora = QLabel("")
            btn_editar = self._boton("Editar", 'fa5s.pen', "secundario")
            principal.addWidget(logo)
            principal.addLayout(textos, 1)
            principal.addWidget(pildora)
            principal.addWidget(btn_editar)
            col.addLayout(principal)

            edicion = QWidget()
            fila_edicion = QHBoxLayout(edicion)
            fila_edicion.setContentsMargins(52, 0, 0, 0)
            txt_input = QLineEdit(self.config.get(clave, ""))
            txt_input.setPlaceholderText("https://...")
            btn_listo = self._boton("Listo", 'fa5s.check', "primario")
            fila_edicion.addWidget(txt_input, 1)
            fila_edicion.addWidget(btn_listo)
            edicion.setVisible(False)
            col.addWidget(edicion)
            caja.addWidget(fila)

            self.url_inputs[clave] = txt_input
            self._filas_enlace[clave] = (lbl_dominio, pildora, txt_input)
            btn_editar.clicked.connect(lambda _=False, e=edicion, t=txt_input: (e.setVisible(not e.isVisible()), t.setFocus()))
            btn_listo.clicked.connect(lambda _=False, e=edicion: e.setVisible(False))
            txt_input.textChanged.connect(lambda _t, c=clave: self._actualizar_enlace(c))
            self._actualizar_enlace(clave)
        lay.addWidget(marco)
        lay.addStretch()
        return pagina

    def _actualizar_enlace(self, clave):
        lbl_dominio, pildora, txt_input = self._filas_enlace[clave]
        url = txt_input.text().strip()
        lbl_dominio.setText(self._resumen_url(url))
        if not url:
            texto, fondo, color = "Sin enlace", estilo_com.SUPERFICIE_2, estilo_com.APAGADO
        elif url.startswith(("http://", "https://")):
            texto, fondo, color = "✓ Enlace válido", "#E3F6EE", SUCCESS
        else:
            texto, fondo, color = "✗ Debe empezar con https://", "#FDECEC", DANGER
        pildora.setText(texto)
        pildora.setStyleSheet(f"background: {fondo}; color: {color}; border-radius: 11px; padding: 3px 11px; "
                              "font-size: 11px; font-weight: 600; border: none;")
        borde = DANGER if (url and not url.startswith(("http://", "https://"))) else VOLVO
        txt_input.setStyleSheet(
            f"QLineEdit {{ background: {estilo_com.SUPERFICIE}; color: {estilo_com.TINTA}; border: 1px solid {borde}; "
            f"border-radius: 9px; padding: 8px 12px; font-size: 13px; }}")
        txt_input.setToolTip(url)

    def _crear_pagina_avanzado(self):
        pagina, lay = self._pagina()
        self._titulo_pagina(lay, "Herramientas de programador",
                            "Solo visible en modo programador (MODO_PROGRAMADOR en config/settings.py).")

        marco, caja = self._marco_lista("grupoRMJaviera")
        interior = QVBoxLayout()
        interior.setContentsMargins(16, 14, 16, 14)
        interior.setSpacing(8)
        lbl_rm = QLabel("Archivo de referencia (RM Javiera)")
        lbl_rm.setStyleSheet(f"font-weight: 600; font-size: 13px; color: {estilo_com.TINTA}; {_SIN_FONDO}")
        interior.addWidget(lbl_rm)
        fila_rm = QHBoxLayout()
        campo_rm = self.path_inputs["RM_JAVI_PATH"]
        campo_rm.setPlaceholderText("RM_JAVI_VOLVO CONNECT.xlsx")
        campo_rm.setStyleSheet(self._qss_campo())
        fila_rm.addWidget(campo_rm, 1)
        btn_rm = self._boton("Seleccionar archivo", 'fa5s.folder-open', "secundario")
        btn_rm.clicked.connect(self._elegir_rm_javiera)
        fila_rm.addWidget(btn_rm)
        interior.addLayout(fila_rm)
        self.lbl_estado_rm = QLabel("")
        interior.addWidget(self.lbl_estado_rm)
        caja.addLayout(interior)
        lay.addWidget(marco)
        campo_rm.textChanged.connect(lambda *_: self._refrescar_estado_rm())
        self._refrescar_estado_rm()

        fila = QHBoxLayout()
        fila.setSpacing(10)
        self.btn_sync_javi = self._boton("Sincronizar RM_JAVI", 'fa5s.cloud-download-alt', "primario")
        self.btn_sync_javi.clicked.connect(self._abrir_sync_rm_javi)
        self.btn_refresh = self._boton("Actualizar diagnóstico", 'fa5s.sync-alt', "secundario")
        self.btn_refresh.clicked.connect(self.force_refresh)
        self.btn_view_errors = self._boton("Ver errores", 'fa5s.search', "primario")
        self.btn_view_errors.setStyleSheet(
            f"QPushButton {{ background: {DANGER}; color: white; border: none; border-radius: 9px; padding: 8px 14px; "
            "font-size: 13px; font-weight: 600; }")
        self.btn_view_errors.setVisible(False)
        self.btn_view_errors.clicked.connect(self._show_error_dialog)
        for btn in (self.btn_sync_javi, self.btn_refresh, self.btn_view_errors):
            fila.addWidget(btn)
        fila.addStretch()
        lay.addLayout(fila)
        lay.addStretch()
        return pagina

    @staticmethod
    def _qss_campo():
        return (f"QLineEdit {{ background: {estilo_com.SUPERFICIE_2}; color: {estilo_com.TINTA}; border: 1px solid {estilo_com.LINEA}; "
                "border-radius: 10px; padding: 9px 12px; font-size: 13px; }"
                f"QLineEdit:focus {{ border-color: {VOLVO}; }}")

    def _elegir_rm_javiera(self):
        actual = self.path_inputs["RM_JAVI_PATH"].text().strip()
        ruta, _ = QFileDialog.getOpenFileName(
            self, "Selecciona el archivo de referencia (RM Javiera)",
            os.path.dirname(actual) if actual else "", "Archivos Excel (*.xlsx *.xlsm)")
        if ruta:
            self._asignar_ruta("RM_JAVI_PATH", ruta)

    def _refrescar_estado_rm(self):
        ruta = self.path_inputs["RM_JAVI_PATH"].text().strip()
        if not ruta:
            self.lbl_estado_rm.setText("Sin archivo seleccionado")
            color = TEXT_SECONDARY
        elif os.path.exists(ruta):
            self.lbl_estado_rm.setText("✓ Archivo encontrado")
            color = SUCCESS
        else:
            self.lbl_estado_rm.setText("✗ No se encuentra en la ruta indicada")
            color = DANGER
        self.lbl_estado_rm.setStyleSheet(f"color: {color}; font-size: 12px; font-weight: 600;")

    # ------------------------------------------------------------------
    # Carpeta de datos: una sola elección y el resto de rutas se completa solo
    # ------------------------------------------------------------------
    def _crear_grupo_carpeta_datos(self):
        grupo, lay = self._pagina()
        self._titulo_pagina(lay, "Carpeta de datos",
                            "Elige la carpeta una sola vez y la app encuentra cada archivo. "
                            "Guarda al final para aplicar los cambios.")

        carpeta_inicial = self.config.get("DATA_DIR") or os.path.dirname(self.config.get("EXCEL_PATH", "") or "")
        fila = QHBoxLayout()
        fila.setSpacing(8)
        self.data_dir_input = QLineEdit(carpeta_inicial)
        self.data_dir_input.setPlaceholderText("Carpeta que contiene RM_JAVI.csv y los demás archivos")
        self.data_dir_input.setStyleSheet(self._qss_campo())
        self.data_dir_input.addAction(qta.icon('fa5s.folder', color=estilo_com.APAGADO), QLineEdit.ActionPosition.LeadingPosition)
        fila.addWidget(self.data_dir_input, 1)
        btn_examinar = self._boton("Examinar", 'fa5s.folder-open', "secundario")
        btn_examinar.clicked.connect(self._elegir_carpeta_datos)
        btn_auto = self._boton("Autodetectar", 'fa5s.magic', "primario")
        btn_auto.clicked.connect(lambda: self._autodetectar())
        btn_abrir = self._boton("Abrir carpeta", 'fa5s.external-link-alt', "secundario")
        btn_abrir.clicked.connect(self._abrir_carpeta_datos)
        btn_ayuda = self._boton("¿Cómo me conecto?", 'fa5s.question-circle', "texto")
        btn_ayuda.clicked.connect(lambda: GuiaConexionDialog(self).exec())
        for btn in (btn_examinar, btn_auto, btn_abrir, btn_ayuda):
            fila.addWidget(btn)
        lay.addLayout(fila)

        fila_resumen = QHBoxLayout()
        self.lbl_resumen_rutas = QLabel("")
        fila_resumen.addWidget(self.lbl_resumen_rutas)
        fila_resumen.addStretch()
        btn_copia = self._boton("Actualizar copia de BBDD Clientes CM", 'fa5s.file-upload', "texto")
        btn_copia.clicked.connect(
            lambda: self._reemplazar_archivo_copia("POWERBI_EXCEL_PATH", self.path_inputs["POWERBI_EXCEL_PATH"]))
        fila_resumen.addWidget(btn_copia)
        lay.addLayout(fila_resumen)

        self.caja_estado_rutas = QVBoxLayout()
        self.caja_estado_rutas.setSpacing(12)
        lay.addLayout(self.caja_estado_rutas)
        lay.addStretch()
        return grupo

    def _elegir_carpeta_datos(self):
        actual = self.data_dir_input.text().strip()
        carpeta = QFileDialog.getExistingDirectory(
            self, "Selecciona la carpeta de datos", actual if os.path.isdir(actual) else "")
        if carpeta:
            self.data_dir_input.setText(os.path.normpath(carpeta))
            self._autodetectar()

    def _abrir_carpeta_datos(self):
        carpeta = self.data_dir_input.text().strip()
        if os.path.isdir(carpeta):
            os.startfile(carpeta)

    def _asignar_ruta(self, clave, ruta):
        ruta = os.path.normpath(ruta)
        if clave in self.path_inputs:
            self.path_inputs[clave].setText(ruta)
        else:
            self._rutas_extra[clave] = ruta

    def _autodetectar(self):
        """Completa todas las rutas a partir de la carpeta de datos; si está
        vacía o no existe, primero la busca desde la carpeta del usuario."""
        base = self.data_dir_input.text().strip()
        QApplication.setOverrideCursor(Qt.CursorShape.WaitCursor)
        try:
            if not os.path.isdir(base):
                base = rutas.buscar_carpeta_de_datos(os.path.expanduser("~")) or ""
            encontrados = rutas.descubrir_archivos(base, subir=2) if base else {}
        finally:
            QApplication.restoreOverrideCursor()
        if not encontrados:
            QMessageBox.information(
                self, "Autodetectar archivos",
                "No encontré los archivos de la aplicación. Elige a mano la carpeta que contiene RM_JAVI.csv.")
            return
        self.data_dir_input.setText(os.path.normpath(base))
        for clave, ruta in encontrados.items():
            self._asignar_ruta(clave, ruta)
        self._refrescar_estado_rutas()

    def _ruta_de(self, clave):
        if clave in self.path_inputs:
            return self.path_inputs[clave].text().strip()
        return self._rutas_extra.get(clave, "")

    @staticmethod
    def _tamano_legible(bytes_):
        if bytes_ >= 1024 * 1024:
            return f"{bytes_ / (1024 * 1024):.1f} MB"
        return f"{max(1, round(bytes_ / 1024))} KB"

    _GRUPOS_ARCHIVOS = (
        ("Base de datos principal", ("EXCEL_PATH", "PROVISION_EXCEL_PATH", "INSTALACIONES_PATH", "POWERBI_EXCEL_PATH")),
        ("Registros y contactos", ("LOG_CSV_PATH", "DELETED_LOG_EXCEL_PATH", "CONTACTOS_EXCEL_PATH", "CORREOS_EXCEL_PATH",
                                    "EJECUTIVOS_PATH", "METAS_VC_PATH", "CASOS_BAJA_CM_PATH", "FACTURACION_MANUAL_PATH")),
    )

    def _refrescar_estado_rutas(self):
        """Dos tablas de estado: una fila por archivo con check, nombre, fecha y tamaño."""
        caja = self.caja_estado_rutas
        while caja.count():
            item = caja.takeAt(0)
            if item.widget() is not None:
                item.widget().hide()
                item.widget().setParent(None)
                item.widget().deleteLater()
        faltan, total = [], 0
        for n_grupo, (titulo_grupo, claves) in enumerate(self._GRUPOS_ARCHIVOS):
            marco, lay_marco = self._marco_lista(f"grupoArchivos{n_grupo}")
            encabezado = QWidget()
            encabezado.setStyleSheet(f"background: {estilo_com.SUPERFICIE_2}; border-top-left-radius: 12px; border-top-right-radius: 12px;")
            fila_enc = QHBoxLayout(encabezado)
            fila_enc.setContentsMargins(16, 7, 16, 7)
            for texto, alinea in ((titulo_grupo.upper(), Qt.AlignmentFlag.AlignLeft), ("MODIFICADO · TAMAÑO", Qt.AlignmentFlag.AlignRight)):
                lbl = QLabel(texto)
                lbl.setAlignment(alinea)
                lbl.setStyleSheet(f"font-size: 10px; font-weight: 700; color: {estilo_com.APAGADO}; {_SIN_FONDO}")
                fila_enc.addWidget(lbl, 1)
            lay_marco.addWidget(encabezado)
            for clave in claves:
                etiqueta, nombre_estandar, obligatorio = rutas.ARCHIVOS_DE_DATOS[clave]
                ruta = self._ruta_de(clave)
                existe, tamano, modificado = rutas.estado_archivo(ruta) if ruta else (False, 0, 0.0)
                if obligatorio or ruta:
                    total += 1
                    if not existe:
                        faltan.append(etiqueta)
                if existe:
                    marca, color, fondo = "✓", SUCCESS, "#E3F6EE"
                    fecha = datetime.fromtimestamp(modificado).strftime('%d-%m-%Y %H:%M')
                    detalle = f"{fecha}   ·   {self._tamano_legible(tamano)}"
                elif ruta:
                    marca, color, fondo, detalle = "✕", DANGER, "#FDECEC", "No se encuentra en la ruta indicada"
                elif obligatorio:
                    marca, color, fondo, detalle = "✕", DANGER, "#FDECEC", "Sin ruta"
                else:
                    marca, color, fondo, detalle = "–", estilo_com.APAGADO, estilo_com.SUPERFICIE_2, "Sin ruta"
                lay_marco.addWidget(self._linea())
                fila_w = QWidget()
                fila_w.setMinimumHeight(50)
                fila = QHBoxLayout(fila_w)
                fila.setContentsMargins(16, 8, 16, 8)
                fila.setSpacing(12)
                lbl_marca = QLabel(marca)
                lbl_marca.setFixedSize(22, 22)
                lbl_marca.setAlignment(Qt.AlignmentFlag.AlignCenter)
                lbl_marca.setStyleSheet(f"background: {fondo}; color: {color}; border-radius: 11px; font-size: 11px; font-weight: 700; border: none;")
                textos = QVBoxLayout()
                textos.setSpacing(0)
                lbl_nombre = QLabel(etiqueta)
                lbl_nombre.setStyleSheet(f"font-size: 13px; font-weight: 600; color: {estilo_com.TINTA}; {_SIN_FONDO}")
                lbl_archivo = QLabel(os.path.basename(ruta) if ruta else nombre_estandar)
                lbl_archivo.setStyleSheet(f"font-size: 11px; color: {estilo_com.APAGADO}; {_SIN_FONDO}")
                textos.addWidget(lbl_nombre)
                textos.addWidget(lbl_archivo)
                lbl_detalle = QLabel(detalle)
                lbl_detalle.setAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
                lbl_detalle.setStyleSheet(f"font-size: 12px; color: {estilo_com.APAGADO if existe or not ruta else DANGER}; {_SIN_FONDO}")
                fila.addWidget(lbl_marca)
                fila.addLayout(textos, 1)
                if not obligatorio:
                    lbl_op = QLabel("Opcional")
                    lbl_op.setStyleSheet(f"font-size: 10px; font-weight: 600; color: {estilo_com.APAGADO}; border: 1px solid {estilo_com.LINEA}; "
                                         "border-radius: 6px; padding: 0 6px; background: transparent;")
                    fila.addWidget(lbl_op)
                fila.addWidget(lbl_detalle)
                lay_marco.addWidget(fila_w)
            caja.addWidget(marco)
        n_ok = total - len(faltan)
        if faltan:
            self.lbl_resumen_rutas.setText(f"✕ Faltan {len(faltan)} de {total}: " + ", ".join(faltan))
            self.lbl_estado_barra.setText(f"⚠ Faltan {len(faltan)} archivos")
            fondo_estado, color_estado = "#FEF3DC", "#B26B00"
            self.lbl_resumen_rutas.setStyleSheet(f"color: {DANGER}; font-weight: 600; font-size: 12px;")
        else:
            self.lbl_resumen_rutas.setText(f"✓ Los {total} archivos están disponibles")
            self.lbl_estado_barra.setText("●  Conectado a la carpeta de datos")
            fondo_estado, color_estado = "#E3F6EE", SUCCESS
            self.lbl_resumen_rutas.setStyleSheet(f"color: {SUCCESS}; font-weight: 600; font-size: 12px;")
        self.lbl_estado_barra.setStyleSheet(
            f"background: {fondo_estado}; color: {color_estado}; font-size: 12px; font-weight: 600; "
            "border-radius: 14px; padding: 6px 14px; border: none;")
        admin_estilo.actualizar_contador(self.tabbar, self._badges, 0, n_ok)

    @staticmethod
    def _vista_tema(oscuro):
        fondo, barra, linea, acento = (("#16202E", "#0B1A2B", "#33465F", "#5AA2DC") if oscuro
                                       else ("#F5F7FA", VOLVO, "#D3DBE6", VOLVO))
        vista = QFrame()
        vista.setFixedHeight(78)
        vista.setStyleSheet(f"QFrame {{ background: {fondo}; border: 1px solid {estilo_com.LINEA}; border-radius: 8px; }}")
        lay = QHBoxLayout(vista)
        lay.setContentsMargins(0, 0, 0, 0)
        lay.setSpacing(0)
        lateral = QFrame()
        lateral.setFixedWidth(34)
        lateral.setStyleSheet(f"background: {barra}; border: none; border-top-left-radius: 8px; border-bottom-left-radius: 8px;")
        lay.addWidget(lateral)
        cuerpo = QVBoxLayout()
        cuerpo.setContentsMargins(10, 10, 10, 10)
        cuerpo.setSpacing(6)
        for ancho, color in ((90, acento), (140, linea), (110, linea)):
            raya = QFrame()
            raya.setFixedSize(ancho, 7)
            raya.setStyleSheet(f"background: {color}; border: none; border-radius: 3px;")
            cuerpo.addWidget(raya)
        cuerpo.addStretch()
        lay.addLayout(cuerpo)
        lay.addStretch()
        return vista

    @staticmethod
    def _vista_tamano(punto_titulo, punto_texto):
        vista = QFrame()
        vista.setFixedHeight(78)
        vista.setStyleSheet(f"QFrame {{ background: {estilo_com.SUPERFICIE_2}; border: 1px solid {estilo_com.LINEA}; border-radius: 8px; }}")
        lay = QVBoxLayout(vista)
        lay.setContentsMargins(12, 8, 12, 8)
        lay.setSpacing(2)
        lay.addStretch()
        titulo = QLabel("Aa · Contratos activos")
        titulo.setStyleSheet(f"font-size: {punto_titulo}px; font-weight: 600; color: {estilo_com.TINTA}; {_SIN_FONDO}")
        texto = QLabel("Así se leerá la aplicación")
        texto.setStyleSheet(f"font-size: {punto_texto}px; color: {estilo_com.APAGADO}; {_SIN_FONDO}")
        lay.addWidget(titulo)
        lay.addWidget(texto)
        return vista

    def _crear_grupo_escala(self):
        grupo, lay = self._pagina()
        gestor = GestorTema.instancia()

        self._titulo_pagina(lay, "Tema", "Colores de la aplicación. El cambio se aplica al instante y se recuerda.")
        fila_tema = QHBoxLayout()
        fila_tema.setSpacing(14)
        self._tarjetas_tema = {}
        for modo, titulo, detalle in ((CLARO, "Claro", "Fondo blanco, ideal con buena luz."),
                                      (OSCURO, "Oscuro", "Menos brillo para jornadas largas.")):
            tarjeta = TarjetaOpcion(titulo, detalle, self._vista_tema(modo == OSCURO))
            tarjeta.clicked.connect(lambda m=modo: gestor.aplicar(m))
            self._tarjetas_tema[modo] = tarjeta
            fila_tema.addWidget(tarjeta, 1)
        fila_tema.addStretch(1)
        lay.addLayout(fila_tema)
        gestor.tema_cambiado.connect(self._marcar_tema)
        self._marcar_tema(gestor.modo)

        lay.addSpacing(6)
        self._titulo_pagina(lay, "Tamaño de la interfaz",
                            "Agranda o achica letras, tablas y botones de toda la aplicación. "
                            "Al cambiarlo, la aplicación se reinicia para aplicarlo.")
        nivel_actual = get_nivel_escala()
        fila = QHBoxLayout()
        fila.setSpacing(14)
        tamanos = {"pequeno": (11, 10), "mediano": (13, 11), "grande": (16, 13)}
        for clave in NIVELES_ESCALA:
            nombre, factor = NIVELES_ESCALA[clave]
            tarjeta = TarjetaOpcion(nombre, f"{round(factor * 100)} % del tamaño base.", self._vista_tamano(*tamanos[clave]))
            tarjeta.marcar(clave == nivel_actual)
            if clave != nivel_actual:
                tarjeta.clicked.connect(lambda c=clave: self._cambiar_escala(c))
            else:
                tarjeta.setCursor(Qt.CursorShape.ArrowCursor)
            fila.addWidget(tarjeta, 1)
        lay.addLayout(fila)
        lay.addStretch()
        return grupo

    def _marcar_tema(self, modo):
        for clave, tarjeta in self._tarjetas_tema.items():
            tarjeta.marcar(clave == modo)

    def _cambiar_escala(self, nivel):
        nombre = NIVELES_ESCALA[nivel][0]
        reply = QMessageBox.question(self, "Cambiar tamaño de la interfaz",
            f"La aplicación se reiniciará para aplicar el tamaño \"{nombre}\".\n\n"
            "Asegúrate de no tener cambios sin guardar. ¿Continuar?",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No, QMessageBox.StandardButton.No)
        if reply != QMessageBox.StandardButton.Yes:
            return
        set_nivel_escala(nivel)
        # close() pasa por MainWindow.closeEvent: guarda geometría/columnas y detiene los hilos.
        self.window().close()
        QProcess.startDetached(sys.executable, sys.argv)
        QApplication.quit()

    def _validate_path_input(self, txt_input):
        ruta = txt_input.text().strip()
        if not ruta or os.path.exists(ruta):
            txt_input.setStyleSheet("""
                QLineEdit { border: 1.5px solid #10B981; background-color: #F0FDF4; color: #166534; border-radius: 6px; padding: 6px 10px; font-weight: bold; font-size: 12px; }
            """)
            txt_input.setToolTip(" Archivo Excel encontrado en esta ruta." if ruta else "Ruta local vacía.")
        else:
            txt_input.setStyleSheet("""
                QLineEdit { border: 2px solid #EF4444; background-color: #FEF2F2; color: #991B1B; border-radius: 6px; padding: 6px 10px; font-weight: bold; font-size: 12px; }
            """)
            txt_input.setToolTip(" ALERTA: No se encontró el archivo en esta ruta local.")

    def _create_status_card(self, title, icon_name, icon_color, initial_value):
        frame = QFrame()
        frame.setObjectName("tarjetaEstadoConfig")
        frame.setStyleSheet("QFrame#tarjetaEstadoConfig { background: transparent; border: none; }")
        layout = QHBoxLayout(frame)
        layout.setContentsMargins(6, 2, 6, 2)
        layout.setSpacing(8)

        icon_lbl = QLabel()
        icon_lbl.setFixedSize(26, 26)
        icon_lbl.setAlignment(Qt.AlignmentFlag.AlignCenter)
        icon_lbl.setPixmap(qta.icon(icon_name, color=icon_color).pixmap(QSize(13, 13)))
        tono = QColor(icon_color)
        icon_lbl.setStyleSheet(
            f"background-color: rgba({tono.red()},{tono.green()},{tono.blue()},34); border-radius: 8px; border: none;")

        title_lbl = QLabel(title)
        title_lbl.setStyleSheet(f"color: {estilo_com.APAGADO}; font-size: 10px; font-weight: 600; border: none; background: transparent;")
        val_lbl = QLabel(initial_value)
        val_lbl.setStyleSheet(f"color: {estilo_com.TINTA}; font-size: 12px; font-weight: 800; border: none; background: transparent;")
        texto = QVBoxLayout()
        texto.setSpacing(0)
        texto.addWidget(title_lbl)
        texto.addWidget(val_lbl)

        layout.addWidget(icon_lbl)
        layout.addLayout(texto, 1)
        return {'frame': frame, 'val_lbl': val_lbl, 'icon_color': icon_color}

    def _abrir_sync_rm_javi(self):
        # Solo modo desarrollador: en producción esta opción desaparece de la app.
        from src.ui.comercial import SyncRMJaviDialog
        if SyncRMJaviDialog(self).exec() == QDialog.DialogCode.Accepted:
            self.data_manager.data_ready.emit()

    def update_health_panel(self):
        if hasattr(self.data_manager, 'get_system_health'):
            health = self.data_manager.get_system_health()
            self.card_db['val_lbl'].setText(health.get('db_status', 'N/A'))
            self.card_rows['val_lbl'].setText(str(health.get('db_rows', 0)))
            self.card_sync['val_lbl'].setText(health.get('last_modified', 'N/A'))
            self.card_ram['val_lbl'].setText(health.get('ram_usage', 'N/A'))
            self.card_load_time['val_lbl'].setText(health.get('load_time', 'N/A'))
            self.card_logs['val_lbl'].setText(health.get('log_status', 'N/A'))
            self.card_logs['val_lbl'].setStyleSheet(f"color: {health.get('log_color', TEXT_PRIMARY)}; font-size: 12px; font-weight: 800; border: none; background: transparent;")
            self.last_errors = health.get('recent_errors', [])
            self.btn_view_errors.setVisible(len(self.last_errors) > 0)

    def force_refresh(self):
        self.btn_refresh.setEnabled(False)
        self.btn_refresh.setText(" Calculando...")
        self.update_health_panel()
        self.btn_refresh.setText(" Actualizar Diagnóstico")
        self.btn_refresh.setEnabled(True)

    def _show_error_dialog(self):
        dialog = QDialog(self)
        dialog.setWindowTitle("Diagnóstico Detallado de Errores")
        dialog.resize(800, 500)
        dialog.setStyleSheet(f"background-color: {SURFACE};")
        
        layout = QVBoxLayout(dialog)
        lbl_info = QLabel("El sistema ha detectado problemas recientes:")
        lbl_info.setStyleSheet("font-size: 14px; font-weight: bold; color: #1F2937;")
        layout.addWidget(lbl_info)
        
        txt_display = QTextEdit()
        txt_display.setReadOnly(True)
        txt_display.setStyleSheet("background-color: #111827; color: #F3F4F6; font-family: Consolas; font-size: 12px; padding: 10px; border-radius: 6px;")
        
        reporte = ""
        for err in self.last_errors:
            partes = err.split(" - ERROR - ")
            if len(partes) > 1:
                reporte += f"[{partes[0]}] FALLO:\n> {partes[1]}\n\n"
            else:
                reporte += f"[Momento desconocido] FALLO:\n> {err}\n\n"
        txt_display.setText(reporte)
        layout.addWidget(txt_display)
        
        footer = QHBoxLayout()
        btn_resolve = QPushButton(" Limpiar Registro")
        btn_resolve.setIcon(qta.icon('fa5s.broom', color='white'))
        btn_resolve.setStyleSheet(f"background-color: {SUCCESS}; color: white; border-radius: 4px; font-weight: bold; padding: 5px 15px;")
        btn_resolve.clicked.connect(lambda: self._clear_error_logs(dialog))
        
        btn_close = QPushButton("Cerrar")
        ThemeManager.apply_secondary_button(btn_close)
        btn_close.clicked.connect(dialog.accept)
        
        footer.addWidget(btn_resolve)
        footer.addStretch()
        footer.addWidget(btn_close)
        layout.addLayout(footer)
        dialog.exec()

    def _clear_error_logs(self, dialog):
        log_path = os.path.join(os.getcwd(), 'logs', 'sistema.log')
        if os.path.exists(log_path):
            try:
                with open(log_path, 'w', encoding='utf-8') as f: f.write("")
                QMessageBox.information(self, "Registro Limpio", "El sistema vuelve a estar en estado Saludable.")
            except Exception as e:
                QMessageBox.warning(self, "Aviso", f"Error: {e}")
        dialog.accept()        
        self.force_refresh()   

    def _browse_file(self, config_key, text_input):
        # LOG_CSV_PATH y CASOS_BAJA_CM_PATH apuntan a CSV (historial y casos
        # de Revision CM, ver cm_bajas_service.py); CORREOS_EXCEL_PATH tambien
        # desde la migracion B-5.2 (18-09-2026, Correos.xlsx -> Correos.csv) --
        # el nombre de la clave de configuracion se mantiene igual para no
        # romper app_config.json existentes, solo cambio el formato del
        # archivo al que apunta. El resto sigue siendo *.xlsx/*.xlsm como siempre.
        campos_csv = ("LOG_CSV_PATH", "CASOS_BAJA_CM_PATH", "CORREOS_EXCEL_PATH", "METAS_VC_PATH")
        filtro = "Archivo CSV (*.csv)" if config_key in campos_csv else "Archivos Excel (*.xlsx *.xlsm)"
        ruta, _ = QFileDialog.getOpenFileName(self, "Seleccionar archivo", os.path.dirname(text_input.text()) if text_input.text() else "", filtro)
        if ruta:
            text_input.setText(ruta.replace("/", "\\"))
            self._validate_path_input(text_input)

    def _reemplazar_archivo_copia(self, config_key, text_input):
        """Reemplaza el CONTENIDO del archivo ya configurado en `config_key`
        por una copia nueva elegida por el usuario, sin cambiar la ruta
        guardada (a diferencia de _browse_file). Pensado para BBDD Clientes
        CM: la ruta apunta siempre al mismo lugar (SharePoint/OneDrive
        sincronizado), y periódicamente llega una copia nueva de Power BI
        que hay que dejar ahí mismo."""
        ruta_actual = text_input.text().strip()
        if not ruta_actual:
            QMessageBox.warning(self, "Ruta no configurada",
                "Primero elige la carpeta de datos (o usa 'Autodetectar archivos') y guarda los cambios, "
                "antes de poder reemplazar su contenido.")
            return

        ruta_nueva, _ = QFileDialog.getOpenFileName(self, "Selecciona la copia nueva",
            os.path.dirname(ruta_actual) if os.path.exists(os.path.dirname(ruta_actual)) else "",
            "Archivos Excel (*.xlsx *.xlsm)")
        if not ruta_nueva:
            return

        if os.path.abspath(ruta_nueva) == os.path.abspath(ruta_actual):
            QMessageBox.warning(self, "Mismo archivo", "El archivo seleccionado es el mismo que ya está configurado.")
            return

        # Validacion ligera: confirma que la copia nueva es un Excel legible
        # y que su primera hoja (misma que usa powerbi_cache_service, NO
        # wb.active) tiene columna 'CDB' -- evita reemplazar el archivo real
        # por algo con el formato equivocado.
        try:
            wb = openpyxl.load_workbook(ruta_nueva, read_only=True, data_only=True)
            ws = wb[wb.sheetnames[0]]
            primera_fila = next(ws.iter_rows(values_only=True, max_row=1), ())
            wb.close()
            headers = [str(h).strip() if h is not None else "" for h in primera_fila]
        except Exception as e:
            QMessageBox.critical(self, "Archivo inválido", f"No se pudo leer el archivo seleccionado como Excel:\n\n{e}")
            return

        if "CDB" not in headers:
            QMessageBox.critical(self, "Formato inesperado",
                "El archivo seleccionado no tiene una columna 'CDB' en su primera hoja.\n\n"
                "No parece ser una copia válida de BBDD Clientes CM. No se reemplazó nada.")
            return

        reply = QMessageBox.question(self, "Confirmar actualización",
            f"Esto reemplazará el contenido de:\n{ruta_actual}\n\ncon:\n{ruta_nueva}\n\n"
            "Se guardará una copia de seguridad del archivo actual antes de reemplazarlo.\n\n¿Continuar?",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No, QMessageBox.StandardButton.No)
        if reply != QMessageBox.StandardButton.Yes:
            return

        try:
            if os.path.exists(ruta_actual):
                writer_backup = ExcelWriter()
                writer_backup._crear_backup(ruta_actual)
                if not writer_backup.last_backup_ok:
                    QMessageBox.critical(self, "Actualización cancelada",
                        "No se pudo crear un respaldo del archivo actual antes de reemplazarlo.\n\n"
                        "Para evitar perder el archivo actual sin posibilidad de deshacerlo, "
                        "la actualización fue cancelada. No se modificó ningún archivo.")
                    return

            # shutil.copy (NO copy2): necesitamos que el archivo destino
            # quede con fecha de modificación de AHORA, no la fecha original
            # de la copia nueva -- el cache CSV de powerbi_cache_service se
            # invalida comparando esa fecha contra la del cache existente.
            shutil.copy(ruta_nueva, ruta_actual)

            QProcess.startDetached(sys.executable, sys.argv)
            QApplication.quit()
        except Exception as e:
            QMessageBox.critical(self, "Error", f"No se pudo reemplazar el archivo:\n\n{e}")

    def _restore_backup(self):
        backup_dir = os.path.join(os.getcwd(), 'data', 'backups')
        if not os.path.exists(backup_dir): return QMessageBox.information(self, "Sin Backups", "Aún no hay copias de seguridad.")

        # Fase 2 (migracion PruebasFM.xlsx -> .csv, 23-09-2026): el filtro y
        # la extension esperada se derivan de EXCEL_PATH en vez de
        # hardcodear ".xlsx", asi soporta ambos formatos sin duplicar esta
        # pantalla. Filtrar por la extension ACTUAL de EXCEL_PATH (no ambas
        # a la vez) evita mezclar un backup .csv con un destino .xlsx o
        # viceversa -- un backup nunca cambia de formato al restaurarse.
        extension_actual = os.path.splitext(EXCEL_PATH)[1] or ".xlsx"
        ruta_backup, _ = QFileDialog.getOpenFileName(
            self, "Selecciona copia de seguridad", backup_dir,
            f"Copias de seguridad (*{extension_actual})",
        )
        if ruta_backup:
            # La carpeta de backups mezcla copias de PruebasFM, Provisiones,
            # Contactos, etc. Esta funcion solo restaura PruebasFM (xlsx o
            # csv, segun EXCEL_PATH), asi que se bloquea si el backup
            # elegido pertenece a otro archivo (evita sobrescribir la base
            # de datos con la estructura equivocada).
            nombre_base_excel = os.path.splitext(os.path.basename(EXCEL_PATH))[0]
            nombre_backup = os.path.basename(ruta_backup)
            if not (nombre_backup.startswith(f"{nombre_base_excel}_backup_")
                    or (nombre_base_excel in ("RM_JAVI", "Registro Maestro")
                        and nombre_backup.startswith(("RM_JAVI_backup_", "Registro Maestro_backup_", "PruebasFM_backup_")))):
                QMessageBox.critical(self, "Backup incompatible",
                    f"El archivo seleccionado ('{nombre_backup}') no es una copia de seguridad de la base de datos principal ({os.path.basename(EXCEL_PATH)}).\n\n"
                    "Restaurarlo aquí sobrescribiría la base de datos con datos de otro archivo. Selecciona un backup que comience con "
                    f"'{nombre_base_excel}_backup_'.")
                return

            reply = QMessageBox.question(self, 'Advertencia Crítica', f"¿Restaurar?\n\nArchivo: {nombre_backup}\nEsto sobrescribirá tu base de datos actual.", QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No, QMessageBox.StandardButton.No)
            if reply == QMessageBox.StandardButton.Yes:
                try:
                    # Antes de sobrescribir EXCEL_PATH con el backup elegido,
                    # se respalda el estado ACTUAL (mismo mecanismo que ya
                    # usa ExcelWriter en el resto de la app) para poder
                    # deshacer la restauracion si hiciera falta. Si ese
                    # backup previo falla, se cancela todo sin tocar el
                    # archivo real.
                    writer_backup = ExcelWriter()
                    writer_backup._crear_backup(EXCEL_PATH)
                    if not writer_backup.last_backup_ok:
                        QMessageBox.critical(self, "Restauración cancelada",
                            "No se pudo crear un respaldo del estado actual antes de restaurar.\n\n"
                            "Para evitar perder el estado actual sin posibilidad de deshacerlo, "
                            "la restauración fue cancelada. No se modificó ningún archivo.")
                        return

                    shutil.copy2(ruta_backup, EXCEL_PATH)
                    QProcess.startDetached(sys.executable, sys.argv)
                    QApplication.quit()
                except Exception as e: QMessageBox.critical(self, "Error", str(e))

    def _save_changes(self):
        hay_error = False
        
        for key, txt_input in self.path_inputs.items(): 
            ruta = txt_input.text().strip()
            if ruta and not os.path.exists(ruta):
                hay_error = True
                
        for key, txt_input in self.url_inputs.items():
            url = txt_input.text().strip()
            if url and not (url.startswith("http://") or url.startswith("https://")):
                hay_error = True
                
        if hay_error:
            reply = QMessageBox.warning(self, "Rutas o URLs Inválidas", 
                                        "Atención: Una o más casillas tienen datos incorrectos (marcadas en rojo).\n\n"
                                        "Si guardas y reinicias, el sistema puede fallar al cargar esas fuentes.\n¿Deseas guardar de todos modos?",
                                        QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
                                        QMessageBox.StandardButton.No)
            if reply == QMessageBox.StandardButton.No: return

        for key, txt_input in self.path_inputs.items():
            self.config[key] = txt_input.text().strip()
        self.config.update(self._rutas_extra)
        self.config["DATA_DIR"] = self.data_dir_input.text().strip()
        for key, txt_input in self.url_inputs.items():
            self.config[key] = txt_input.text().strip()
            
        save_config(self.config)
        
        msg = QMessageBox(self)
        msg.setWindowTitle("Configuración Guardada")
        msg.setIcon(QMessageBox.Icon.Information)
        msg.setText("Las rutas y enlaces se han actualizado exitosamente.\n\nLa aplicación se reiniciará automáticamente para aplicar los cambios.")
        msg.setStandardButtons(QMessageBox.StandardButton.Ok)
        msg.exec()

        QProcess.startDetached(sys.executable, sys.argv)
        QApplication.quit()
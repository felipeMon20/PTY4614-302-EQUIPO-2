"""Excel 'Detalle de facturación EDP – Volvo Connect' de una flota de
Facturación Manual: mismo contenido que el PDF (facturacion_manual_pdf), pero con
los totales como fórmulas que dependen de la UF y del precio por unidad."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from src.services.facturacion_manual_service import MESES_ES
from src.ui.facturacion_manual_pdf import FIRMA

_NEGRO = PatternFill("solid", fgColor="000000")
_ENTRADA = PatternFill("solid", fgColor="FFF8DC")
_FINO = Side(style="thin", color="000000")
_BORDE = Border(left=_FINO, right=_FINO, top=_FINO, bottom=_FINO)
_FORMATO_UF = '#,##0.00'
_FORMATO_CLP = '"$ "#,##0'
_ANCHOS = {"A": 46, "B": 16, "C": 24, "D": 14, "E": 14}


def nombre_sugerido(cliente, anio=None, mes=None):
    from src.ui.facturacion_manual_pdf import nombre_sugerido as nombre_pdf
    return nombre_pdf(cliente, anio, mes)[:-4] + ".xlsx"


def _encabezado(ws, fila, textos):
    for col, texto in enumerate(textos, start=1):
        c = ws.cell(row=fila, column=col, value=texto)
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = _NEGRO
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = _BORDE


def generar_excel(ruta, datos):
    """`datos` es el mismo diccionario que recibe facturacion_manual_pdf.generar_pdf
    más `precio_uf` (precio por unidad). La cantidad es COUNTA de los VIN cuando
    coincide con `unidades_n`; si no, queda como valor ingresado."""
    mes, anio = datos["mes"], datos["anio"]
    filas = datos["filas"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Detalle"
    base = Font(name="Arial", size=10)
    negrita = Font(name="Arial", size=10, bold=True)
    for letra, ancho in _ANCHOS.items():
        ws.column_dimensions[letra].width = ancho

    ws.merge_cells("A1:E1")
    ws["A1"] = "DETALLE DE FACTURACIÓN EDP – VOLVO CONNECT"
    ws["A1"].font = Font(name="Arial", size=13, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells("A2:E2")
    ws["A2"] = f"Periodo: {MESES_ES[mes - 1]} {anio}"
    ws["A2"].font = negrita
    ws["A2"].alignment = Alignment(horizontal="center")

    ws["A4"], ws["A5"] = "Cliente:", "Fecha de emisión:"
    ws["B4"], ws["B5"] = datos["cliente"], datos["fecha_emision"]
    for celda in ("A4", "A5"):
        ws[celda].font = negrita
    for celda in ("B4", "B5"):
        ws[celda].font = base
        ws[celda].alignment = Alignment(horizontal="left")

    ws.merge_cells("A7:E7")
    ws["A7"] = (f"Junto con saludar, por medio del presente se informa el detalle correspondiente a la "
                f"facturación de los servicios EDP asociados al período de "
                f"{MESES_ES[mes - 1].lower()} de {anio}.")
    ws["A7"].font = base
    ws["A7"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[7].height = 30

    ws["A9"] = "1.  Detalle de Servicios Facturados"
    ws["A9"].font = negrita
    _encabezado(ws, 10, ["Resumen precios", ""])
    ws.merge_cells("A10:B10")

    fila_detalle = 18
    primera, ultima = fila_detalle + 1, fila_detalle + len(filas)
    cantidad = f"=COUNTA(C{primera}:C{ultima})" if datos["unidades_n"] == len(filas) and filas else datos["unidades_n"]
    fijo = bool(datos.get("precio_fijo"))
    if fijo:
        resumen = [
            ("Unidades conectadas", cantidad, "0", False),
            ("Modalidad", "Precio fijo mensual", "@", False),
            ("Precio Neto Total CLP", datos["neto_clp"], _FORMATO_CLP, False),
        ]
    else:
        resumen = [
            ("Unidades conectadas", cantidad, "0", False),
            (datos.get("uf_etiqueta", "UF hoy"), datos["uf"] or None, _FORMATO_UF, True),
            ("Precio UF por unidad", datos["precio_uf"], _FORMATO_UF, True),
            ("Precio Neto Total UF", "=ROUND(B11*B13,2)", _FORMATO_UF, False),
            ("Precio Neto Total CLP", '=IF(B12>0,ROUND(B11*B13*B12,0),"-")', _FORMATO_CLP, False),
        ]
    for i, (etiqueta, valor, formato, entrada) in enumerate(resumen):
        fila = 11 + i
        a, b = ws.cell(row=fila, column=1, value=etiqueta), ws.cell(row=fila, column=2, value=valor)
        a.font, b.font = base, base
        a.border = b.border = _BORDE
        b.number_format = formato
        b.alignment = Alignment(horizontal="right")
        if entrada:
            b.fill = _ENTRADA
            b.font = Font(name="Arial", size=10, color="0000FF")
    if not fijo:
        ws["C14"] = "Celdas amarillas: UF y precio editables; los totales se recalculan."
        ws["C14"].font = Font(name="Arial", size=9, italic=True, color="666666")

    ws["A17"] = "2.  Detalle"
    ws["A17"].font = negrita
    _encabezado(ws, fila_detalle, ["Flota", "Rut", "VIN", "Plan", "Precio UF"])
    for i, (flota, rut, vin, plan, _precio) in enumerate(filas):
        fila = primera + i
        for col, valor in enumerate((flota, rut, vin, plan, "-" if fijo else "=$B$13"), start=1):
            c = ws.cell(row=fila, column=col, value=valor)
            c.font = base
            c.border = _BORDE
            c.alignment = Alignment(horizontal="left" if col == 1 else "center")
        ws.cell(row=fila, column=2).number_format = "@"
        ws.cell(row=fila, column=3).number_format = "@"
        ws.cell(row=fila, column=5).number_format = _FORMATO_UF

    fila = ultima + 2
    ws.cell(row=fila, column=1, value="3. Consideraciones").font = negrita
    consideraciones = [
        "Los precios indicados corresponden a los servicios efectivamente prestados durante el período señalado.",
        "El impuesto al valor agregado (IVA) se aplica conforme a la normativa vigente en Chile.",
        "Este documento tiene carácter informativo y respalda el detalle de la facturación emitida.",
    ]
    for texto in consideraciones:
        fila += 1
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=5)
        ws.cell(row=fila, column=1, value="•  " + texto).font = base
        ws.cell(row=fila, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[fila].height = 28
    fila += 2
    ws.cell(row=fila, column=1, value="Sin otro particular, saluda atentamente,").font = base
    for i, texto in enumerate(FIRMA):
        fila += 1 if i else 2
        ws.cell(row=fila, column=1, value=texto).font = negrita if i < 2 else base

    ws.freeze_panes = None
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{fila_detalle}:{fila_detalle}"
    wb.save(ruta)

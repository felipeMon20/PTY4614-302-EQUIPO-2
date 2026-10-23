"""Plantilla base del documento de Facturacion Manual (fase 'Crear plantilla
base', ver Notas del proyecto.txt, sección 7). Solo estructura visual del DOCX: textos
fijos y tabla de flota reconstruidos a partir de los documentos reales
encontrados en la carpeta de Facturacion Manual (ver informe de
investigacion, 19-09-2026) y de los campos ya confirmados en df_master
(construir_contexto_empresa, facturacion_manual_service.py).

Sin calculos ni conexion a 'Facturación Manual - BBDD.xlsx': todo campo
financiero (Valor UF, Precio Total UF, Monto Total CLP, N° OC, HES) queda
como el texto literal PLACEHOLDER_PENDIENTE hasta que se resuelva con
Javiera (reunion 22-09-2026, ver informe de viabilidad). No usa python-docx
como motor de reemplazo de variables (eso es tarea de la fase siguiente,
'Preparar generación DOCX/PDF'): aqui solo se arma el documento con
docx.Document directamente, con tokens {{...}} como marcadores de texto
para esa fase futura.

Nota (ver informe de investigacion): en los documentos reales, este tipo de
carta EDP solo existe hoy para clientes Tipo=OC en la BBDD real; para
Tipo=FACTURA MANUAL no se encontro ningun documento equivalente. Esta
plantilla base es unica y generica a proposito -- si Javiera confirma que
ambos tipos de cliente deben recibir el mismo documento, no hace falta
tocarla; si confirma que deben ser distintos, se decide recien ahi.
"""

import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

PLACEHOLDER_PENDIENTE = "PENDIENTE JAVIERA"

TITULO = "DETALLE DE FACTURACIÓN – VOLVO CONNECT"

TEXTO_SALUDO = "Estimado(a) cliente:"

TEXTO_INTRODUCCION = (
    "Junto con saludar, por medio del presente se informa el detalle "
    "correspondiente a la facturación de los servicios Volvo Connect."
)

TEXTO_IVA = "El impuesto al valor agregado (IVA) se aplica conforme a la normativa vigente en Chile."
TEXTO_CARACTER_INFORMATIVO = "Este documento tiene carácter informativo y respalda el detalle de la facturación emitida."

TEXTO_DESPEDIDA = "Sin otro particular, saluda atentamente,"
FIRMA_NOMBRE = "Javiera González Maldonado"
FIRMA_CARGO = "Encargada de Servicios Conectados"

# Filas del cuadro "Resumen de Facturación": (etiqueta, valor). Las filas
# financieras usan PLACEHOLDER_PENDIENTE a proposito -- no se calculan aqui.
FILAS_RESUMEN_FACTURACION = [
    ("Unidades conectadas", "{{CANTIDAD_UNIDADES}}"),
    ("Valor UF", PLACEHOLDER_PENDIENTE),
    ("Precio Total UF", PLACEHOLDER_PENDIENTE),
    ("Monto Total CLP", PLACEHOLDER_PENDIENTE),
    ("N° OC", PLACEHOLDER_PENDIENTE),
    ("HES", PLACEHOLDER_PENDIENTE),
]

# Encabezados de la tabla de flota: solo columnas confirmadas en df_master
# (ver COLUMNAS_CONTEXTO_UNIDAD en facturacion_manual_service.py). "Flota" se
# muestra por fila porque asi aparece en los documentos reales encontrados
# (repite el nombre de la empresa/flota en cada linea de la tabla).
ENCABEZADOS_TABLA_FLOTA = [
    "Flota", "RUT", "VIN", "Chasis", "Plan", "Estado Comercial", "Estado de Activación", "Modelo",
]

TOKENS_FILA_FLOTA = [
    "{{FLOTA}}", "{{RUT}}", "{{VIN}}", "{{CHASIS}}", "{{PLAN}}",
    "{{ESTADO_COMERCIAL}}", "{{ESTADO_ACTIVACION}}", "{{MODELO}}",
]

NOTA_FILA_REPETIDA = "Esta fila se repite por cada unidad de la empresa (ver contexto de facturación)."


def construir_plantilla_base():
    """Arma el Document base con docx.Document (python-docx) y lo devuelve
    en memoria. No escribe a disco -- eso lo decide quien llama (ver
    generar_plantilla_base)."""
    documento = docx.Document()

    titulo = documento.add_heading(TITULO, level=1)
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p_cliente = documento.add_paragraph()
    p_cliente.add_run("Cliente: {{RAZON_SOCIAL}}").bold = True
    documento.add_paragraph("RUT: {{RUT}}")
    documento.add_paragraph("Fecha de emisión: {{FECHA_EMISION}}")

    documento.add_paragraph(TEXTO_SALUDO)
    documento.add_paragraph(TEXTO_INTRODUCCION)

    documento.add_heading("1. Resumen de Facturación", level=2)
    tabla_resumen = documento.add_table(rows=0, cols=2)
    tabla_resumen.style = "Table Grid"
    for etiqueta, valor in FILAS_RESUMEN_FACTURACION:
        fila = tabla_resumen.add_row().cells
        fila[0].text = etiqueta
        fila[1].text = valor

    documento.add_heading("2. Detalle de Unidades", level=2)
    tabla_flota = documento.add_table(rows=0, cols=len(ENCABEZADOS_TABLA_FLOTA))
    tabla_flota.style = "Table Grid"
    fila_encabezado = tabla_flota.add_row().cells
    for i, encabezado in enumerate(ENCABEZADOS_TABLA_FLOTA):
        fila_encabezado[i].text = encabezado
        for run in fila_encabezado[i].paragraphs[0].runs:
            run.bold = True
    fila_placeholder = tabla_flota.add_row().cells
    for i, token in enumerate(TOKENS_FILA_FLOTA):
        fila_placeholder[i].text = token

    nota = documento.add_paragraph()
    nota_run = nota.add_run(NOTA_FILA_REPETIDA)
    nota_run.italic = True
    nota_run.font.size = Pt(8)

    documento.add_heading("3. Consideraciones", level=2)
    documento.add_paragraph(TEXTO_IVA)
    documento.add_paragraph(TEXTO_CARACTER_INFORMATIVO)

    documento.add_paragraph()
    documento.add_paragraph(TEXTO_DESPEDIDA)
    documento.add_paragraph(FIRMA_NOMBRE)
    documento.add_paragraph(FIRMA_CARGO)

    return documento


def generar_plantilla_base(output_path):
    """Genera el archivo .docx de la plantilla base en output_path (crea
    carpetas intermedias si no existen). Devuelve output_path."""
    import os
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    documento = construir_plantilla_base()
    documento.save(output_path)
    return output_path
